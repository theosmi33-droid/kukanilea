#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
KUKANILEA Core (v2.4 - A2 Multi-Tenant + More Formats) — DROP-IN replacement
============================================================================

Das ist dein `kukanilea_core.py` als kompletter, kopierfertiger Code.

Wichtig (Kompatibilität):
- Exportiert die gleichen Funktionen/Variablen wie vorher (Contract).
- Behält **TOPHANDWERK_*** ENV-Variablen für Abwärtskompatibilität.
- Zusätzlich werden **KUKANILEA_*** ENV-Variablen als Alias akzeptiert (haben Vorrang).

Hinweis:
- In `kukanilea_upload.py` musst du nur den Import anpassen:
    import kukanilea_core as core
  (statt `import tophandwerk_core as core`)
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import time
import unicodedata
import uuid
import zipfile
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from difflib import SequenceMatcher
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.eventlog.core import event_append

# Optional libs
try:
    from pypdf import PdfReader  # type: ignore
except Exception:
    PdfReader = None

try:
    import fitz  # PyMuPDF  # type: ignore
except Exception:
    fitz = None

try:
    from PIL import Image  # type: ignore
except Exception:
    Image = None

try:
    import pytesseract  # type: ignore
except Exception:
    pytesseract = None

try:
    from docx import Document as DocxDocument  # type: ignore
except Exception:
    DocxDocument = None

try:
    import openpyxl  # type: ignore
except Exception:
    openpyxl = None

# Optional for Outlook .msg
try:
    import extract_msg  # type: ignore
except Exception:
    extract_msg = None


# ============================================================
# ENV helpers (KUKANILEA_* overrides TOPHANDWERK_*)
# ============================================================
def _env(key: str, default: str = "") -> str:
    """
    Read env with alias support:
      KUKANILEA_<key> overrides TOPHANDWERK_<key>.
    key is expected WITHOUT prefix, e.g. "TENANT_DEFAULT".
    """
    k1 = f"KUKANILEA_{key}"
    k2 = f"TOPHANDWERK_{key}"
    v = os.environ.get(k1)
    if v is not None:
        return str(v)
    v = os.environ.get(k2)
    if v is not None:
        return str(v)
    return default


def _env_bool(key: str, default: str = "0") -> bool:
    v = _env(key, default).strip()
    return v in ("1", "true", "TRUE", "yes", "YES", "on", "ON")


# ============================================================
# CONFIG / PATHS
# ============================================================
# (bewusst: Default-Pfade bleiben "Tophandwerk_*", damit bestehende Daten nicht verloren gehen)
EINGANG = Path.home() / _env("EINGANG_DIRNAME", "Tophandwerk_Eingang")
BASE_PATH = Path.home() / _env("BASE_DIRNAME", "Tophandwerk_Kundenablage")
PENDING_DIR = Path.home() / _env("PENDING_DIRNAME", "Tophandwerk_Pending")
DONE_DIR = Path.home() / _env("DONE_DIRNAME", "Tophandwerk_Done")
DB_PATH = Path.home() / _env("DB_FILENAME", "Tophandwerk_DB.sqlite3")
SCHEMA_VERSION = 4

# Multi-tenant behavior
TENANT_DEFAULT = _env("TENANT_DEFAULT", "").strip()  # e.g. "FIRMA_X"
TENANT_REQUIRE = _env_bool("TENANT_REQUIRE", "0")

# A2: include more formats (some are store-only if extraction returns "")
SUPPORTED_EXT = {
    # Documents
    ".pdf",
    ".txt",
    ".md",
    ".rtf",
    ".docx",
    ".xlsx",
    ".csv",
    ".eml",
    ".html",
    ".htm",
    ".json",
    ".xml",
    # Images (OCR)
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    # Outlook (best-effort)
    ".msg",
    # Containers / non-text (accepted, but extraction may be empty; store-only)
    ".zip",
    ".7z",
    ".rar",
    ".dwg",
    ".dxf",
    ".ifc",
    ".p7m",
    ".p7s",
    ".psd",
    ".ai",
    ".mp4",
    ".mov",
    ".mp3",
}

# OCR / Extraction limits
OCR_MAX_PAGES = 2
MIN_TEXT_LEN_BEFORE_OCR = 200

# Duplicate-detection for object folders (name/addr/plzort variations)
DEFAULT_DUP_SIM_THRESHOLD = 0.93

# Assistant search result size
ASSISTANT_DEFAULT_LIMIT = 50

# Extraction size guards (prevent huge RAM / DB bloat)
MAX_EXTRACT_CHARS = 200_000
MAX_CSV_ROWS = 2000
MAX_CSV_COLS = 60
MAX_XLSX_ROWS = 2000
MAX_XLSX_COLS = 60
MAX_DOCX_PARAS = 4000


# ============================================================
# UTIL
# ============================================================
def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _sha256_bytes(b: bytes) -> str:
    h = hashlib.sha256()
    h.update(b)
    return h.hexdigest()


def _read_bytes(fp: Path) -> bytes:
    with open(fp, "rb") as f:
        return f.read()


def _token() -> str:
    raw = f"{time.time_ns()}:{os.getpid()}:{os.urandom(8).hex()}".encode("utf-8")
    return base64.urlsafe_b64encode(hashlib.sha256(raw).digest())[:22].decode("ascii")


def normalize_component(s: Any) -> str:
    """
    Strong normalization:
    - normalize unicode (NFKC)
    - collapse whitespace
    """
    if s is None:
        return ""
    s = str(s).strip()
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _to_jsonable(x: Any) -> Any:
    if isinstance(x, dict):
        return {str(k): _to_jsonable(v) for k, v in x.items()}
    if isinstance(x, (list, tuple, set)):
        return [_to_jsonable(v) for v in x]
    if isinstance(x, (datetime, date)):
        return x.isoformat()
    if isinstance(x, Decimal):
        return float(x)
    if isinstance(x, Path):
        return str(x)
    if isinstance(x, bytes):
        return base64.b64encode(x).decode("ascii")
    if x is None or isinstance(x, (str, int, float, bool)):
        return x
    return str(x)


def _payload(**kwargs: Any) -> Dict[str, Any]:
    return {k: _to_jsonable(v) for k, v in kwargs.items()}


def _norm_for_match(s: Any) -> str:
    """
    Aggressive matching normalization:
    - lower
    - replace umlauts
    - drop non-alnum
    """
    s = normalize_component(s).lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = re.sub(r"[^a-z0-9]+", "", s)
    return s


def _atomic_write_text(fp: Path, text: str, encoding: str = "utf-8") -> None:
    fp.parent.mkdir(parents=True, exist_ok=True)
    tmp = fp.with_suffix(fp.suffix + f".tmp_{os.getpid()}_{time.time_ns()}")
    tmp.write_text(text, encoding=encoding)
    tmp.replace(fp)


def _clip_text(s: str, limit: int = MAX_EXTRACT_CHARS) -> str:
    s = s or ""
    if len(s) <= limit:
        return s
    return s[:limit]


def _html_to_text(html: str) -> str:
    """
    Very pragmatic HTML -> text (no external deps).
    """
    if not html:
        return ""
    html = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>", "\n", html)
    html = re.sub(r"(?i)</p\s*>", "\n", html)
    html = re.sub(r"(?is)<[^>]+>", " ", html)
    html = (
        html.replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
    )
    html = re.sub(r"[ \t]+", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()


# --------------------------
# Tenant helpers (no DB migration)
# --------------------------
def _infer_tenant_from_path(fp: Path) -> str:
    """
    Convention:
      BASE_PATH/<TENANT>/<KUNDENORDNER>/...
    Returns tenant or "".
    """
    try:
        fp = Path(fp).resolve()
        base = Path(BASE_PATH).resolve()
        parts = fp.parts
        bparts = base.parts

        for i in range(len(parts) - len(bparts) + 1):
            if parts[i : i + len(bparts)] == bparts:
                if i + len(bparts) < len(parts):
                    tenant = normalize_component(parts[i + len(bparts)])
                    if tenant and not re.match(r"^\d{3,}_", tenant):
                        return tenant
                break
    except Exception:
        pass
    return ""


def _tenant_prefix_kdnr(tenant: str, kdnr: str) -> str:
    tenant = normalize_component(tenant)
    kdnr = normalize_component(kdnr)
    if not kdnr:
        return ""
    if ":" in kdnr:
        return kdnr
    if tenant:
        return f"{tenant}:{kdnr}"
    return kdnr


def _tenant_object_folder_tag(tenant: str, object_folder: str) -> str:
    tenant = normalize_component(tenant)
    object_folder = normalize_component(object_folder)
    if tenant and object_folder:
        return f"{tenant}/{object_folder}"
    return object_folder


def _effective_tenant(*candidates: Any) -> str:
    """
    Choose first non-empty tenant from candidates, else env default.
    """
    for c in candidates:
        t = normalize_component(c)
        if t:
            return t
    return normalize_component(TENANT_DEFAULT)


def _safe_fs(s: Any) -> str:
    s = normalize_component(s)
    if not s:
        return ""
    s = re.sub(r"[^\wäöüÄÖÜß\-\.\, ]+", "", s)
    s = s.replace(" ", "_")
    s = re.sub(r"_+", "_", s).strip("_")
    return s[:60]


_DATE_PATTERNS = [
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
]


def parse_excel_like_date(s: Any) -> str:
    """
    Accept common Excel-like date strings -> normalized YYYY-MM-DD or "" if invalid.
    """
    if not s:
        return ""
    s = str(s).strip()
    if not s:
        return ""
    s = re.sub(r"\s+", " ", s).replace("T", " ")

    for pat in _DATE_PATTERNS:
        try:
            dt = datetime.strptime(s, pat)
            d = dt.date()
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    m = re.search(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        try:
            y = int(m.group(1))
            mo = int(m.group(2))
            da = int(m.group(3))
            d = date(y, mo, da)
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
    if m:
        try:
            da = int(m.group(1))
            mo = int(m.group(2))
            y = int(m.group(3))
            if y < 100:
                y = 2000 + y if y <= 68 else 1900 + y
            d = date(y, mo, da)
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    return ""


# ============================================================
# PENDING / DONE store (JSON files)
# ============================================================
def _pending_path(token: str) -> Path:
    return PENDING_DIR / f"{token}.json"


def _done_path(token: str) -> Path:
    return DONE_DIR / f"{token}.json"


def read_pending(token: str) -> Optional[Dict[str, Any]]:
    fp = _pending_path(token)
    if not fp.exists():
        return None
    try:
        return json.loads(fp.read_text(encoding="utf-8"))
    except Exception:
        return None


def write_pending(token: str, payload: Dict[str, Any]) -> None:
    fp = _pending_path(token)
    _atomic_write_text(
        fp, json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def delete_pending(token: str) -> None:
    fp = _pending_path(token)
    try:
        fp.unlink()
    except Exception:
        pass


def list_pending() -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    PENDING_DIR.mkdir(parents=True, exist_ok=True)
    for fp in sorted(
        PENDING_DIR.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True
    ):
        try:
            j = json.loads(fp.read_text(encoding="utf-8"))
            j["_token"] = fp.stem
            out.append(j)
        except Exception:
            continue
    return out


def write_done(token: str, payload: Dict[str, Any]) -> None:
    fp = _done_path(token)
    _atomic_write_text(
        fp, json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def read_done(token: str) -> Optional[Dict[str, Any]]:
    fp = _done_path(token)
    if not fp.exists():
        return None
    try:
        return json.loads(fp.read_text(encoding="utf-8"))
    except Exception:
        return None


# ============================================================
# SQLITE DB
# ============================================================
_DB_LOCK = threading.Lock()
_FTS5_AVAILABLE: Optional[bool] = None


def _db() -> sqlite3.Connection:
    con = sqlite3.connect(str(DB_PATH), timeout=5.0)
    con.row_factory = sqlite3.Row
    try:
        con.execute("PRAGMA journal_mode=WAL;")
        con.execute("PRAGMA synchronous=NORMAL;")
    except Exception:
        pass
    con.execute("PRAGMA foreign_keys=ON;")
    con.execute("PRAGMA busy_timeout=5000;")
    return con


def _table_exists(con: sqlite3.Connection, name: str) -> bool:
    row = con.execute(
        "SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name=?",
        (name,),
    ).fetchone()
    return bool(row)


def _column_exists(con: sqlite3.Connection, table: str, column: str) -> bool:
    rows = con.execute(f"PRAGMA table_info({table})").fetchall()
    return any(r[1] == column for r in rows)


def _ensure_column(
    con: sqlite3.Connection, table: str, column: str, col_type: str
) -> None:
    if not _table_exists(con, table):
        return
    if not _column_exists(con, table, column):
        con.execute(f"ALTER TABLE {table} ADD COLUMN {column} {col_type}")


def _add_column_if_missing(
    con: sqlite3.Connection, table: str, col_name: str, col_def_sql: str
) -> None:
    _ensure_column(con, table, col_name, col_def_sql)


def _has_fts5(con: sqlite3.Connection) -> bool:
    global _FTS5_AVAILABLE
    if _FTS5_AVAILABLE is not None:
        return _FTS5_AVAILABLE
    try:
        con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS __fts5_test USING fts5(x);")
        con.execute("DROP TABLE IF EXISTS __fts5_test;")
        _FTS5_AVAILABLE = True
    except Exception:
        _FTS5_AVAILABLE = False
    return bool(_FTS5_AVAILABLE)


def _init_entity_link_tables(con: sqlite3.Connection) -> None:
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS entity_links (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          a_type TEXT NOT NULL,
          a_id TEXT NOT NULL,
          b_type TEXT NOT NULL,
          b_id TEXT NOT NULL,
          link_type TEXT NOT NULL,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          CHECK (NOT (a_type = b_type AND a_id = b_id)),
          UNIQUE(tenant_id, a_type, a_id, b_type, b_id, link_type)
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_entity_links_tenant_a ON entity_links(tenant_id, a_type, a_id, created_at DESC);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_entity_links_tenant_b ON entity_links(tenant_id, b_type, b_id, created_at DESC);"
    )


def _init_knowledge_tables(con: sqlite3.Connection) -> None:
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_chunks (
          id INTEGER PRIMARY KEY,
          chunk_id TEXT NOT NULL,
          tenant_id TEXT NOT NULL,
          owner_user_id TEXT,
          source_type TEXT NOT NULL,
          source_ref TEXT NOT NULL,
          title TEXT,
          body TEXT NOT NULL,
          tags TEXT,
          content_hash TEXT NOT NULL,
          is_redacted INTEGER NOT NULL DEFAULT 1,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(tenant_id, chunk_id),
          UNIQUE(tenant_id, source_type, source_ref, content_hash)
        );
        """
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_source_policies (
          tenant_id TEXT PRIMARY KEY,
          allow_manual INTEGER NOT NULL DEFAULT 1,
          allow_tasks INTEGER NOT NULL DEFAULT 1,
          allow_projects INTEGER NOT NULL DEFAULT 1,
          allow_documents INTEGER NOT NULL DEFAULT 0,
          allow_leads INTEGER NOT NULL DEFAULT 0,
          allow_email INTEGER NOT NULL DEFAULT 0,
          allow_calendar INTEGER NOT NULL DEFAULT 0,
          allow_ocr INTEGER NOT NULL DEFAULT 0,
          allow_customer_pii INTEGER NOT NULL DEFAULT 0,
          updated_at TEXT NOT NULL
        );
        """
    )
    _add_column_if_missing(
        con, "knowledge_source_policies", "allow_ocr", "INTEGER NOT NULL DEFAULT 0"
    )

    if _has_fts5(con):
        con.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_fts USING fts5(
              title,
              body,
              tags,
              content='knowledge_chunks',
              content_rowid='id'
            );
            """
        )
    else:
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS knowledge_fts_fallback (
              rowid INTEGER PRIMARY KEY,
              title TEXT,
              body TEXT,
              tags TEXT
            );
            """
        )

    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kchunks_tenant_updated ON knowledge_chunks(tenant_id, updated_at DESC);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kchunks_tenant_source ON knowledge_chunks(tenant_id, source_type, source_ref);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kchunks_tenant_owner ON knowledge_chunks(tenant_id, owner_user_id, updated_at DESC);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kpolicy_updated ON knowledge_source_policies(updated_at);"
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_email_sources (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          content_sha256 TEXT NOT NULL,
          received_at TEXT,
          subject_redacted TEXT,
          from_domain TEXT,
          to_domains_json TEXT,
          has_attachments INTEGER NOT NULL DEFAULT 0,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(tenant_id, content_sha256)
        );
        """
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_email_ingest_log (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          email_source_id TEXT NOT NULL,
          status TEXT NOT NULL,
          reason_code TEXT,
          created_at TEXT NOT NULL
        );
        """
    )

    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kes_tenant_created ON knowledge_email_sources(tenant_id, created_at DESC);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kes_tenant_sha ON knowledge_email_sources(tenant_id, content_sha256);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_keil_tenant_created ON knowledge_email_ingest_log(tenant_id, created_at DESC);"
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_ics_sources (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          content_sha256 TEXT NOT NULL,
          filename TEXT,
          event_count INTEGER NOT NULL DEFAULT 0,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(tenant_id, content_sha256)
        );
        """
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS knowledge_ics_ingest_log (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          ics_source_id TEXT NOT NULL,
          status TEXT NOT NULL,
          reason_code TEXT,
          created_at TEXT NOT NULL
        );
        """
    )

    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kis_tenant_created ON knowledge_ics_sources(tenant_id, created_at DESC);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kis_tenant_sha ON knowledge_ics_sources(tenant_id, content_sha256);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_kiil_tenant_created ON knowledge_ics_ingest_log(tenant_id, created_at DESC);"
    )


def _init_autonomy_tables(con: sqlite3.Connection) -> None:
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS source_watch_config (
          tenant_id TEXT PRIMARY KEY,
          documents_inbox_dir TEXT,
          email_inbox_dir TEXT,
          calendar_inbox_dir TEXT,
          exclude_globs TEXT,
          enabled INTEGER NOT NULL DEFAULT 1,
          max_bytes_per_file INTEGER NOT NULL DEFAULT 262144,
          max_files_per_scan INTEGER NOT NULL DEFAULT 200,
          updated_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS source_files (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          source_kind TEXT NOT NULL,
          basename TEXT,
          path_hash TEXT NOT NULL,
          fingerprint TEXT NOT NULL,
          metadata_json TEXT,
          status TEXT NOT NULL,
          last_seen_at TEXT NOT NULL,
          first_seen_at TEXT NOT NULL,
          last_ingested_at TEXT,
          last_error_code TEXT,
          UNIQUE(tenant_id, source_kind, path_hash)
        );
        """
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS source_ingest_log (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          source_kind TEXT NOT NULL,
          path_hash TEXT NOT NULL,
          action TEXT NOT NULL,
          detail_code TEXT NOT NULL,
          created_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_source_files_tenant_kind_status ON source_files(tenant_id, source_kind, status);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_source_files_tenant_seen ON source_files(tenant_id, last_seen_at);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_source_ingest_log_tenant_time ON source_ingest_log(tenant_id, created_at);"
    )
    _add_column_if_missing(con, "source_watch_config", "exclude_globs", "TEXT")
    _add_column_if_missing(con, "source_files", "metadata_json", "TEXT")
    _add_column_if_missing(con, "source_files", "basename", "TEXT")
    _add_column_if_missing(con, "source_files", "duplicate_of_file_id", "TEXT")
    _add_column_if_missing(con, "source_files", "doctype_token", "TEXT")
    _add_column_if_missing(con, "source_files", "correspondent_token", "TEXT")
    _add_column_if_missing(con, "source_files", "autotag_applied_at", "TEXT")
    _add_column_if_missing(con, "source_files", "sha256", "TEXT")
    _add_column_if_missing(con, "source_files", "size_bytes", "INTEGER")
    _add_column_if_missing(con, "source_files", "knowledge_chunk_id", "TEXT")
    _add_column_if_missing(con, "source_files", "ocr_knowledge_chunk_id", "TEXT")

    _add_column_if_missing(con, "knowledge_chunks", "doctype_token", "TEXT")
    _add_column_if_missing(con, "knowledge_chunks", "correspondent_token", "TEXT")

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS auto_tagging_rules (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          name TEXT NOT NULL,
          enabled INTEGER NOT NULL DEFAULT 1,
          priority INTEGER NOT NULL DEFAULT 0,
          condition_json TEXT NOT NULL,
          action_json TEXT NOT NULL,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(tenant_id, name)
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_autotag_rules_tenant_enabled_prio ON auto_tagging_rules(tenant_id, enabled, priority);"
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS autonomy_maintenance_status (
          tenant_id TEXT PRIMARY KEY,
          last_backup_at TEXT,
          last_backup_size_bytes INTEGER,
          last_backup_verified INTEGER,
          last_log_rotation_at TEXT,
          last_smoke_test_at TEXT,
          last_smoke_test_result TEXT,
          config_json TEXT,
          updated_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS autonomy_scan_history (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          started_at TEXT NOT NULL,
          finished_at TEXT,
          status TEXT NOT NULL,
          files_scanned INTEGER DEFAULT 0,
          files_ingested INTEGER DEFAULT 0,
          files_skipped_dedup INTEGER DEFAULT 0,
          files_skipped_exclude INTEGER DEFAULT 0,
          files_failed INTEGER DEFAULT 0,
          error_summary TEXT,
          created_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_scan_history_tenant_start ON autonomy_scan_history(tenant_id, started_at DESC);"
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS autonomy_ocr_jobs (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          source_file_id TEXT NOT NULL,
          status TEXT NOT NULL,
          started_at TEXT,
          finished_at TEXT,
          duration_ms INTEGER DEFAULT 0,
          bytes_in INTEGER DEFAULT 0,
          chars_out INTEGER DEFAULT 0,
          error_code TEXT,
          created_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_autonomy_ocr_jobs_tenant_status ON autonomy_ocr_jobs(tenant_id, status, created_at DESC);"
    )

    con.execute(
        """
        CREATE TABLE IF NOT EXISTS tags (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          name TEXT NOT NULL,
          name_norm TEXT NOT NULL,
          color TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(tenant_id, name_norm)
        );
        """
    )
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS tag_assignments (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          entity_type TEXT NOT NULL,
          entity_id TEXT NOT NULL,
          tag_id TEXT NOT NULL,
          created_at TEXT NOT NULL,
          FOREIGN KEY(tag_id) REFERENCES tags(id) ON DELETE CASCADE,
          UNIQUE(tenant_id, entity_type, entity_id, tag_id)
        );
        """
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_tags_tenant_name_norm ON tags(tenant_id, name_norm);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_tag_assignments_tenant_entity ON tag_assignments(tenant_id, entity_type, entity_id);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_tag_assignments_tenant_tag ON tag_assignments(tenant_id, tag_id);"
    )


def _init_conversation_tables(con: sqlite3.Connection) -> None:
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS conversation_events (
          id TEXT PRIMARY KEY,
          tenant_id TEXT NOT NULL,
          channel TEXT NOT NULL CHECK(channel IN ('email','chat','phone')),
          channel_ref TEXT,
          channel_ref_norm TEXT,
          direction TEXT NOT NULL CHECK(direction IN ('inbound','outbound')),
          occurred_at TEXT NOT NULL,
          redacted_payload_json TEXT NOT NULL,
          redaction_findings_json TEXT,
          audit_hash TEXT NOT NULL,
          created_at TEXT NOT NULL
        );
        """
    )
    con.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_conversation_events_ref
        ON conversation_events(tenant_id, channel, channel_ref_norm)
        WHERE channel_ref_norm IS NOT NULL;
        """
    )
    con.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_conversation_events_hash ON conversation_events(tenant_id, channel, audit_hash);"
    )
    con.execute(
        "CREATE INDEX IF NOT EXISTS idx_conversation_events_tenant_time ON conversation_events(tenant_id, occurred_at DESC);"
    )


def db_init() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS users(
                  username TEXT PRIMARY KEY,
                  pass_sha256 TEXT NOT NULL,
                  created_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS roles(
                  username TEXT NOT NULL,
                  role TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  PRIMARY KEY(username, role),
                  FOREIGN KEY(username) REFERENCES users(username) ON DELETE CASCADE
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS audit(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ts TEXT NOT NULL,
                  user TEXT NOT NULL,
                  role TEXT NOT NULL,
                  action TEXT NOT NULL,
                  target TEXT,
                  meta_json TEXT
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS tasks(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ts TEXT NOT NULL,
                  tenant TEXT NOT NULL,
                  severity TEXT NOT NULL,
                  task_type TEXT NOT NULL,
                  status TEXT NOT NULL,
                  title TEXT NOT NULL,
                  details TEXT,
                  token TEXT,
                  path TEXT,
                  meta_json TEXT,
                  created_by TEXT,
                  resolved_by TEXT,
                  resolved_at TEXT
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_tasks_status ON tasks(status, ts);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_tasks_tenant ON tasks(tenant, status, ts);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS time_projects(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT NOT NULL,
                  name TEXT NOT NULL,
                  description TEXT,
                  status TEXT NOT NULL DEFAULT 'ACTIVE',
                  budget_hours INTEGER NOT NULL DEFAULT 0,
                  budget_cost REAL NOT NULL DEFAULT 0,
                  created_by TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_time_projects_tenant ON time_projects(tenant_id, status);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS time_entries(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT NOT NULL,
                  project_id INTEGER,
                  task_id INTEGER,
                  user_id INTEGER,
                  user TEXT NOT NULL,
                  start_at TEXT NOT NULL,
                  end_at TEXT,
                  start_time TEXT,
                  end_time TEXT,
                  duration_seconds INTEGER,
                  duration INTEGER,
                  note TEXT,
                  approval_status TEXT NOT NULL DEFAULT 'PENDING',
                  approved_by TEXT,
                  approved_at TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(project_id) REFERENCES time_projects(id) ON DELETE SET NULL,
                  FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE SET NULL
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_time_entries_tenant ON time_entries(tenant_id, start_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_time_entries_user ON time_entries(tenant_id, user, start_at);"
            )
            con.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS idx_time_entries_running
                ON time_entries(tenant_id, user)
                WHERE end_at IS NULL;
                """
            )
            _ensure_column(
                con, "time_projects", "budget_hours", "INTEGER NOT NULL DEFAULT 0"
            )
            _ensure_column(
                con, "time_projects", "budget_cost", "REAL NOT NULL DEFAULT 0"
            )
            _ensure_column(con, "time_entries", "task_id", "INTEGER")
            _ensure_column(con, "time_entries", "user_id", "INTEGER")
            _ensure_column(con, "time_entries", "start_time", "TEXT")
            _ensure_column(con, "time_entries", "end_time", "TEXT")
            _ensure_column(con, "time_entries", "duration", "INTEGER")
            _ensure_column(con, "tasks", "project_id", "INTEGER")
            _ensure_column(con, "deals", "project_id", "INTEGER")

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS ai_predictions(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT,
                  project_id INTEGER NOT NULL,
                  predicted_hours REAL,
                  predicted_cost REAL,
                  deviation_ratio REAL,
                  llm_summary TEXT,
                  meta_json TEXT,
                  created_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS ai_insights(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT,
                  project_id INTEGER,
                  insight_type TEXT NOT NULL,
                  title TEXT NOT NULL,
                  message TEXT NOT NULL,
                  meta_json TEXT,
                  created_at TEXT NOT NULL
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS events(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ts TEXT NOT NULL,
                  event_type TEXT NOT NULL,
                  entity_type TEXT NOT NULL,
                  entity_id INTEGER NOT NULL,
                  payload_json TEXT NOT NULL,
                  prev_hash TEXT NOT NULL,
                  hash TEXT NOT NULL UNIQUE
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_events_entity ON events(entity_type, entity_id, id DESC);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_events_entity_created ON events(entity_type, entity_id, ts DESC);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS skills(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  cache_key TEXT NOT NULL UNIQUE,
                  name TEXT NOT NULL,
                  source_url TEXT NOT NULL,
                  ref TEXT NOT NULL,
                  resolved_commit TEXT NOT NULL,
                  status TEXT NOT NULL,
                  fetched_at TEXT NOT NULL,
                  manifest_json TEXT NOT NULL
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_skills_status ON skills(status, fetched_at DESC);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS customers(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  name TEXT NOT NULL,
                  vat_id TEXT,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_customers_tenant ON customers(tenant_id);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS contacts(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  customer_id TEXT NOT NULL,
                  name TEXT NOT NULL,
                  email TEXT,
                  phone TEXT,
                  role TEXT,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(customer_id) REFERENCES customers(id)
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_contacts_tenant_customer ON contacts(tenant_id, customer_id);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS deals(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  customer_id TEXT NOT NULL,
                  title TEXT NOT NULL,
                  stage TEXT NOT NULL,
                  project_id INTEGER,
                  value_cents INTEGER,
                  currency TEXT DEFAULT 'EUR',
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(customer_id) REFERENCES customers(id)
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_deals_tenant_stage ON deals(tenant_id, stage);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS quotes(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  customer_id TEXT NOT NULL,
                  deal_id TEXT,
                  status TEXT NOT NULL,
                  currency TEXT NOT NULL DEFAULT 'EUR',
                  subtotal_cents INTEGER NOT NULL DEFAULT 0,
                  tax_cents INTEGER NOT NULL DEFAULT 0,
                  total_cents INTEGER NOT NULL DEFAULT 0,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(customer_id) REFERENCES customers(id),
                  FOREIGN KEY(deal_id) REFERENCES deals(id)
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_quotes_tenant_customer ON quotes(tenant_id, customer_id);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS quote_items(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  quote_id TEXT NOT NULL,
                  description TEXT NOT NULL,
                  qty REAL NOT NULL DEFAULT 1,
                  unit_price_cents INTEGER NOT NULL DEFAULT 0,
                  line_total_cents INTEGER NOT NULL DEFAULT 0,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(quote_id) REFERENCES quotes(id) ON DELETE CASCADE
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_quote_items_tenant_quote ON quote_items(tenant_id, quote_id);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS emails_cache(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  customer_id TEXT,
                  contact_id TEXT,
                  message_id TEXT,
                  from_addr TEXT,
                  to_addrs TEXT,
                  subject TEXT,
                  received_at TEXT,
                  body_text TEXT,
                  raw_eml BLOB NOT NULL,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(customer_id) REFERENCES customers(id),
                  FOREIGN KEY(contact_id) REFERENCES contacts(id)
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_emails_tenant_received ON emails_cache(tenant_id, received_at DESC);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_deals_tenant_customer ON deals(tenant_id, customer_id);"
            )

            _ensure_column(con, "deals", "probability", "INTEGER")
            _ensure_column(con, "deals", "expected_close_date", "TEXT")
            _ensure_column(con, "quotes", "quote_number", "TEXT")
            _ensure_column(
                con, "quotes", "tax_amount_cents", "INTEGER NOT NULL DEFAULT 0"
            )
            _ensure_column(con, "emails_cache", "attachments_json", "TEXT")

            _ensure_column(con, "leads", "priority", "TEXT NOT NULL DEFAULT 'normal'")
            _ensure_column(con, "leads", "pinned", "INTEGER NOT NULL DEFAULT 0")
            _ensure_column(con, "leads", "assigned_to", "TEXT")
            _ensure_column(con, "leads", "response_due", "TEXT")
            _ensure_column(con, "leads", "screened_at", "TEXT")
            _ensure_column(con, "leads", "blocked_reason", "TEXT")

            con.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS idx_emails_cache_tenant_message
                ON emails_cache(tenant_id, message_id)
                WHERE message_id IS NOT NULL AND message_id != '';
                """
            )
            con.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS idx_quotes_tenant_quote_number
                ON quotes(tenant_id, quote_number)
                WHERE quote_number IS NOT NULL AND quote_number != '';
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS leads(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  status TEXT NOT NULL DEFAULT 'new'
                    CHECK(status IN ('new','contacted','qualified','lost','won','screening','ignored')),
                  source TEXT NOT NULL
                    CHECK(source IN ('call','email','webform','manual')),
                  customer_id TEXT,
                  contact_name TEXT,
                  contact_email TEXT,
                  contact_phone TEXT,
                  subject TEXT,
                  message TEXT,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(customer_id) REFERENCES customers(id)
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS call_logs(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  lead_id TEXT,
                  caller_name TEXT,
                  caller_phone TEXT,
                  direction TEXT NOT NULL CHECK(direction IN ('inbound','outbound')),
                  duration_seconds INTEGER,
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(lead_id) REFERENCES leads(id)
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS appointment_requests(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  lead_id TEXT NOT NULL,
                  requested_date TEXT,
                  status TEXT NOT NULL DEFAULT 'pending'
                    CHECK(status IN ('pending','accepted','declined','rescheduled')),
                  notes TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(lead_id) REFERENCES leads(id)
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS lead_blocklist(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  kind TEXT NOT NULL,
                  value TEXT NOT NULL,
                  reason TEXT,
                  created_at TEXT NOT NULL,
                  created_by TEXT,
                  UNIQUE(tenant_id, kind, value)
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS lead_claims(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  lead_id TEXT NOT NULL,
                  claimed_by TEXT NOT NULL,
                  claimed_at TEXT NOT NULL,
                  claimed_until TEXT NOT NULL,
                  released_at TEXT,
                  release_reason TEXT,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  UNIQUE(tenant_id, lead_id)
                );
                """
            )

            _add_column_if_missing(
                con, "leads", "priority", "TEXT NOT NULL DEFAULT 'normal'"
            )
            _add_column_if_missing(con, "leads", "pinned", "INTEGER NOT NULL DEFAULT 0")
            _add_column_if_missing(con, "leads", "assigned_to", "TEXT")
            _add_column_if_missing(con, "leads", "response_due", "TEXT")
            _add_column_if_missing(con, "leads", "screened_at", "TEXT")
            _add_column_if_missing(con, "leads", "blocked_reason", "TEXT")

            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_status_created ON leads(tenant_id, status, created_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_created ON leads(tenant_id, created_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_call_logs_tenant_lead_created ON call_logs(tenant_id, lead_id, created_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_appt_tenant_lead_updated ON appointment_requests(tenant_id, lead_id, updated_at);"
            )

            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_status_due ON leads(tenant_id, status, response_due);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_priority_created ON leads(tenant_id, priority, created_at DESC);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_pinned_created ON leads(tenant_id, pinned, created_at DESC);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_leads_tenant_assigned_due ON leads(tenant_id, assigned_to, response_due);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_blocklist_tenant_kind_value ON lead_blocklist(tenant_id, kind, value);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_lead_claims_tenant_lead ON lead_claims(tenant_id, lead_id);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_lead_claims_tenant_user ON lead_claims(tenant_id, claimed_by, claimed_until);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_lead_claims_tenant_until ON lead_claims(tenant_id, claimed_until);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_events_entity_created ON events(entity_type, entity_id, ts DESC);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS automation_rules(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  enabled INTEGER NOT NULL DEFAULT 1 CHECK (enabled IN (0,1)),
                  name TEXT NOT NULL,
                  scope TEXT NOT NULL,
                  condition_kind TEXT NOT NULL,
                  condition_json TEXT NOT NULL,
                  action_list_json TEXT NOT NULL,
                  created_by TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  updated_at TEXT NOT NULL,
                  last_error TEXT,
                  last_error_at TEXT
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS automation_runs(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  triggered_by TEXT NOT NULL,
                  started_at TEXT NOT NULL,
                  finished_at TEXT,
                  status TEXT NOT NULL DEFAULT 'running',
                  max_actions INTEGER NOT NULL,
                  actions_executed INTEGER NOT NULL DEFAULT 0,
                  aborted_reason TEXT,
                  warnings_json TEXT
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS automation_run_actions(
                  id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  run_id TEXT NOT NULL,
                  rule_id TEXT NOT NULL,
                  target_entity_type TEXT NOT NULL,
                  target_entity_id_int INTEGER NOT NULL,
                  action_kind TEXT NOT NULL,
                  action_hash TEXT NOT NULL,
                  status TEXT NOT NULL,
                  error TEXT,
                  created_at TEXT NOT NULL,
                  UNIQUE(tenant_id, run_id, rule_id, target_entity_type, target_entity_id_int, action_kind, action_hash)
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS daily_insights_cache(
                  tenant_id TEXT NOT NULL,
                  day TEXT NOT NULL,
                  payload_json TEXT NOT NULL,
                  generated_at TEXT NOT NULL,
                  PRIMARY KEY (tenant_id, day)
                );
                """
            )

            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_automation_rules_tenant_enabled ON automation_rules(tenant_id, enabled, updated_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_automation_runs_tenant_started ON automation_runs(tenant_id, started_at DESC);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_automation_run_actions_tenant_run ON automation_run_actions(tenant_id, run_id, created_at);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_daily_insights_cache_day ON daily_insights_cache(day);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS ontology_types(
                  type_name TEXT PRIMARY KEY,
                  table_name TEXT NOT NULL,
                  pk_field TEXT NOT NULL DEFAULT 'id',
                  title_field TEXT,
                  description_field TEXT,
                  created_at TEXT
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS derived_active_timers(
                  user_id INTEGER PRIMARY KEY,
                  task_id INTEGER NOT NULL,
                  start_time TEXT NOT NULL,
                  last_event_id INTEGER
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS derived_budget_progress(
                  project_id INTEGER PRIMARY KEY,
                  total_hours REAL,
                  total_cost REAL,
                  budget_hours REAL,
                  budget_cost REAL,
                  hours_percent REAL,
                  cost_percent REAL,
                  last_event_id INTEGER
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS benchmarks(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ts TEXT NOT NULL,
                  metric TEXT NOT NULL,
                  scope_type TEXT NOT NULL,
                  scope_id INTEGER,
                  n INTEGER NOT NULL,
                  p50 REAL NOT NULL,
                  p75 REAL NOT NULL,
                  p90 REAL NOT NULL,
                  min REAL,
                  max REAL,
                  last_event_id INTEGER
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS review_locks(
                  token TEXT PRIMARY KEY,
                  tenant TEXT NOT NULL,
                  locked_by TEXT NOT NULL,
                  locked_roles TEXT,
                  locked_at TEXT NOT NULL,
                  heartbeat_at TEXT NOT NULL,
                  expires_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_locks_tenant ON review_locks(tenant);"
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS docs(
                  doc_id TEXT PRIMARY KEY,                  -- sha256(file bytes)
                  group_key TEXT NOT NULL,                  -- heuristic group for versioning
                  tenant_id TEXT,
                  kdnr TEXT,
                  object_folder TEXT,
                  doctype TEXT,
                  doc_date TEXT,                            -- YYYY-MM-DD (OPTIONAL)
                  created_at TEXT NOT NULL
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS versions(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  doc_id TEXT NOT NULL,
                  version_no INTEGER NOT NULL,
                  bytes_sha256 TEXT NOT NULL,
                  file_name TEXT NOT NULL,
                  file_path TEXT NOT NULL,
                  extracted_text TEXT,
                  used_ocr INTEGER NOT NULL DEFAULT 0,
                  note TEXT,
                  created_at TEXT NOT NULL,
                  tenant_id TEXT,
                  FOREIGN KEY(doc_id) REFERENCES docs(doc_id) ON DELETE CASCADE
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS docs_index(
                  doc_id TEXT PRIMARY KEY,
                  tenant_id TEXT NOT NULL,
                  kdnr TEXT,
                  doctype TEXT,
                  customer_name TEXT,
                  address TEXT,
                  doc_date TEXT,
                  doc_number TEXT,
                  file_name TEXT,
                  file_path TEXT,
                  tokens TEXT,
                  snippet TEXT,
                  updated_at TEXT NOT NULL,
                  FOREIGN KEY(doc_id) REFERENCES docs(doc_id) ON DELETE CASCADE
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS tenants(
                  tenant_id TEXT PRIMARY KEY,
                  display_name TEXT NOT NULL,
                  created_at TEXT NOT NULL
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS tenant_users(
                  tenant_id TEXT NOT NULL,
                  username TEXT NOT NULL,
                  role TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  PRIMARY KEY(tenant_id, username),
                  FOREIGN KEY(username) REFERENCES users(username) ON DELETE CASCADE
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS entities(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT NOT NULL,
                  doc_id TEXT NOT NULL,
                  entity_type TEXT NOT NULL,
                  value TEXT NOT NULL,
                  norm_value TEXT,
                  meta_json TEXT,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(doc_id) REFERENCES docs(doc_id) ON DELETE CASCADE
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS links(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  tenant_id TEXT NOT NULL,
                  doc_id TEXT NOT NULL,
                  entity_id INTEGER,
                  link_type TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(doc_id) REFERENCES docs(doc_id) ON DELETE CASCADE,
                  FOREIGN KEY(entity_id) REFERENCES entities(id) ON DELETE CASCADE
                );
                """
            )

            _ensure_column(con, "docs", "tenant_id", "TEXT")
            _ensure_column(con, "versions", "tenant_id", "TEXT")
            _ensure_column(con, "audit", "tenant_id", "TEXT")
            _ensure_column(con, "tasks", "tenant_id", "TEXT")

            con.execute("CREATE INDEX IF NOT EXISTS idx_docs_group ON docs(group_key);")
            con.execute("CREATE INDEX IF NOT EXISTS idx_docs_kdnr ON docs(kdnr);")
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_docs_tenant ON docs(tenant_id, kdnr);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_versions_doc ON versions(doc_id);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_versions_path ON versions(file_path);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_versions_tenant ON versions(tenant_id);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_entities_doc ON entities(doc_id, tenant_id);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_entities_type ON entities(entity_type, norm_value);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_docs_index_tenant ON docs_index(tenant_id, kdnr, doctype);"
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_docs_index_tokens ON docs_index(tokens);"
            )

            if _has_fts5(con):
                con.execute(
                    """
                    CREATE VIRTUAL TABLE IF NOT EXISTS docs_fts
                    USING fts5(
                        doc_id UNINDEXED,
                        tenant_id UNINDEXED,
                        kdnr UNINDEXED,
                        doctype UNINDEXED,
                        doc_date UNINDEXED,
                        file_name UNINDEXED,
                        file_path UNINDEXED,
                        content,
                        tokenize='unicode61'
                    );
                    """
                )

            _init_entity_link_tables(con)
            _init_knowledge_tables(con)
            _init_autonomy_tables(con)
            _init_conversation_tables(con)

            row = con.execute("PRAGMA user_version").fetchone()
            cur_ver = int(row[0] if row else 0)
            if cur_ver < SCHEMA_VERSION:
                con.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")

            con.commit()
        finally:
            con.close()


# ============================================================
# RBAC
# ============================================================
def _pw_hash(pw: str) -> str:
    return _sha256_bytes((pw or "").encode("utf-8"))


def rbac_create_user(username: str, password: str) -> str:
    username = normalize_component(username).lower()
    if not username:
        raise ValueError("username required")
    if not password:
        raise ValueError("password required")

    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                "INSERT OR REPLACE INTO users(username, pass_sha256, created_at) VALUES (?,?,?)",
                (username, _pw_hash(password), _now_iso()),
            )
            con.commit()
            return username
        finally:
            con.close()


def rbac_verify_user(username: str, password: str) -> bool:
    username = normalize_component(username).lower()
    if not username or not password:
        return False
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT pass_sha256 FROM users WHERE username=?", (username,)
            ).fetchone()
            if not row:
                return False
            return str(row["pass_sha256"]) == _pw_hash(password)
        finally:
            con.close()


def rbac_assign_role(username: str, role: str) -> None:
    username = normalize_component(username).lower()
    role = normalize_component(role).upper()
    if not username or not role:
        return
    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                "INSERT OR IGNORE INTO roles(username, role, created_at) VALUES (?,?,?)",
                (username, role, _now_iso()),
            )
            con.commit()
        finally:
            con.close()


def rbac_get_user_roles(username: str) -> List[str]:
    username = normalize_component(username).lower()
    if not username:
        return []
    with _DB_LOCK:
        con = _db()
        try:
            rows = con.execute(
                "SELECT role FROM roles WHERE username=? ORDER BY role", (username,)
            ).fetchall()
            return [str(r["role"]) for r in rows]
        finally:
            con.close()


# ============================================================
# AUDIT
# ============================================================
def audit_log(
    user: str,
    role: str,
    action: str,
    target: str = "",
    meta: Optional[dict] = None,
    tenant_id: str = "",
) -> None:
    user = normalize_component(user).lower()
    role = normalize_component(role).upper()
    action = normalize_component(action)
    target = normalize_component(target)
    meta_json = json.dumps(meta or {}, ensure_ascii=False)
    tenant_id = (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )

    with _DB_LOCK:
        con = _db()
        try:
            if _column_exists(con, "audit", "tenant_id"):
                con.execute(
                    "INSERT INTO audit(ts, user, role, action, target, meta_json, tenant_id) VALUES (?,?,?,?,?,?,?)",
                    (_now_iso(), user, role, action, target, meta_json, tenant_id),
                )
            else:
                con.execute(
                    "INSERT INTO audit(ts, user, role, action, target, meta_json) VALUES (?,?,?,?,?,?)",
                    (_now_iso(), user, role, action, target, meta_json),
                )
            con.commit()
        finally:
            con.close()


# ============================================================
# TASKS
# ============================================================
def task_create(
    *,
    tenant: str,
    severity: str,
    task_type: str,
    title: str,
    details: str = "",
    token: str = "",
    path: str = "",
    meta: Optional[dict] = None,
    created_by: str = "",
) -> int:
    tenant = normalize_component(tenant).lower()
    severity = normalize_component(severity).upper() or "INFO"
    task_type = normalize_component(task_type).upper() or "GENERAL"
    title = normalize_component(title)
    details = (details or "").strip()
    token = normalize_component(token)
    path = normalize_component(path)
    created_by = normalize_component(created_by).lower()
    meta_json = json.dumps(meta or {}, ensure_ascii=False)

    with _DB_LOCK:
        con = _db()
        try:
            cur = con.execute(
                """
                INSERT INTO tasks(ts, tenant, severity, task_type, status, title, details, token, path, meta_json, created_by)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    _now_iso(),
                    tenant,
                    severity,
                    task_type,
                    "OPEN",
                    title,
                    details,
                    token,
                    path,
                    meta_json,
                    created_by,
                ),
            )
            con.commit()
            return int(cur.lastrowid or 0)
        finally:
            con.close()


def task_list(
    *, tenant: str = "", status: str = "OPEN", limit: int = 200
) -> List[Dict[str, Any]]:
    tenant = normalize_component(tenant).lower()
    status = normalize_component(status).upper()
    limit = max(1, min(int(limit), 2000))

    with _DB_LOCK:
        con = _db()
        try:
            if tenant:
                rows = con.execute(
                    """
                    SELECT * FROM tasks
                    WHERE tenant=? AND status=?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (tenant, status, limit),
                ).fetchall()
            else:
                rows = con.execute(
                    """
                    SELECT * FROM tasks
                    WHERE status=?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (status, limit),
                ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()


def task_set_status(task_id: int, status: str, resolved_by: str = "") -> bool:
    status = normalize_component(status).upper()
    resolved_by = normalize_component(resolved_by).lower()
    if status not in {"OPEN", "RESOLVED", "DISMISSED"}:
        return False

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT id FROM tasks WHERE id=?", (int(task_id),)
            ).fetchone()
            if not row:
                return False
            con.execute(
                "UPDATE tasks SET status=?, resolved_by=?, resolved_at=? WHERE id=?",
                (status, resolved_by, _now_iso(), int(task_id)),
            )
            con.commit()
            return True
        finally:
            con.close()


# ============================================================
# TIME TRACKING
# ============================================================
def _time_tenant(tenant_id: str) -> str:
    return (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )


def _parse_iso(value: str) -> datetime:
    return datetime.fromisoformat(value)


def _duration_seconds(start_at: str, end_at: str) -> int:
    return max(0, int((_parse_iso(end_at) - _parse_iso(start_at)).total_seconds()))


def time_project_create(
    *,
    tenant_id: str,
    name: str,
    description: str = "",
    budget_hours: int = 0,
    budget_cost: float = 0.0,
    created_by: str = "",
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    name = normalize_component(name)
    description = (description or "").strip()
    created_by = normalize_component(created_by).lower()
    budget_hours = max(0, int(budget_hours or 0))
    budget_cost = max(0.0, float(budget_cost or 0.0))
    if not name:
        raise ValueError("project_name_required")

    now = _now_iso()
    project_id = 0
    with _DB_LOCK:
        con = _db()
        try:
            cur = con.execute(
                """
                INSERT INTO time_projects(
                    tenant_id, name, description, status, budget_hours, budget_cost,
                    created_by, created_at, updated_at
                )
                VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (
                    tenant_id,
                    name,
                    description,
                    "ACTIVE",
                    budget_hours,
                    budget_cost,
                    created_by,
                    now,
                    now,
                ),
            )
            con.commit()
            project_id = int(cur.lastrowid or 0)
        finally:
            con.close()
    audit_log(
        user=created_by or "system",
        role="SYSTEM",
        action="TIME_PROJECT_CREATE",
        target=str(project_id),
        meta={"name": name, "budget_hours": budget_hours, "budget_cost": budget_cost},
        tenant_id=tenant_id,
    )
    return {
        "id": project_id,
        "tenant_id": tenant_id,
        "name": name,
        "description": description,
        "status": "ACTIVE",
        "budget_hours": budget_hours,
        "budget_cost": budget_cost,
        "created_by": created_by,
        "created_at": now,
        "updated_at": now,
    }


def time_project_list(
    *, tenant_id: str, status: str = "ACTIVE"
) -> List[Dict[str, Any]]:
    tenant_id = _time_tenant(tenant_id)
    status = normalize_component(status).upper() or "ACTIVE"
    with _DB_LOCK:
        con = _db()
        try:
            rows = con.execute(
                """
                SELECT * FROM time_projects
                WHERE tenant_id=? AND status=?
                ORDER BY name
                """,
                (tenant_id, status),
            ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()


def time_project_update_budget(
    *,
    tenant_id: str,
    project_id: int,
    budget_hours: Optional[int] = None,
    budget_cost: Optional[float] = None,
    user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    project_id = int(project_id)
    user_id_norm = int(user_id) if user_id is not None else None

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT id, budget_hours, budget_cost FROM time_projects WHERE id=? AND tenant_id=?",
                (project_id, tenant_id),
            ).fetchone()
            if not row:
                raise ValueError("project_not_found")
            before = {
                "budget_hours": int(row["budget_hours"] or 0),
                "budget_cost": float(row["budget_cost"] or 0.0),
            }
            new_hours = (
                max(0, int(budget_hours))
                if budget_hours is not None
                else before["budget_hours"]
            )
            new_cost = (
                max(0.0, float(budget_cost))
                if budget_cost is not None
                else before["budget_cost"]
            )
            con.execute(
                """
                UPDATE time_projects
                SET budget_hours=?, budget_cost=?, updated_at=?
                WHERE id=? AND tenant_id=?
                """,
                (new_hours, new_cost, _now_iso(), project_id, tenant_id),
            )
            event_append(
                event_type="project_budget_updated",
                entity_type="project",
                entity_id=project_id,
                payload={
                    "schema_version": 1,
                    "source": "core/project_budget_update",
                    "actor_user_id": user_id_norm,
                    "data": _payload(
                        before=before,
                        after={"budget_hours": new_hours, "budget_cost": new_cost},
                        user_id=user_id_norm,
                    ),
                },
                con=con,
            )
            con.commit()
        finally:
            con.close()

    summary = time_entries_summary_by_project(
        tenant_id=tenant_id, project_id=project_id
    )
    summary["updated_budget_hours"] = new_hours
    summary["updated_budget_cost"] = new_cost
    return summary


def _time_project_lookup(
    con: sqlite3.Connection, tenant_id: str, project_id: Optional[int]
) -> Optional[dict]:
    if project_id is None:
        return None
    row = con.execute(
        "SELECT * FROM time_projects WHERE id=? AND tenant_id=?",
        (int(project_id), tenant_id),
    ).fetchone()
    return dict(row) if row else None


def _task_lookup(
    con: sqlite3.Connection, tenant_id: str, task_id: Optional[int]
) -> Optional[dict]:
    if task_id is None:
        return None
    row = con.execute(
        "SELECT * FROM tasks WHERE id=? AND LOWER(tenant)=LOWER(?)",
        (int(task_id), tenant_id),
    ).fetchone()
    return dict(row) if row else None


def time_entry_start(
    *,
    tenant_id: str,
    user: str,
    project_id: Optional[int] = None,
    task_id: Optional[int] = None,
    user_id: Optional[int] = None,
    note: str = "",
    started_at: Optional[str] = None,
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    user = normalize_component(user).lower()
    note = (note or "").strip()
    if not user:
        raise ValueError("user_required")

    now = started_at or _now_iso()
    entry_id = 0
    user_id_norm = int(user_id) if user_id is not None else None
    with _DB_LOCK:
        con = _db()
        try:
            if project_id is not None and not _time_project_lookup(
                con, tenant_id, project_id
            ):
                raise ValueError("project_not_found")
            if task_id is not None and not _task_lookup(con, tenant_id, task_id):
                raise ValueError("task_not_found")
            row = con.execute(
                "SELECT id FROM time_entries WHERE tenant_id=? AND user=? AND end_at IS NULL",
                (tenant_id, user),
            ).fetchone()
            if row:
                raise ValueError("running_timer_exists")
            cur = con.execute(
                """
                INSERT INTO time_entries(
                    tenant_id, project_id, task_id, user_id, user,
                    start_at, end_at, start_time, end_time,
                    duration_seconds, duration, note,
                    approval_status, created_at, updated_at
                )
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    tenant_id,
                    project_id,
                    task_id,
                    user_id_norm,
                    user,
                    now,
                    None,
                    now,
                    None,
                    None,
                    None,
                    note,
                    "PENDING",
                    now,
                    now,
                ),
            )
            entry_id = int(cur.lastrowid or 0)
            event_append(
                event_type="timer_started",
                entity_type="time_entry",
                entity_id=entry_id,
                payload={
                    "schema_version": 1,
                    "source": "core/timer_start",
                    "actor_user_id": user_id_norm,
                    "data": _payload(
                        user_id=user_id_norm,
                        task_id=task_id,
                        start_time=now,
                        end_time=None,
                        duration=None,
                    ),
                },
                con=con,
            )
            con.commit()
        finally:
            con.close()
    audit_log(
        user=user,
        role="OPERATOR",
        action="TIME_ENTRY_START",
        target=str(entry_id),
        meta={"project_id": project_id or "", "task_id": task_id or ""},
        tenant_id=tenant_id,
    )
    return time_entry_get(tenant_id=tenant_id, entry_id=entry_id) or {}


def time_entry_get(*, tenant_id: str, entry_id: int) -> Optional[Dict[str, Any]]:
    tenant_id = _time_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                """
                SELECT te.*, tp.name AS project_name, t.title AS task_title
                FROM time_entries te
                LEFT JOIN time_projects tp ON tp.id = te.project_id
                LEFT JOIN tasks t ON t.id = te.task_id
                WHERE te.tenant_id=? AND te.id=?
                """,
                (tenant_id, int(entry_id)),
            ).fetchone()
            if not row:
                return None
            entry = dict(row)
            if entry.get("end_at"):
                entry["duration_seconds"] = _duration_seconds(
                    entry["start_at"], entry["end_at"]
                )
            else:
                entry["duration_seconds"] = _duration_seconds(
                    entry["start_at"], _now_iso()
                )
            entry["start_time"] = entry.get("start_time") or entry.get("start_at")
            entry["end_time"] = entry.get("end_time") or entry.get("end_at")
            entry["duration"] = entry.get("duration") or entry.get("duration_seconds")
            return entry
        finally:
            con.close()


def time_entry_stop(
    *,
    tenant_id: str,
    user: str,
    entry_id: Optional[int] = None,
    ended_at: Optional[str] = None,
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    user = normalize_component(user).lower()
    if not user:
        raise ValueError("user_required")
    end_ts = ended_at or _now_iso()

    with _DB_LOCK:
        con = _db()
        try:
            if entry_id is None:
                row = con.execute(
                    """
                    SELECT id, user_id, task_id, start_at FROM time_entries
                    WHERE tenant_id=? AND user=? AND end_at IS NULL
                    """,
                    (tenant_id, user),
                ).fetchone()
            else:
                row = con.execute(
                    """
                    SELECT id, user_id, task_id, start_at FROM time_entries
                    WHERE tenant_id=? AND user=? AND id=?
                    """,
                    (tenant_id, user, int(entry_id)),
                ).fetchone()
            if not row:
                raise ValueError("no_running_timer")
            start_at = str(row["start_at"])
            if _parse_iso(end_ts) < _parse_iso(start_at):
                raise ValueError("invalid_time_range")
            duration = _duration_seconds(start_at, end_ts)
            con.execute(
                """
                UPDATE time_entries
                SET end_at=?, end_time=?, duration_seconds=?, duration=?, updated_at=?
                WHERE id=? AND tenant_id=?
                """,
                (
                    end_ts,
                    end_ts,
                    duration,
                    duration,
                    _now_iso(),
                    int(row["id"]),
                    tenant_id,
                ),
            )
            stopped_id = int(row["id"])
            post = con.execute(
                """
                SELECT id, user_id, task_id, start_at, end_at, duration_seconds, duration
                FROM time_entries
                WHERE id=? AND tenant_id=?
                """,
                (stopped_id, tenant_id),
            ).fetchone()
            actor_user_id = (
                int(post["user_id"]) if post and post["user_id"] is not None else None
            )
            event_append(
                event_type="timer_stopped",
                entity_type="time_entry",
                entity_id=stopped_id,
                payload={
                    "schema_version": 1,
                    "source": "core/timer_stop",
                    "actor_user_id": actor_user_id,
                    "data": _payload(
                        user_id=actor_user_id,
                        task_id=post["task_id"] if post else row["task_id"],
                        start_time=post["start_at"] if post else start_at,
                        end_time=post["end_at"] if post else end_ts,
                        duration=post["duration"] if post else duration,
                    ),
                },
                con=con,
            )
            con.commit()
        finally:
            con.close()
    audit_log(
        user=user,
        role="OPERATOR",
        action="TIME_ENTRY_STOP",
        target=str(stopped_id),
        meta={"duration_seconds": duration},
        tenant_id=tenant_id,
    )
    return time_entry_get(tenant_id=tenant_id, entry_id=stopped_id) or {}


def time_entry_update(
    *,
    tenant_id: str,
    entry_id: int,
    project_id: Optional[int] = None,
    task_id: Optional[int] = None,
    user_id: Optional[int] = None,
    start_at: Optional[str] = None,
    end_at: Optional[str] = None,
    note: Optional[str] = None,
    user: str = "",
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    user = normalize_component(user).lower()

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT * FROM time_entries WHERE id=? AND tenant_id=?",
                (int(entry_id), tenant_id),
            ).fetchone()
            if not row:
                raise ValueError("entry_not_found")
            before = dict(row)
            if project_id is not None and not _time_project_lookup(
                con, tenant_id, project_id
            ):
                raise ValueError("project_not_found")
            if task_id is not None and not _task_lookup(con, tenant_id, task_id):
                raise ValueError("task_not_found")
            start_val = start_at or row["start_at"]
            end_val = end_at if end_at is not None else row["end_at"]
            duration_val = None
            if end_val:
                if _parse_iso(end_val) < _parse_iso(start_val):
                    raise ValueError("invalid_time_range")
                duration_val = _duration_seconds(start_val, end_val)
            con.execute(
                """
                UPDATE time_entries
                SET project_id=?, task_id=?, user_id=?,
                    start_at=?, start_time=?,
                    end_at=?, end_time=?,
                    duration_seconds=?, duration=?, note=?, updated_at=?
                WHERE id=? AND tenant_id=?
                """,
                (
                    project_id if project_id is not None else row["project_id"],
                    task_id if task_id is not None else row["task_id"],
                    user_id if user_id is not None else row["user_id"],
                    start_val,
                    start_val,
                    end_val,
                    end_val,
                    duration_val,
                    duration_val,
                    note if note is not None else row["note"],
                    _now_iso(),
                    int(entry_id),
                    tenant_id,
                ),
            )
            after_row = con.execute(
                "SELECT * FROM time_entries WHERE id=? AND tenant_id=?",
                (int(entry_id), tenant_id),
            ).fetchone()
            after = dict(after_row) if after_row else {}
            actor_user_id = (
                int(after.get("user_id")) if after.get("user_id") is not None else None
            )
            event_append(
                event_type="time_entry_edited",
                entity_type="time_entry",
                entity_id=int(entry_id),
                payload={
                    "schema_version": 1,
                    "source": "core/time_entry_update",
                    "actor_user_id": actor_user_id,
                    "data": _payload(before=before, after=after),
                },
                con=con,
            )
            con.commit()
        finally:
            con.close()
    audit_log(
        user=user or "system",
        role="OPERATOR",
        action="TIME_ENTRY_EDIT",
        target=str(entry_id),
        meta={
            "project_id": project_id or "",
            "task_id": task_id or "",
            "note_changed": note is not None,
        },
        tenant_id=tenant_id,
    )
    return time_entry_get(tenant_id=tenant_id, entry_id=int(entry_id)) or {}


def time_entry_approve(
    *, tenant_id: str, entry_id: int, approved_by: str
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    approved_by = normalize_component(approved_by).lower()
    if not approved_by:
        raise ValueError("approved_by_required")
    now = _now_iso()
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT id FROM time_entries WHERE id=? AND tenant_id=?",
                (int(entry_id), tenant_id),
            ).fetchone()
            if not row:
                raise ValueError("entry_not_found")
            con.execute(
                """
                UPDATE time_entries
                SET approval_status=?, approved_by=?, approved_at=?, updated_at=?
                WHERE id=? AND tenant_id=?
                """,
                ("APPROVED", approved_by, now, now, int(entry_id), tenant_id),
            )
            con.commit()
        finally:
            con.close()
    audit_log(
        user=approved_by,
        role="ADMIN",
        action="TIME_ENTRY_APPROVE",
        target=str(entry_id),
        meta={},
        tenant_id=tenant_id,
    )
    return time_entry_get(tenant_id=tenant_id, entry_id=int(entry_id)) or {}


def time_entries_list(
    *,
    tenant_id: str,
    user: Optional[str] = None,
    start_at: Optional[str] = None,
    end_at: Optional[str] = None,
    limit: int = 500,
) -> List[Dict[str, Any]]:
    tenant_id = _time_tenant(tenant_id)
    user = normalize_component(user or "").lower()
    limit = max(1, min(int(limit), 2000))

    clauses = ["te.tenant_id=?"]
    params: List[Any] = [tenant_id]
    if user:
        clauses.append("te.user=?")
        params.append(user)
    if start_at:
        clauses.append("te.start_at>=?")
        params.append(start_at)
    if end_at:
        clauses.append("te.start_at<=?")
        params.append(end_at)

    where_sql = " AND ".join(clauses)
    with _DB_LOCK:
        con = _db()
        try:
            rows = con.execute(
                f"""
                SELECT te.*, tp.name AS project_name, t.title AS task_title
                FROM time_entries te
                LEFT JOIN time_projects tp ON tp.id = te.project_id
                LEFT JOIN tasks t ON t.id = te.task_id
                WHERE {where_sql}
                ORDER BY te.start_at DESC
                LIMIT ?
                """,
                (*params, limit),
            ).fetchall()
            entries = [dict(r) for r in rows]
            now = _now_iso()
            for entry in entries:
                if entry.get("end_at"):
                    entry["duration_seconds"] = _duration_seconds(
                        entry["start_at"], entry["end_at"]
                    )
                else:
                    entry["duration_seconds"] = _duration_seconds(
                        entry["start_at"], now
                    )
                entry["start_time"] = entry.get("start_time") or entry.get("start_at")
                entry["end_time"] = entry.get("end_time") or entry.get("end_at")
                entry["duration"] = entry.get("duration") or entry.get(
                    "duration_seconds"
                )
            return entries
        finally:
            con.close()


def time_entries_summary_by_task(*, tenant_id: str, task_id: int) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                """
                SELECT COALESCE(SUM(COALESCE(duration_seconds, duration, 0)), 0) AS total_seconds,
                       COUNT(*) AS total_entries
                FROM time_entries
                WHERE tenant_id=? AND task_id=?
                """,
                (tenant_id, int(task_id)),
            ).fetchone()
            total_seconds = int((row["total_seconds"] if row else 0) or 0)
            total_entries = int((row["total_entries"] if row else 0) or 0)
            return {
                "tenant_id": tenant_id,
                "task_id": int(task_id),
                "total_entries": total_entries,
                "total_seconds": total_seconds,
                "total_hours": round(total_seconds / 3600.0, 2),
            }
        finally:
            con.close()


def time_entries_summary_by_project(
    *, tenant_id: str, project_id: int
) -> Dict[str, Any]:
    tenant_id = _time_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                """
                SELECT p.id, p.name, p.budget_hours, p.budget_cost,
                       COALESCE(SUM(COALESCE(te.duration_seconds, te.duration, 0)), 0) AS total_seconds
                FROM time_projects p
                LEFT JOIN time_entries te ON te.project_id = p.id AND te.tenant_id = p.tenant_id
                WHERE p.tenant_id=? AND p.id=?
                GROUP BY p.id, p.name, p.budget_hours, p.budget_cost
                """,
                (tenant_id, int(project_id)),
            ).fetchone()
            if not row:
                raise ValueError("project_not_found")
            total_seconds = int(row["total_seconds"] or 0)
            total_hours = round(total_seconds / 3600.0, 2)
            budget_hours = int(row["budget_hours"] or 0)
            budget_cost = float(row["budget_cost"] or 0.0)
            pct_hours = (
                min(100.0, (total_hours / budget_hours) * 100.0)
                if budget_hours > 0
                else 0.0
            )
            pct_cost = pct_hours if budget_cost > 0 else 0.0
            return {
                "tenant_id": tenant_id,
                "project_id": int(project_id),
                "project_name": str(row["name"] or ""),
                "budget_hours": budget_hours,
                "budget_cost": budget_cost,
                "spent_hours": total_hours,
                "spent_cost": round((total_hours / budget_hours) * budget_cost, 2)
                if budget_hours > 0
                else 0.0,
                "progress_hours_pct": round(pct_hours, 2),
                "progress_cost_pct": round(pct_cost, 2),
                "warning": bool(pct_hours >= 80.0 or pct_cost >= 80.0),
            }
        finally:
            con.close()


def ai_prediction_add(
    *,
    tenant_id: str,
    project_id: int,
    predicted_hours: float,
    predicted_cost: float,
    deviation_ratio: float,
    llm_summary: str,
    meta_json: str,
) -> int:
    tenant_id = _time_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            cur = con.execute(
                """
                INSERT INTO ai_predictions(
                  tenant_id, project_id, predicted_hours, predicted_cost,
                  deviation_ratio, llm_summary, meta_json, created_at
                ) VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    tenant_id,
                    int(project_id),
                    float(predicted_hours or 0.0),
                    float(predicted_cost or 0.0),
                    float(deviation_ratio or 0.0),
                    llm_summary or "",
                    meta_json or "{}",
                    _now_iso(),
                ),
            )
            con.commit()
            return int(cur.lastrowid or 0)
        finally:
            con.close()


def ai_insight_add(
    *,
    tenant_id: str,
    project_id: Optional[int],
    insight_type: str,
    title: str,
    message: str,
    meta_json: str,
) -> int:
    tenant_id = _time_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            cur = con.execute(
                """
                INSERT INTO ai_insights(tenant_id, project_id, insight_type, title, message, meta_json, created_at)
                VALUES (?,?,?,?,?,?,?)
                """,
                (
                    tenant_id,
                    int(project_id) if project_id is not None else None,
                    normalize_component(insight_type) or "info",
                    title or "Insight",
                    message or "",
                    meta_json or "{}",
                    _now_iso(),
                ),
            )
            con.commit()
            return int(cur.lastrowid or 0)
        finally:
            con.close()


def time_entries_export_csv(
    *,
    tenant_id: str,
    user: Optional[str] = None,
    start_at: Optional[str] = None,
    end_at: Optional[str] = None,
    limit: int = 2000,
) -> str:
    entries = time_entries_list(
        tenant_id=tenant_id,
        user=user,
        start_at=start_at,
        end_at=end_at,
        limit=limit,
    )
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "entry_id",
            "project_id",
            "project_name",
            "user",
            "start_at",
            "end_at",
            "duration_seconds",
            "duration_hours",
            "note",
            "approval_status",
            "approved_by",
            "approved_at",
        ]
    )
    for entry in entries[:MAX_CSV_ROWS]:
        duration_seconds = int(entry.get("duration_seconds") or 0)
        writer.writerow(
            [
                entry.get("id"),
                entry.get("project_id"),
                entry.get("project_name") or "",
                entry.get("user"),
                entry.get("start_at"),
                entry.get("end_at") or "",
                duration_seconds,
                round(duration_seconds / 3600.0, 2),
                entry.get("note") or "",
                entry.get("approval_status") or "",
                entry.get("approved_by") or "",
                entry.get("approved_at") or "",
            ]
        )
    return output.getvalue()


# ============================================================
# REVIEW LOCKS (soft locking for concurrent review)
# ============================================================
_REVIEW_LOCK_TTL_SECONDS = int(_env("REVIEW_LOCK_TTL_SECONDS", "90") or 90)


def lock_prune_expired() -> int:
    now = _now_iso()
    with _DB_LOCK:
        con = _db()
        try:
            cur = con.execute("DELETE FROM review_locks WHERE expires_at < ?", (now,))
            con.commit()
            return int(cur.rowcount or 0)
        finally:
            con.close()


def lock_acquire(
    token: str, tenant: str, user: str, roles: List[str]
) -> Dict[str, Any]:
    token = normalize_component(token)
    tenant = normalize_component(tenant).lower()
    user = normalize_component(user).lower()
    roles_s = ",".join([normalize_component(r).upper() for r in (roles or [])])

    if not token or not tenant or not user:
        return {"ok": False, "status": "invalid"}

    lock_prune_expired()

    now = _now_iso()
    exp = (datetime.now() + timedelta(seconds=_REVIEW_LOCK_TTL_SECONDS)).isoformat(
        timespec="seconds"
    )

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT locked_by, expires_at FROM review_locks WHERE token=?", (token,)
            ).fetchone()
            if row:
                locked_by = str(row["locked_by"])
                if locked_by == user:
                    con.execute(
                        "UPDATE review_locks SET heartbeat_at=?, expires_at=?, locked_roles=? WHERE token=?",
                        (now, exp, roles_s, token),
                    )
                    con.commit()
                    return {
                        "ok": True,
                        "status": "renewed",
                        "locked_by": user,
                        "expires_at": exp,
                    }
                return {
                    "ok": False,
                    "status": "conflict",
                    "locked_by": locked_by,
                    "expires_at": str(row["expires_at"]),
                }
            con.execute(
                "INSERT INTO review_locks(token, tenant, locked_by, locked_roles, locked_at, heartbeat_at, expires_at) VALUES (?,?,?,?,?,?,?)",
                (token, tenant, user, roles_s, now, now, exp),
            )
            con.commit()
            return {
                "ok": True,
                "status": "acquired",
                "locked_by": user,
                "expires_at": exp,
            }
        finally:
            con.close()


def lock_heartbeat(token: str, user: str, roles: List[str]) -> Dict[str, Any]:
    token = normalize_component(token)
    user = normalize_component(user).lower()
    roles_s = ",".join([normalize_component(r).upper() for r in (roles or [])])

    if not token or not user:
        return {"ok": False, "status": "invalid"}

    lock_prune_expired()
    now = _now_iso()
    exp = (datetime.now() + timedelta(seconds=_REVIEW_LOCK_TTL_SECONDS)).isoformat(
        timespec="seconds"
    )

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT locked_by, expires_at FROM review_locks WHERE token=?", (token,)
            ).fetchone()
            if not row:
                return {"ok": False, "status": "missing"}
            if str(row["locked_by"]) != user:
                return {
                    "ok": False,
                    "status": "conflict",
                    "locked_by": str(row["locked_by"]),
                    "expires_at": str(row["expires_at"]),
                }
            con.execute(
                "UPDATE review_locks SET heartbeat_at=?, expires_at=?, locked_roles=? WHERE token=?",
                (now, exp, roles_s, token),
            )
            con.commit()
            return {
                "ok": True,
                "status": "renewed",
                "locked_by": user,
                "expires_at": exp,
            }
        finally:
            con.close()


def lock_release(token: str, user: str) -> Dict[str, Any]:
    token = normalize_component(token)
    user = normalize_component(user).lower()
    if not token or not user:
        return {"ok": False, "status": "invalid"}

    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT locked_by FROM review_locks WHERE token=?", (token,)
            ).fetchone()
            if not row:
                return {"ok": True, "status": "missing"}
            if str(row["locked_by"]) != user:
                return {
                    "ok": False,
                    "status": "conflict",
                    "locked_by": str(row["locked_by"]),
                }
            con.execute("DELETE FROM review_locks WHERE token=?", (token,))
            con.commit()
            return {"ok": True, "status": "released"}
        finally:
            con.close()


# ============================================================
# FOLDER HELPERS
# ============================================================
def parse_folder_fields(folder_name: str) -> Dict[str, str]:
    folder_name = folder_name.strip()
    parts = folder_name.split("_")
    out = {"kdnr": "", "name": "", "addr": "", "plzort": ""}

    if not parts:
        return out

    if re.match(r"^\d{3,}$", parts[0]):
        out["kdnr"] = parts[0]
        rest = parts[1:]
    else:
        rest = parts

    plz_idx = None
    for i, t in enumerate(rest):
        if re.match(r"^\d{5}$", t):
            plz_idx = i
            break

    if plz_idx is not None:
        plz = rest[plz_idx]
        ort = "_".join(rest[plz_idx + 1 :]) if plz_idx + 1 < len(rest) else ""
        out["plzort"] = normalize_component(f"{plz} {ort.replace('_', ' ')}").strip()
        before = rest[:plz_idx]
    else:
        before = rest

    addr_start = None
    for i, t in enumerate(before):
        low = t.lower()
        if any(
            x in low
            for x in [
                "str",
                "straße",
                "strasse",
                "weg",
                "allee",
                "platz",
                "ring",
                "damm",
                "ufer",
                "gasse",
            ]
        ):
            addr_start = i
            break

    if addr_start is not None:
        out["name"] = normalize_component(
            " ".join(before[:addr_start]).replace("_", " ")
        )
        out["addr"] = normalize_component(
            " ".join(before[addr_start:]).replace("_", " ")
        )
    else:
        out["name"] = normalize_component(" ".join(before).replace("_", " "))

    return out


def find_existing_customer_folders(base_path: Path, kdnr: str) -> List[Path]:
    """
    Tenant-aware search without changing signature:
    - First: search direct child dirs: <base_path>/<kdnr>_*
    - If none: search one tenant level deep: <base_path>/<tenant>/<kdnr>_*
    """
    kdnr = normalize_component(kdnr)
    if not kdnr:
        return []
    base_path = Path(base_path)
    if not base_path.exists():
        return []

    out: List[Path] = []
    prefix = f"{kdnr}_"

    try:
        for p in base_path.iterdir():
            if p.is_dir() and p.name.startswith(prefix):
                out.append(p)
    except Exception:
        pass

    if out:
        return sorted(out)

    try:
        for tdir in base_path.iterdir():
            if not tdir.is_dir():
                continue
            if re.match(r"^\d{3,}_", tdir.name):
                continue
            for p in tdir.iterdir():
                if p.is_dir() and p.name.startswith(prefix):
                    out.append(p)
    except Exception:
        pass

    return sorted(out)


def best_match_object_folder(
    existing: List[Path], addr: str, plzort: str
) -> Tuple[Optional[Path], float]:
    addr_n = _norm_for_match(addr)
    plz_n = _norm_for_match(plzort)

    best: Optional[Path] = None
    best_score = 0.0

    for f in existing:
        fields = parse_folder_fields(f.name)
        a2 = _norm_for_match(fields.get("addr", ""))
        p2 = _norm_for_match(fields.get("plzort", ""))

        s1 = SequenceMatcher(None, addr_n, a2).ratio() if addr_n or a2 else 0.0
        s2 = SequenceMatcher(None, plz_n, p2).ratio() if plz_n or p2 else 0.0
        score = (0.6 * s2) + (0.4 * s1)

        if score > best_score:
            best_score = score
            best = f

    return best, best_score


def detect_object_duplicates_for_kdnr(
    kdnr: str, threshold: float = DEFAULT_DUP_SIM_THRESHOLD
) -> List[Dict[str, Any]]:
    kdnr = normalize_component(kdnr)
    folders = find_existing_customer_folders(BASE_PATH, kdnr)
    names = [(f, _norm_for_match(f.name)) for f in folders]
    out: List[Dict[str, Any]] = []

    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            f1, n1 = names[i]
            f2, n2 = names[j]
            if not n1 or not n2:
                continue
            sim = SequenceMatcher(None, n1, n2).ratio()
            if sim >= float(threshold):
                out.append(
                    {
                        "a": str(f1),
                        "b": str(f2),
                        "name_a": f1.name,
                        "name_b": f2.name,
                        "similarity": round(sim, 4),
                        "hint": "Mögliche Dublette (Admin: mergen/repair ohne Datenverlust).",
                    }
                )
    out.sort(key=lambda x: x["similarity"], reverse=True)
    return out


# ============================================================
# EXTRACTION / OCR
# ============================================================
def _extract_pdf_text(fp: Path) -> str:
    if PdfReader is None:
        return ""
    try:
        reader = PdfReader(str(fp))
        texts: List[str] = []
        for page in reader.pages[: max(1, OCR_MAX_PAGES)]:
            try:
                t = page.extract_text() or ""
                if t:
                    texts.append(t)
            except Exception:
                continue
        return "\n".join(texts).strip()
    except Exception:
        return ""


def _ocr_pdf(fp: Path) -> str:
    if fitz is None or pytesseract is None or Image is None:
        return ""
    try:
        doc = fitz.open(str(fp))
        texts: List[str] = []
        for i in range(min(len(doc), OCR_MAX_PAGES)):
            page = doc.load_page(i)
            pix = page.get_pixmap(dpi=250)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            txt = pytesseract.image_to_string(img, lang="deu+eng")
            if txt:
                texts.append(txt)
        return "\n".join(texts).strip()
    except Exception:
        return ""


def _ocr_image(fp: Path) -> str:
    if pytesseract is None or Image is None:
        return ""
    try:
        img = Image.open(str(fp))
        txt = pytesseract.image_to_string(img, lang="deu+eng")
        return (txt or "").strip()
    except Exception:
        return ""


def _extract_docx_text(fp: Path) -> str:
    if DocxDocument is not None:
        try:
            doc = DocxDocument(str(fp))
            paras = []
            for i, p in enumerate(doc.paragraphs):
                if i >= MAX_DOCX_PARAS:
                    break
                t = (p.text or "").strip()
                if t:
                    paras.append(t)
            return "\n".join(paras).strip()
        except Exception:
            pass

    try:
        with zipfile.ZipFile(str(fp), "r") as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        xml = (
            xml.replace("</w:p>", "\n").replace("</w:tr>", "\n").replace("</w:tc>", " ")
        )
        xml = re.sub(r"<[^>]+>", "", xml)
        xml = re.sub(r"[ \t]+", " ", xml)
        xml = re.sub(r"\n{3,}", "\n\n", xml)
        return xml.strip()
    except Exception:
        return ""


def _xlsx_shared_strings(z: zipfile.ZipFile) -> List[str]:
    try:
        sxml = z.read("xl/sharedStrings.xml").decode("utf-8", errors="ignore")
    except Exception:
        return []
    out = re.findall(r"<t[^>]*>(.*?)</t>", sxml, flags=re.IGNORECASE | re.DOTALL)
    return [re.sub(r"\s+", " ", x).strip() for x in out]


def _extract_xlsx_text(fp: Path) -> str:
    if openpyxl is not None:
        try:
            wb = openpyxl.load_workbook(str(fp), read_only=True, data_only=True)
            lines: List[str] = []
            for ws in wb.worksheets[:3]:
                lines.append(f"[Sheet] {ws.title}")
                rcount = 0
                for row in ws.iter_rows(values_only=True):
                    rcount += 1
                    if rcount > MAX_XLSX_ROWS:
                        break
                    vals = []
                    for v in row[:MAX_XLSX_COLS]:
                        if v is None:
                            continue
                        s = normalize_component(v)
                        if s:
                            vals.append(s)
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines).strip()
        except Exception:
            pass

    try:
        with zipfile.ZipFile(str(fp), "r") as z:
            sst = _xlsx_shared_strings(z)
            sheet_names = [
                n
                for n in z.namelist()
                if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
            ]
            sheet_names = sorted(sheet_names)[:3]
            out_lines: List[str] = []
            for sname in sheet_names:
                xml = z.read(sname).decode("utf-8", errors="ignore")
                out_lines.append(f"[SheetXML] {Path(sname).name}")
                vals = re.findall(r"<v>(.*?)</v>", xml, flags=re.IGNORECASE | re.DOTALL)
                cleaned: List[str] = []
                for v in vals[: MAX_XLSX_ROWS * 3]:
                    v = re.sub(r"\s+", " ", v).strip()
                    if not v:
                        continue
                    if v.isdigit():
                        idx = int(v)
                        if 0 <= idx < len(sst):
                            v = sst[idx]
                    cleaned.append(v)
                for i in range(0, min(len(cleaned), MAX_XLSX_ROWS * 10), MAX_XLSX_COLS):
                    chunk = cleaned[i : i + MAX_XLSX_COLS]
                    if chunk:
                        out_lines.append(" | ".join(chunk))
            return "\n".join(out_lines).strip()
    except Exception:
        return ""


def _extract_csv_text(fp: Path) -> str:
    raw = None
    for enc in ("utf-8-sig", "utf-8", "latin1"):
        try:
            raw = fp.read_text(encoding=enc, errors="strict")
            break
        except Exception:
            continue
    if raw is None:
        try:
            raw = fp.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""

    sio = io.StringIO(raw)
    try:
        sample = raw[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except Exception:
            dialect = csv.excel
        reader = csv.reader(sio, dialect)
        lines: List[str] = []
        for r_i, row in enumerate(reader):
            if r_i >= MAX_CSV_ROWS:
                break
            cols = [normalize_component(c) for c in row[:MAX_CSV_COLS]]
            cols = [c for c in cols if c]
            if cols:
                lines.append(" | ".join(cols))
        return "\n".join(lines).strip()
    except Exception:
        return _clip_text(raw.strip(), 50_000)


def _extract_eml_text(fp: Path) -> str:
    try:
        msg = BytesParser(policy=policy.default).parsebytes(_read_bytes(fp))
    except Exception:
        return ""

    hdr = []
    for k in ("Subject", "From", "To", "Cc", "Date"):
        try:
            v = msg.get(k)
            if v:
                hdr.append(f"{k}: {v}")
        except Exception:
            pass

    parts_text: List[str] = []

    def add_text(t: str) -> None:
        t = (t or "").strip()
        if t:
            parts_text.append(t)

    try:
        if msg.is_multipart():
            for part in msg.walk():
                ctype = (part.get_content_type() or "").lower()
                disp = (part.get_content_disposition() or "").lower()
                if disp == "attachment":
                    continue
                try:
                    payload = part.get_content()
                except Exception:
                    payload = None

                if ctype == "text/plain":
                    if isinstance(payload, str):
                        add_text(payload)
                elif ctype == "text/html":
                    if isinstance(payload, str):
                        add_text(_html_to_text(payload))
        else:
            ctype = (msg.get_content_type() or "").lower()
            try:
                payload = msg.get_content()
            except Exception:
                payload = None
            if isinstance(payload, str):
                add_text(_html_to_text(payload) if ctype == "text/html" else payload)
    except Exception:
        pass

    body = "\n\n".join(parts_text).strip()
    all_text = "\n".join(hdr + (["", body] if body else [])).strip()
    return all_text


def _extract_html_file(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
        return _html_to_text(raw)
    except Exception:
        return ""


def _extract_md_text(fp: Path) -> str:
    try:
        return (fp.read_text(encoding="utf-8", errors="ignore") or "").strip()
    except Exception:
        return ""


def _extract_rtf_text(fp: Path) -> str:
    raw = ""
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        try:
            raw = fp.read_text(encoding="latin1", errors="ignore")
        except Exception:
            return ""
    raw = re.sub(r"{\\\*[^}]*}", " ", raw)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def _extract_json_text(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""
    try:
        obj = json.loads(raw)
    except Exception:
        return _clip_text(raw.strip(), 50_000)

    out: List[str] = []

    def walk(x: Any, prefix: str = "") -> None:
        if isinstance(x, dict):
            for k, v in x.items():
                walk(v, f"{prefix}{k}.")
        elif isinstance(x, list):
            for i, v in enumerate(x[:2000]):
                walk(v, f"{prefix}{i}.")
        else:
            s = normalize_component(x)
            if s:
                out.append(f"{prefix}{s}")

    walk(obj)
    return "\n".join(out).strip()


def _extract_xml_text(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""
    raw = re.sub(r"(?is)<[^>]+>", " ", raw)
    raw = (
        raw.replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
    )
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def _extract_msg_text(fp: Path) -> str:
    if extract_msg is None:
        return ""
    try:
        m = extract_msg.Message(str(fp))
        m.process()
        parts: List[str] = []
        if getattr(m, "subject", None):
            parts.append(f"Subject: {m.subject}")
        if getattr(m, "sender", None):
            parts.append(f"From: {m.sender}")
        if getattr(m, "to", None):
            parts.append(f"To: {m.to}")
        if getattr(m, "date", None):
            parts.append(f"Date: {m.date}")
        if getattr(m, "body", None):
            parts.append(str(m.body))
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _extract_text(fp: Path) -> Tuple[str, bool]:
    """
    Returns (text, used_ocr)
    """
    ext = fp.suffix.lower()

    if ext == ".txt":
        try:
            t = fp.read_text(encoding="utf-8", errors="ignore")
            return (t or "").strip(), False
        except Exception:
            return "", False

    if ext == ".md":
        t = _extract_md_text(fp)
        return _clip_text(t), False

    if ext == ".rtf":
        t = _extract_rtf_text(fp)
        return _clip_text(t), False

    if ext in (".html", ".htm"):
        t = _extract_html_file(fp)
        return _clip_text(t), False

    if ext == ".xml":
        t = _extract_xml_text(fp)
        return _clip_text(t), False

    if ext == ".json":
        t = _extract_json_text(fp)
        return _clip_text(t), False

    if ext == ".csv":
        t = _extract_csv_text(fp)
        return _clip_text(t), False

    if ext == ".docx":
        t = _extract_docx_text(fp)
        return _clip_text(t), False

    if ext == ".xlsx":
        t = _extract_xlsx_text(fp)
        return _clip_text(t), False

    if ext == ".eml":
        t = _extract_eml_text(fp)
        return _clip_text(t), False

    if ext == ".msg":
        t = _extract_msg_text(fp)
        return _clip_text(t), False

    if ext == ".pdf":
        t = _extract_pdf_text(fp)
        if len(t) >= MIN_TEXT_LEN_BEFORE_OCR:
            return _clip_text(t), False
        o = _ocr_pdf(fp)
        if o:
            return _clip_text(o), True
        return _clip_text(t), False

    if ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"):
        o = _ocr_image(fp)
        return _clip_text(o), True if o else False

    return "", False


# ============================================================
# HEURISTIC PARSING (SUGGESTIONS)
# ============================================================
_DOCTYPE_KEYWORDS = [
    (
        "H_RECHNUNG",
        [
            r"\bh[ _-]?rechnung\b",
            r"\bhändler[ _-]?rechnung\b",
            r"\bhaendler[ _-]?rechnung\b",
            r"\bdealer[ _-]?invoice\b",
        ],
    ),
    (
        "H_ANGEBOT",
        [
            r"\bh[ _-]?angebot\b",
            r"\bhändler[ _-]?angebot\b",
            r"\bhaendler[ _-]?angebot\b",
            r"\bdealer[ _-]?offer\b",
        ],
    ),
    ("RECHNUNG", [r"\brechnung\b", r"\binvoice\b"]),
    ("ANGEBOT", [r"\bangebot\b", r"\bquotation\b", r"\boffer\b"]),
    ("AUFTRAGSBESTAETIGUNG", [r"\bauftragsbest", r"\border confirmation\b"]),
    ("MAHNUNG", [r"\bmahnung\b", r"\breminder\b"]),
    ("NACHTRAG", [r"\bnachtrag\b"]),
    ("AW", [r"\baufmaß\b", r"\baufmass\b", r"\bmaß\b", r"\bmass\b"]),
    ("FOTO", [r"\bfoto\b", r"\bimage\b"]),
]


def _detect_doctype(text: str, filename: str) -> str:
    hay = f"{filename}\n{text}".lower()
    best = ("SONSTIGES", 0)
    for dt, pats in _DOCTYPE_KEYWORDS:
        score = 0
        for p in pats:
            if re.search(p, hay, flags=re.IGNORECASE):
                score += 1
        if score > best[1]:
            best = (dt, score)
    return best[0]


def _find_kdnr_candidates(text: str) -> List[Tuple[str, float]]:
    cands: List[str] = []
    for m in re.finditer(
        r"(kunden[\s\-]*nr\.?\s*[:#]?\s*)(\d{3,})", text, flags=re.IGNORECASE
    ):
        cands.append(m.group(2))
    for m in re.finditer(r"(kdnr\.?\s*[:#]?\s*)(\d{3,})", text, flags=re.IGNORECASE):
        cands.append(m.group(2))

    if not cands:
        for m in re.finditer(r"\b(\d{4,6})\b", text):
            v = m.group(1)
            if v.startswith("20"):
                continue
            cands.append(v)

    freq: Dict[str, int] = {}
    for c in cands:
        freq[c] = freq.get(c, 0) + 1

    ranked = sorted(freq.items(), key=lambda x: (-x[1], x[0]))
    out: List[Tuple[str, float]] = []
    for num, f in ranked[:8]:
        score = round(min(1.0, 0.55 + 0.1 * f), 2)
        out.append((num, score))
    return out


def _find_dates(text: str) -> Tuple[str, List[Dict[str, Any]]]:
    cands: List[Dict[str, Any]] = []

    for m in re.finditer(r"\b(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})\b", text):
        raw = m.group(0)
        norm = parse_excel_like_date(raw)
        if norm:
            cands.append({"raw": raw, "date": norm, "reason": "DMY"})

    for m in re.finditer(r"\b(20\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})\b", text):
        raw = m.group(0)
        norm = parse_excel_like_date(raw)
        if norm:
            cands.append({"raw": raw, "date": norm, "reason": "YMD"})

    best = ""
    best_score = -1.0
    low = text.lower()

    for c in cands:
        raw = c["raw"]
        d = c["date"]
        idx = low.find(raw.lower())
        window = low[max(0, idx - 25) : idx + 25] if idx >= 0 else ""
        score = 0.1
        if "datum" in window:
            score += 0.9
        if "rechnung" in window or "angebot" in window:
            score += 0.2
        if score > best_score:
            best_score = score
            best = d

    seen = set()
    uniq: List[Dict[str, Any]] = []
    for c in cands:
        if c["date"] in seen:
            continue
        seen.add(c["date"])
        uniq.append(c)

    return best, uniq[:12]


def _find_name_addr_plzort(text: str) -> Tuple[List[str], List[str], List[str]]:
    lines = [normalize_component(x) for x in (text or "").splitlines()]
    lines = [line for line in lines if line]

    plzort: List[str] = []
    addr: List[str] = []
    name: List[str] = []

    for line in lines:
        m = re.search(r"\b(\d{5})\s+([A-Za-zÄÖÜäöüß\- ]{2,})\b", line)
        if m:
            candidate = normalize_component(f"{m.group(1)} {m.group(2)}")
            if candidate not in plzort:
                plzort.append(candidate)

    for line in lines:
        if re.search(
            r"\b(str\.?|straße|strasse|weg|allee|platz|ring|damm|ufer|gasse)\b",
            line,
            flags=re.IGNORECASE,
        ) and re.search(r"\b\d{1,4}[a-zA-Z]?\b", line):
            if line not in addr:
                addr.append(line)

    for line in lines[:15]:
        if len(line) < 3:
            continue
        if any(
            x in line.lower()
            for x in [
                "angebot",
                "rechnung",
                "datum:",
                "kunden-nr",
                "kunden nr",
                "projekt-nr",
                "bearbeiter",
            ]
        ):
            continue
        if re.search(r"(www\.|http|tel|fax|@)", line, flags=re.IGNORECASE):
            continue
        name.append(line)
        break

    return name[:8], addr[:8], plzort[:8]


def _normalize_entity(value: str) -> str:
    return re.sub(r"\\s+", " ", value.strip().lower())


def extract_entities(text: str) -> List[Dict[str, Any]]:
    entities: List[Dict[str, Any]] = []
    if not text:
        return entities
    emails = set(re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\\.[A-Za-z]{2,}", text))
    for email in emails:
        entities.append({"entity_type": "email", "value": email})

    phones = set(re.findall(r"(\\+?\\d[\\d\\s()/.-]{6,}\\d)", text))
    for phone in phones:
        entities.append({"entity_type": "phone", "value": phone})

    kdnr_matches = re.findall(
        r"\\b(?:KDNR|Kundennr|KundenNr)\\s*[:#]?\\s*(\\d{3,})\\b", text, re.IGNORECASE
    )
    for kdnr in set(kdnr_matches):
        entities.append({"entity_type": "kdnr", "value": kdnr})

    invoice_matches = re.findall(
        r"\\b(?:Rechnung|Angebot|Auftrag|Lieferschein|Bestellung)\\D{0,8}(\\d{3,}[\\-/]?\\d*)\\b",
        text,
        re.IGNORECASE,
    )
    for inv in set(invoice_matches):
        entities.append({"entity_type": "doc_number", "value": inv})

    date, _ = _find_dates(text)
    if date:
        entities.append({"entity_type": "date", "value": date})

    names, addrs, plzort = _find_name_addr_plzort(text)
    for name in names[:2]:
        entities.append({"entity_type": "customer_name", "value": name})
    for addr in addrs[:2]:
        entities.append({"entity_type": "address", "value": addr})
    for plz in plzort[:2]:
        entities.append({"entity_type": "postal_city", "value": plz})

    return entities


def _store_entities(
    con: sqlite3.Connection, tenant_id: str, doc_id: str, text: str
) -> None:
    entities = extract_entities(text)
    if not entities:
        return
    seen = set()
    for ent in entities:
        value = str(ent.get("value", "") or "")
        if not value:
            continue
        norm = _normalize_entity(value)
        key = (ent.get("entity_type"), norm)
        if key in seen:
            continue
        seen.add(key)
        con.execute(
            """
            INSERT INTO entities(tenant_id, doc_id, entity_type, value, norm_value, meta_json, created_at)
            VALUES (?,?,?,?,?,?,?)
            """,
            (
                tenant_id,
                doc_id,
                str(ent.get("entity_type", "")),
                value,
                norm,
                json.dumps(ent.get("meta", {})),
                _now_iso(),
            ),
        )


# ============================================================
# INDEX / SEARCH
# ============================================================
def _index_tokens(
    text: str, extra: Optional[List[str]] = None, limit: int = 120
) -> str:
    tokens: List[str] = []
    if extra:
        tokens.extend(extra)
    if text:
        tokens.extend(re.split(r"[\\s,;:/()\\[\\]{}<>]+", text))
    cleaned = []
    seen = set()
    for tok in tokens:
        norm = _norm_for_match(tok)
        if len(norm) < 2:
            continue
        if norm in seen:
            continue
        seen.add(norm)
        cleaned.append(norm)
        if len(cleaned) >= limit:
            break
    return " ".join(cleaned)


def _index_extract_fields(text: str, file_name: str) -> Dict[str, str]:
    names, addrs, plzort = _find_name_addr_plzort(text)
    entities = extract_entities(text)
    doc_number = ""
    for ent in entities:
        if ent.get("entity_type") == "doc_number":
            doc_number = str(ent.get("value", ""))
            break
    if not doc_number:
        match = re.search(r"\\b(\\d{4,}[\\-/]?\\d*)\\b", file_name)
        if match:
            doc_number = match.group(1)
    address = " ".join([*addrs[:1], *plzort[:1]]).strip()
    return {
        "customer_name": names[0] if names else "",
        "address": address,
        "doc_number": doc_number,
    }


def _index_put(con: sqlite3.Connection, row: Dict[str, Any]) -> None:
    doc_id = str(row.get("doc_id", "") or "")
    if not doc_id:
        return
    tokens = _index_tokens(
        " ".join(
            [
                str(row.get("file_name", "") or ""),
                str(row.get("doctype", "") or ""),
                str(row.get("kdnr", "") or ""),
                str(row.get("customer_name", "") or ""),
                str(row.get("address", "") or ""),
                str(row.get("doc_number", "") or ""),
                str(row.get("content", "") or ""),
            ]
        ),
        extra=[row.get("doc_date", "")],
    )
    con.execute("DELETE FROM docs_index WHERE doc_id = ?", (doc_id,))
    con.execute(
        """
        INSERT INTO docs_index(
            doc_id, tenant_id, kdnr, doctype, customer_name, address,
            doc_date, doc_number, file_name, file_path, tokens, snippet, updated_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            doc_id,
            str(row.get("tenant_id", "") or ""),
            str(row.get("kdnr", "") or ""),
            str(row.get("doctype", "") or ""),
            str(row.get("customer_name", "") or ""),
            str(row.get("address", "") or ""),
            str(row.get("doc_date", "") or ""),
            str(row.get("doc_number", "") or ""),
            str(row.get("file_name", "") or ""),
            str(row.get("file_path", "") or ""),
            tokens,
            _clip_text(str(row.get("snippet", "") or ""), 240),
            _now_iso(),
        ),
    )


def _fts_put(con: sqlite3.Connection, row: Dict[str, Any]) -> None:
    if not (_has_fts5(con) and _table_exists(con, "docs_fts")):
        return

    doc_id = str(row.get("doc_id", "") or "")
    if not doc_id:
        return

    con.execute("DELETE FROM docs_fts WHERE doc_id = ?", (doc_id,))
    if _column_exists(con, "docs_fts", "tenant_id"):
        con.execute(
            """
            INSERT INTO docs_fts(doc_id, tenant_id, kdnr, doctype, doc_date, file_name, file_path, content)
            VALUES (?,?,?,?,?,?,?,?)
            """,
            (
                doc_id,
                str(row.get("tenant_id", "") or ""),
                str(row.get("kdnr", "") or ""),
                str(row.get("doctype", "") or ""),
                str(row.get("doc_date", "") or ""),
                str(row.get("file_name", "") or ""),
                str(row.get("file_path", "") or ""),
                _clip_text(str(row.get("content", "") or ""), MAX_EXTRACT_CHARS),
            ),
        )
    else:
        con.execute(
            """
            INSERT INTO docs_fts(doc_id, kdnr, doctype, doc_date, file_name, file_path, content)
            VALUES (?,?,?,?,?,?,?)
            """,
            (
                doc_id,
                str(row.get("kdnr", "") or ""),
                str(row.get("doctype", "") or ""),
                str(row.get("doc_date", "") or ""),
                str(row.get("file_name", "") or ""),
                str(row.get("file_path", "") or ""),
                _clip_text(str(row.get("content", "") or ""), MAX_EXTRACT_CHARS),
            ),
        )


def _compute_group_key(kdnr: str, doctype: str, doc_date: str, file_name: str) -> str:
    k = normalize_component(kdnr)
    t = normalize_component(doctype).upper()
    d = parse_excel_like_date(doc_date) or ""
    stem = Path(file_name).stem
    stem = re.sub(r"_(v\d+|\d{6})$", "", stem, flags=re.IGNORECASE)
    stem_n = _norm_for_match(stem)
    raw = f"{k}|{t}|{d}|{stem_n}"
    return _sha256_bytes(raw.encode("utf-8"))


def index_upsert_document(
    *,
    doc_id: str,
    group_key: str,
    kdnr: str,
    object_folder: str,
    doctype: str,
    doc_date: str,
    file_name: str,
    file_path: str,
    extracted_text: str,
    used_ocr: bool,
    note: str = "",
    tenant_id: str = "",
) -> None:
    tenant_id = (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )
    with _DB_LOCK:
        con = _db()
        try:
            exists = con.execute(
                "SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)
            ).fetchone()
            if not exists:
                con.execute(
                    "INSERT INTO docs(doc_id, group_key, tenant_id, kdnr, object_folder, doctype, doc_date, created_at) VALUES (?,?,?,?,?,?,?,?)",
                    (
                        doc_id,
                        group_key,
                        tenant_id,
                        kdnr,
                        object_folder,
                        doctype,
                        doc_date or "",
                        _now_iso(),
                    ),
                )

            row = con.execute(
                "SELECT MAX(version_no) AS mx FROM versions WHERE doc_id=?", (doc_id,)
            ).fetchone()
            mx = int(row["mx"] or 0) if row else 0
            version_no = mx + 1

            con.execute(
                """
                INSERT INTO versions(doc_id, version_no, bytes_sha256, file_name, file_path, extracted_text, used_ocr, note, created_at, tenant_id)
                VALUES (?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    doc_id,
                    version_no,
                    doc_id,
                    file_name,
                    file_path,
                    _clip_text(extracted_text, MAX_EXTRACT_CHARS),
                    1 if used_ocr else 0,
                    note,
                    _now_iso(),
                    tenant_id,
                ),
            )

            _fts_put(
                con,
                {
                    "doc_id": doc_id,
                    "tenant_id": tenant_id,
                    "kdnr": kdnr,
                    "doctype": doctype,
                    "doc_date": doc_date or "",
                    "file_name": file_name,
                    "file_path": file_path,
                    "content": extracted_text or "",
                },
            )
            extra = _index_extract_fields(extracted_text or "", file_name)
            _index_put(
                con,
                {
                    "doc_id": doc_id,
                    "tenant_id": tenant_id,
                    "kdnr": kdnr,
                    "doctype": doctype,
                    "doc_date": doc_date or "",
                    "file_name": file_name,
                    "file_path": file_path,
                    "content": extracted_text or "",
                    "snippet": extracted_text or "",
                    **extra,
                },
            )
            _store_entities(con, tenant_id, doc_id, extracted_text)
            con.commit()
        finally:
            con.close()


def assistant_search(
    query: str,
    kdnr: str = "",
    limit: int = ASSISTANT_DEFAULT_LIMIT,
    role: str = "ADMIN",
    tenant_id: str = "",
) -> List[Dict[str, Any]]:
    """
    Tenant note:
    - If you stored kdnr as "TENANT:1234", search by the same.
    - If you pass only "1234" and TENANT_DEFAULT is set, it will auto-prefix.
    """
    query = normalize_component(query)
    kdnr_in = normalize_component(kdnr)
    tenant_id = (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )

    if kdnr_in and ":" not in kdnr_in:
        kdnr_in = _tenant_prefix_kdnr(TENANT_DEFAULT, kdnr_in)

    if not query:
        return []

    with _DB_LOCK:
        con = _db()
        try:
            use_fts = _has_fts5(con) and _table_exists(con, "docs_fts")
            rows: List[sqlite3.Row] = []

            if use_fts:
                tokens = [t for t in re.split(r"\\s+", query) if t]
                q = " OR ".join(tokens) if tokens else query

                if kdnr_in:
                    rows = con.execute(
                        """
                        SELECT f.doc_id, d.kdnr, d.doctype, d.doc_date, f.file_name, f.file_path,
                               snippet(docs_fts, 7, '', '', ' … ', 12) AS snip
                        FROM docs_fts f
                        JOIN docs d ON d.doc_id=f.doc_id
                        WHERE docs_fts MATCH ? AND d.kdnr=? AND d.tenant_id=?
                        LIMIT ?
                        """,
                        (q, kdnr_in, tenant_id, int(limit)),
                    ).fetchall()
                else:
                    rows = con.execute(
                        """
                        SELECT f.doc_id, d.kdnr, d.doctype, d.doc_date, f.file_name, f.file_path,
                               snippet(docs_fts, 7, '', '', ' … ', 12) AS snip
                        FROM docs_fts f
                        JOIN docs d ON d.doc_id=f.doc_id
                        WHERE docs_fts MATCH ? AND d.tenant_id=?
                        LIMIT ?
                        """,
                        (q, tenant_id, int(limit)),
                    ).fetchall()
            else:
                if not _table_exists(con, "docs_index"):
                    return []
                tokens = [t for t in re.split(r"\\s+", query) if t]
                like = f"%{_norm_for_match(query)}%"
                if kdnr_in:
                    rows = con.execute(
                        """
                        SELECT d.doc_id, d.kdnr, d.doctype, d.doc_date,
                               x.file_name, x.file_path, x.snippet AS snip
                        FROM docs d
                        JOIN docs_index x ON x.doc_id=d.doc_id
                        WHERE d.kdnr=? AND x.tenant_id=?
                          AND x.tokens LIKE ?
                        LIMIT ?
                        """,
                        (kdnr_in, tenant_id, like, int(limit)),
                    ).fetchall()
                else:
                    rows = con.execute(
                        """
                        SELECT d.doc_id, d.kdnr, d.doctype, d.doc_date,
                               x.file_name, x.file_path, x.snippet AS snip
                        FROM docs d
                        JOIN docs_index x ON x.doc_id=d.doc_id
                        WHERE x.tenant_id=?
                          AND x.tokens LIKE ?
                        LIMIT ?
                        """,
                        (tenant_id, like, int(limit)),
                    ).fetchall()

                if not rows and tokens:
                    for tok in tokens[:3]:
                        like_tok = f"%{_norm_for_match(tok)}%"
                        rows = con.execute(
                            """
                            SELECT d.doc_id, d.kdnr, d.doctype, d.doc_date,
                                   x.file_name, x.file_path, x.snippet AS snip
                            FROM docs d
                            JOIN docs_index x ON x.doc_id=d.doc_id
                            WHERE x.tenant_id=? AND x.tokens LIKE ?
                            LIMIT ?
                            """,
                            (tenant_id, like_tok, int(limit)),
                        ).fetchall()
                        if rows:
                            break

            out: List[Dict[str, Any]] = []
            for r in rows:
                doc_id = str(r["doc_id"])
                vc = con.execute(
                    "SELECT COUNT(*) AS c FROM versions WHERE doc_id=?", (doc_id,)
                ).fetchone()
                version_count = int(vc["c"] or 0) if vc else 0
                out.append(
                    {
                        "doc_id": doc_id,
                        "token": doc_id,
                        "kdnr": str(r["kdnr"] or ""),
                        "doctype": str(r["doctype"] or ""),
                        "doc_date": str(r["doc_date"] or ""),
                        "file_name": str(r["file_name"] or ""),
                        "file_path": str(r["file_path"] or ""),
                        "version_count": version_count,
                        "note": "",
                        "preview": str(r["snip"] or ""),
                    }
                )
            return out
        finally:
            con.close()


def assistant_suggest(query: str, tenant_id: str = "", limit: int = 3) -> List[str]:
    try:
        from rapidfuzz import fuzz, process  # type: ignore
    except Exception:
        return []

    query = normalize_component(query)
    tenant_id = (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )
    if not query:
        return []

    with _DB_LOCK:
        con = _db()
        try:
            if not _table_exists(con, "docs_index"):
                return []
            rows = con.execute(
                """
                SELECT customer_name, kdnr, doctype, doc_number, file_name
                FROM docs_index
                WHERE tenant_id=?
                """,
                (tenant_id,),
            ).fetchall()
        finally:
            con.close()

    candidates = set()
    for r in rows:
        for key in ("customer_name", "kdnr", "doctype", "doc_number", "file_name"):
            val = str(r[key] or "").strip()
            if val:
                candidates.add(val)
    if not candidates:
        return []
    matches = process.extract(
        query, list(candidates), scorer=fuzz.partial_ratio, limit=limit
    )
    return [m[0] for m in matches if m[1] >= 70]


def index_run_full(base_path: Optional[Path] = None) -> Dict[str, Any]:
    base = Path(base_path) if base_path else BASE_PATH
    if not base.exists():
        return {
            "ok": True,
            "indexed": 0,
            "skipped": 0,
            "errors": 0,
            "skipped_by_reason": {
                "unsupported_ext": 0,
                "already_indexed": 0,
                "no_text": 0,
                "parse_failed": 0,
            },
        }

    indexed = 0
    skipped = 0
    errors = 0

    skipped_by_reason = {
        "unsupported_ext": 0,
        "already_indexed": 0,
        "no_text": 0,
        "parse_failed": 0,
    }

    try:
        for fp in base.rglob("*"):
            if not fp.is_file():
                continue

            ext = fp.suffix.lower()
            if ext not in SUPPORTED_EXT:
                skipped += 1
                skipped_by_reason["unsupported_ext"] += 1
                continue

            try:
                b = _read_bytes(fp)
                doc_id = _sha256_bytes(b)

                with _DB_LOCK:
                    con = _db()
                    try:
                        exists = con.execute(
                            "SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)
                        ).fetchone()
                    finally:
                        con.close()

                if exists:
                    skipped += 1
                    skipped_by_reason["already_indexed"] += 1
                    continue

                file_name = fp.name

                tenant = _effective_tenant(_infer_tenant_from_path(fp))

                kdnr_raw = ""
                object_folder = ""
                for part in reversed(fp.parts):
                    if re.match(r"^\d{3,}_", part):
                        kdnr_raw = part.split("_", 1)[0]
                        object_folder = part
                        break

                if TENANT_REQUIRE and not tenant:
                    skipped += 1
                    skipped_by_reason["parse_failed"] += 1
                    continue

                kdnr_idx = _tenant_prefix_kdnr(tenant, kdnr_raw) if kdnr_raw else ""
                object_folder_tag = (
                    _tenant_object_folder_tag(tenant, object_folder)
                    if object_folder
                    else ""
                )

                text, used_ocr = _extract_text(fp)
                if not text or len(text.strip()) < 3:
                    skipped += 1
                    skipped_by_reason["no_text"] += 1
                    continue

                doctype = _detect_doctype(text, file_name)
                best_date, _ = _find_dates(text)
                group_key = _compute_group_key(kdnr_idx, doctype, best_date, file_name)

                index_upsert_document(
                    doc_id=doc_id,
                    group_key=group_key,
                    kdnr=kdnr_idx,
                    object_folder=object_folder_tag,
                    doctype=doctype,
                    doc_date=best_date or "",
                    file_name=file_name,
                    file_path=str(fp),
                    extracted_text=text,
                    used_ocr=used_ocr,
                    note="indexed_by_full_scan",
                    tenant_id=tenant,
                )
                indexed += 1

            except Exception:
                skipped += 1
                skipped_by_reason["parse_failed"] += 1
                continue

    except Exception:
        errors += 1

    return {
        "ok": True,
        "indexed": indexed,
        "skipped": skipped,
        "errors": errors,
        "skipped_by_reason": skipped_by_reason,
    }


def import_run(*, import_root: Path, user: str = "", role: str = "") -> Dict[str, Any]:
    root = Path(import_root)
    if not root.exists():
        return {
            "ok": False,
            "indexed": 0,
            "skipped": 0,
            "errors": 0,
            "skipped_by_reason": {},
        }

    indexed = 0
    skipped = 0
    errors = 0
    skipped_by_reason = {
        "unsupported_ext": 0,
        "already_indexed": 0,
        "no_text": 0,
        "parse_failed": 0,
    }

    for fp in root.rglob("*"):
        if not fp.is_file():
            continue

        ext = fp.suffix.lower()
        if ext not in SUPPORTED_EXT:
            skipped += 1
            skipped_by_reason["unsupported_ext"] += 1
            continue

        try:
            b = _read_bytes(fp)
            doc_id = _sha256_bytes(b)

            with _DB_LOCK:
                con = _db()
                try:
                    exists = con.execute(
                        "SELECT 1 FROM versions WHERE doc_id=? AND file_path=? LIMIT 1",
                        (doc_id, str(fp)),
                    ).fetchone()
                finally:
                    con.close()

            if exists:
                skipped += 1
                skipped_by_reason["already_indexed"] += 1
                continue

            tenant = _effective_tenant(_infer_tenant_from_path(fp))
            if TENANT_REQUIRE and not tenant:
                skipped += 1
                skipped_by_reason["parse_failed"] += 1
                continue

            kdnr_raw = ""
            object_folder = ""
            for part in reversed(fp.parts):
                if re.match(r"^\d{3,}_", part):
                    kdnr_raw = part.split("_", 1)[0]
                    object_folder = part
                    break

            kdnr_idx = _tenant_prefix_kdnr(tenant, kdnr_raw) if kdnr_raw else ""
            object_folder_tag = (
                _tenant_object_folder_tag(tenant, object_folder)
                if object_folder
                else ""
            )

            text, used_ocr = _extract_text(fp)
            if not text or len(text.strip()) < 3:
                skipped += 1
                skipped_by_reason["no_text"] += 1
                continue

            doctype = _detect_doctype(text, fp.name)
            best_date, _ = _find_dates(text)
            group_key = _compute_group_key(kdnr_idx, doctype, best_date, fp.name)

            index_upsert_document(
                doc_id=doc_id,
                group_key=group_key,
                kdnr=kdnr_idx,
                object_folder=object_folder_tag,
                doctype=doctype,
                doc_date=best_date or "",
                file_name=fp.name,
                file_path=str(fp),
                extracted_text=text,
                used_ocr=used_ocr,
                note="import_run",
                tenant_id=tenant,
            )
            indexed += 1
            if callable(audit_log):
                audit_log(
                    user=user or "system",
                    role=role or "SYSTEM",
                    action="import_file",
                    target=doc_id,
                    meta={"path": str(fp), "root": str(root)},
                    tenant_id=tenant,
                )
        except Exception:
            errors += 1
            skipped_by_reason["parse_failed"] += 1
            continue

    return {
        "ok": True,
        "indexed": indexed,
        "skipped": skipped,
        "errors": errors,
        "skipped_by_reason": skipped_by_reason,
    }


def index_rebuild(base_path: Optional[Path] = None) -> Dict[str, Any]:
    with _DB_LOCK:
        con = _db()
        try:
            if _table_exists(con, "docs_fts"):
                con.execute("DELETE FROM docs_fts")
            if _table_exists(con, "docs_index"):
                con.execute("DELETE FROM docs_index")
            con.commit()
        finally:
            con.close()
    return index_run_full(base_path=base_path)


def index_warmup(tenant_id: str = "") -> Dict[str, Any]:
    tenant_id = normalize_component(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            fts_enabled = _has_fts5(con)
            tenants: List[str] = []
            if tenant_id:
                tenants = [tenant_id]
            elif _table_exists(con, "tenants"):
                rows = con.execute("SELECT tenant_id FROM tenants").fetchall()
                tenants = [str(r["tenant_id"]) for r in rows if r["tenant_id"]]
            elif _table_exists(con, "docs_index"):
                rows = con.execute(
                    "SELECT DISTINCT tenant_id FROM docs_index"
                ).fetchall()
                tenants = [str(r["tenant_id"]) for r in rows if r["tenant_id"]]
            for t in tenants:
                con.execute(
                    "SELECT COUNT(*) AS c FROM docs_index WHERE tenant_id=?",
                    (t,),
                ).fetchone()
            return {"ok": True, "fts_enabled": fts_enabled, "tenants": len(tenants)}
        finally:
            con.close()


def set_db_path(new_path: Path) -> None:
    global DB_PATH
    DB_PATH = Path(new_path)
    db_init()


def get_db_info() -> Dict[str, Any]:
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute("PRAGMA user_version").fetchone()
            schema_version = int(row[0] if row else 0)
            tenants = 0
            if _table_exists(con, "tenants"):
                trow = con.execute("SELECT COUNT(*) AS c FROM tenants").fetchone()
                tenants = int(trow["c"] or 0) if trow else 0
            return {
                "path": str(DB_PATH),
                "schema_version": schema_version,
                "tenants": tenants,
            }
        finally:
            con.close()


def set_base_path(new_path: Path) -> None:
    global BASE_PATH
    BASE_PATH = Path(new_path)


def get_profile() -> Dict[str, Any]:
    name = _env("KUKANILEA_PROFILE", "").strip()
    if not name:
        name = f"db:{DB_PATH.stem}"
    return {
        "name": name,
        "db_path": str(DB_PATH),
        "base_path": str(BASE_PATH),
    }


def get_health_stats(tenant_id: str = "") -> Dict[str, Any]:
    tenant_id = (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )
    with _DB_LOCK:
        con = _db()
        try:
            fts_enabled = _has_fts5(con) and _table_exists(con, "docs_fts")
            if _table_exists(con, "docs"):
                row = con.execute(
                    "SELECT COUNT(*) AS c FROM docs WHERE tenant_id=?",
                    (tenant_id,),
                ).fetchone()
                doc_count = int(row["c"] or 0) if row else 0
            else:
                doc_count = 0
            last_indexed_at = None
            if _table_exists(con, "docs_index"):
                row = con.execute(
                    "SELECT MAX(updated_at) AS ts FROM docs_index WHERE tenant_id=?",
                    (tenant_id,),
                ).fetchone()
                last_indexed_at = str(row["ts"]) if row and row["ts"] else None
            return {
                "doc_count": doc_count,
                "last_indexed_at": last_indexed_at,
                "fts_enabled": bool(fts_enabled),
            }
        finally:
            con.close()


def audit_list(*, tenant_id: str = "", limit: int = 200) -> List[Dict[str, Any]]:
    tenant_id = normalize_component(tenant_id)
    limit = max(1, min(int(limit), 2000))
    with _DB_LOCK:
        con = _db()
        try:
            if tenant_id and _column_exists(con, "audit", "tenant_id"):
                rows = con.execute(
                    "SELECT * FROM audit WHERE tenant_id=? ORDER BY id DESC LIMIT ?",
                    (tenant_id, limit),
                ).fetchall()
            else:
                rows = con.execute(
                    "SELECT * FROM audit ORDER BY id DESC LIMIT ?",
                    (limit,),
                ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()


# ============================================================
# BACKGROUND ANALYSIS -> PENDING
# ============================================================
def analyze_to_pending(src: Path) -> str:
    PENDING_DIR.mkdir(parents=True, exist_ok=True)

    src = Path(src)
    if not src.exists():
        raise FileNotFoundError(src)

    tenant = _effective_tenant(_infer_tenant_from_path(src))

    t = _token()
    payload: Dict[str, Any] = {
        "status": "ANALYZING",
        "progress": 1.0,
        "progress_phase": "Init…",
        "error": "",
        "path": str(src),
        "filename": src.name,
        "tenant_suggested": tenant,
        "used_ocr": False,
        "extracted_text": "",
        "preview": "",
        "doctype_suggested": "SONSTIGES",
        "doc_date_suggested": "",
        "doc_date_candidates": [],
        "kdnr_ranked": [],
        "name_suggestions": [],
        "addr_suggestions": [],
        "plzort_suggestions": [],
    }
    write_pending(t, payload)

    th = threading.Thread(target=_analyze_worker, args=(t,), daemon=True)
    th.start()
    return t


def start_background_analysis(src: Path) -> str:
    return analyze_to_pending(src)


def _set_progress(token: str, p: float, phase: str) -> None:
    d = read_pending(token) or {}
    d["progress"] = float(max(0.0, min(100.0, p)))
    d["progress_phase"] = phase
    write_pending(token, d)


def _analyze_worker(token: str) -> None:
    try:
        d = read_pending(token)
        if not d:
            return
        src = Path(d.get("path", ""))
        if not src.exists():
            d["status"] = "ERROR"
            d["error"] = "file_missing"
            d["progress"] = 0.0
            d["progress_phase"] = ""
            write_pending(token, d)
            return

        _set_progress(token, 5.0, "Datei lesen…")
        try:
            b = _read_bytes(src)
            doc_id = _sha256_bytes(b)
            d["doc_id"] = doc_id
        except Exception:
            d["doc_id"] = ""

        _set_progress(token, 18.0, "Text extrahieren…")
        text, used_ocr = _extract_text(src)
        d["used_ocr"] = bool(used_ocr)
        d["extracted_text"] = text
        d["preview"] = (text or "").strip()[:900].strip()

        _set_progress(token, 40.0, "Dokumenttyp/Datum erkennen…")
        d["doctype_suggested"] = _detect_doctype(text, src.name)
        best_date, date_cands = _find_dates(text)
        d["doc_date_suggested"] = best_date or ""
        d["doc_date_candidates"] = date_cands

        _set_progress(token, 60.0, "Kundendaten erkennen…")
        d["kdnr_ranked"] = _find_kdnr_candidates(text)
        names, addrs, plzs = _find_name_addr_plzort(text)
        d["name_suggestions"] = names
        d["addr_suggestions"] = addrs
        d["plzort_suggestions"] = plzs

        _set_progress(token, 85.0, "Fertigstellen…")
        d["status"] = "READY"
        d["progress"] = 100.0
        d["progress_phase"] = "Bereit"
        write_pending(token, d)

    except Exception as e:
        d = read_pending(token) or {}
        d["status"] = "ERROR"
        d["error"] = str(e)
        d["progress"] = 0.0
        d["progress_phase"] = ""
        write_pending(token, d)


# ============================================================
# ARCHIVE / PROCESS
# ============================================================
def _doctype_code(doctype: str) -> str:
    doctype = normalize_component(doctype).upper()

    if doctype in {"HAENDLERRECHNUNG", "HÄNDLERRECHNUNG", "H_RE"}:
        doctype = "H_RECHNUNG"
    if doctype in {"HAENDLERANGEBOT", "HÄNDLERANGEBOT", "H_ANG"}:
        doctype = "H_ANGEBOT"

    return {
        "H_RECHNUNG": "H_RE",
        "H_ANGEBOT": "H_ANG",
        "RECHNUNG": "RE",
        "ANGEBOT": "ANG",
        "AUFTRAGSBESTAETIGUNG": "AB",
        "AW": "AW",
        "MAHNUNG": "MAH",
        "NACHTRAG": "NTR",
        "FOTO": "FOTO",
        "SONSTIGES": "DOC",
    }.get(doctype, "DOC")


def _compose_object_folder(kdnr: str, name: str, addr: str, plzort: str) -> str:
    parts = [_safe_fs(kdnr), _safe_fs(name), _safe_fs(addr), _safe_fs(plzort)]
    parts = [p for p in parts if p]
    return "_".join(parts)[:180]


def _compose_filename(
    doctype: str, doc_date: str, kdnr: str, name: str, addr: str, plzort: str, ext: str
) -> str:
    code = _doctype_code(doctype)
    d = parse_excel_like_date(doc_date) or ""
    parts: List[str] = [code]
    if d:
        parts.append(d)
    for x in [kdnr, name, addr, plzort]:
        xs = _safe_fs(x)
        if xs:
            parts.append(xs)

    base = "_".join(parts)
    base = re.sub(r"_+", "_", base).strip("_")[:160]
    return f"{base}{ext}"


def _next_version_suffix(target_dir: Path, base_name: str, ext: str) -> str:
    stem = Path(base_name).stem
    n = 2
    while True:
        cand = f"{stem}_v{n}{ext}"
        if not (target_dir / cand).exists():
            return cand
        n += 1


def _source_allowlist_roots() -> List[Path]:
    roots = [EINGANG, BASE_PATH, PENDING_DIR, DONE_DIR]
    out: List[Path] = []
    for root in roots:
        try:
            out.append(Path(root).resolve())
        except Exception:
            continue
    return out


def is_allowed_source_path(p: Path) -> bool:
    try:
        rp = Path(p).expanduser().resolve()
    except Exception:
        return False
    for root in _source_allowlist_roots():
        if rp == root or str(rp).startswith(str(root) + os.sep):
            return True
    return False


def _db_latest_version_path_for_doc(doc_id: str, tenant_id: str = "") -> str:
    doc_id = normalize_component(doc_id)
    if not doc_id:
        return ""
    tenant_id = _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or ""
    with _DB_LOCK:
        con = _db()
        try:
            order_expr = (
                "version_no DESC, id DESC"
                if _column_exists(con, "versions", "version_no")
                else "id DESC"
            )
            if tenant_id and _column_exists(con, "versions", "tenant_id"):
                row = con.execute(
                    """
                    SELECT file_path FROM versions
                    WHERE doc_id=? AND tenant_id=?
                    ORDER BY """
                    + order_expr
                    + """
                    LIMIT 1
                    """,
                    (doc_id, tenant_id),
                ).fetchone()
            else:
                row = con.execute(
                    """
                    SELECT file_path FROM versions
                    WHERE doc_id=?
                    ORDER BY """
                    + order_expr
                    + """
                    LIMIT 1
                    """,
                    (doc_id,),
                ).fetchone()
            return str(row["file_path"]) if row and row["file_path"] else ""
        finally:
            con.close()


def resolve_source_path(
    token: str, pending: Optional[Dict[str, Any]] = None, tenant_id: str = ""
) -> Optional[Path]:
    token = normalize_component(token)
    pending = pending or read_pending(token) or {}
    if not isinstance(pending, dict):
        pending = {}

    direct_raw = str(pending.get("path", "") or "").strip()
    if direct_raw:
        direct = Path(direct_raw)
        if direct.exists() and is_allowed_source_path(direct):
            return direct

    doc_id = normalize_component(pending.get("doc_id") or token)
    tenant_ctx = _effective_tenant(
        tenant_id, pending.get("tenant_id", ""), pending.get("tenant", "")
    )
    latest_path = db_latest_path_for_doc(doc_id, tenant_id=tenant_ctx) if doc_id else ""

    if latest_path:
        candidate = Path(latest_path)
        if candidate.exists() and is_allowed_source_path(candidate):
            return candidate

    return None


def db_latest_path_for_doc(doc_id: str, tenant_id: str = "") -> str:
    """Return latest file_path for a given doc_id (sha256 bytes), or ''."""
    return _db_latest_version_path_for_doc(doc_id, tenant_id=tenant_id)


def db_path_for_doc(doc_id: str, tenant_id: str = "") -> str:
    """Return a file path for a doc_id by checking versions + docs_index."""
    doc_id = normalize_component(doc_id)
    if not doc_id:
        return ""
    tenant_id = _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or ""
    with _DB_LOCK:
        con = _db()
        try:
            if tenant_id and _column_exists(con, "versions", "tenant_id"):
                row = con.execute(
                    """
                    SELECT file_path FROM versions
                    WHERE doc_id=? AND tenant_id=?
                    ORDER BY id DESC LIMIT 1
                    """,
                    (doc_id, tenant_id),
                ).fetchone()
            else:
                row = con.execute(
                    "SELECT file_path FROM versions WHERE doc_id=? ORDER BY id DESC LIMIT 1",
                    (doc_id,),
                ).fetchone()
            if row and row["file_path"]:
                return str(row["file_path"])
            if _table_exists(con, "docs_index"):
                if tenant_id:
                    row = con.execute(
                        "SELECT file_path FROM docs_index WHERE doc_id=? AND tenant_id=? LIMIT 1",
                        (doc_id, tenant_id),
                    ).fetchone()
                else:
                    row = con.execute(
                        "SELECT file_path FROM docs_index WHERE doc_id=? LIMIT 1",
                        (doc_id,),
                    ).fetchone()
                if row and row["file_path"]:
                    return str(row["file_path"])
            return ""
        finally:
            con.close()


def _db_has_doc(doc_id: str) -> bool:
    with _DB_LOCK:
        con = _db()
        try:
            r = con.execute(
                "SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)
            ).fetchone()
            return bool(r)
        finally:
            con.close()


def process_with_answers(src: Path, answers: Dict[str, Any]) -> Tuple[Path, Path, bool]:
    src = Path(src)
    if not src.exists():
        raise FileNotFoundError(src)

    tenant = _effective_tenant(
        answers.get("tenant"), answers.get("mandant"), _infer_tenant_from_path(src)
    )
    if TENANT_REQUIRE and not tenant:
        raise ValueError("tenant/mandant missing (TENANT_REQUIRE=1)")

    kdnr_raw = normalize_component(answers.get("kdnr", ""))
    use_existing = normalize_component(answers.get("use_existing", ""))
    name = normalize_component(answers.get("name", ""))
    addr = normalize_component(answers.get("addr", ""))
    plzort = normalize_component(answers.get("plzort", ""))
    doctype = normalize_component(answers.get("doctype", "SONSTIGES")).upper()
    doc_date = parse_excel_like_date(answers.get("document_date", "")) or ""

    if not kdnr_raw:
        raise ValueError("kdnr missing")

    kdnr_idx = _tenant_prefix_kdnr(tenant, kdnr_raw)

    BASE_PATH.mkdir(parents=True, exist_ok=True)
    tenant_dir = BASE_PATH / _safe_fs(tenant) if tenant else BASE_PATH
    tenant_dir.mkdir(parents=True, exist_ok=True)

    created_new_object = False
    if use_existing:
        folder = Path(use_existing)
        if not folder.exists() or not folder.is_dir():
            folder = tenant_dir / _compose_object_folder(kdnr_raw, name, addr, plzort)
            created_new_object = True
    else:
        folder_name = _compose_object_folder(kdnr_raw, name, addr, plzort)
        folder = tenant_dir / folder_name
        if not folder.exists():
            created_new_object = True

    folder.mkdir(parents=True, exist_ok=True)

    ext = src.suffix.lower()
    final_name = _compose_filename(doctype, doc_date, kdnr_raw, name, addr, plzort, ext)
    target = folder / final_name

    b = _read_bytes(src)
    doc_id = _sha256_bytes(b)

    object_folder_tag = _tenant_object_folder_tag(tenant, folder.name)

    if target.exists():
        try:
            if _sha256_bytes(_read_bytes(target)) == doc_id:
                try:
                    src.unlink()
                except Exception:
                    pass

                if not _db_has_doc(doc_id):
                    text, used_ocr = _extract_text(target)
                    group_key = _compute_group_key(
                        kdnr_idx, doctype, doc_date, target.name
                    )
                    index_upsert_document(
                        doc_id=doc_id,
                        group_key=group_key,
                        kdnr=kdnr_idx,
                        object_folder=object_folder_tag,
                        doctype=doctype,
                        doc_date=doc_date or "",
                        file_name=target.name,
                        file_path=str(target),
                        extracted_text=text,
                        used_ocr=used_ocr,
                        note="dedupe_same_bytes_existing_target",
                        tenant_id=tenant,
                    )
                return folder, target, created_new_object
        except Exception:
            pass

        final_name = _next_version_suffix(folder, final_name, ext)
        target = folder / final_name

    try:
        src.replace(target)
    except Exception:
        target.write_bytes(b)
        try:
            src.unlink()
        except Exception:
            pass

    text, used_ocr = _extract_text(target)
    group_key = _compute_group_key(kdnr_idx, doctype, doc_date, target.name)

    note = ""
    with _DB_LOCK:
        con = _db()
        try:
            g = con.execute(
                "SELECT doc_id FROM docs WHERE group_key=? LIMIT 1", (group_key,)
            ).fetchone()
            if g and str(g["doc_id"]) != doc_id:
                note = "new_version_same_group_key"
        finally:
            con.close()

    index_upsert_document(
        doc_id=doc_id,
        group_key=group_key,
        kdnr=kdnr_idx,
        object_folder=object_folder_tag,
        doctype=doctype,
        doc_date=doc_date or "",
        file_name=target.name,
        file_path=str(target),
        extracted_text=text,
        used_ocr=used_ocr,
        note=note,
        tenant_id=tenant,
    )

    return folder, target, created_new_object


def _crm_tenant(tenant_id: str) -> str:
    return (
        _effective_tenant(tenant_id) or _effective_tenant(TENANT_DEFAULT) or "default"
    )


def _crm_new_id() -> str:
    return uuid.uuid4().hex


def _crm_event_id(entity_id: str) -> int:
    h = hashlib.sha256((entity_id or "").encode("utf-8")).hexdigest()[:15]
    return max(1, int(h, 16))


def _is_locked_error(exc: Exception) -> bool:
    msg = str(exc).lower()
    return "database is locked" in msg or "database is busy" in msg


def _run_write_txn(fn):
    backoff = [0.05, 0.1, 0.2, 0.4, 0.8]
    for idx, wait in enumerate(backoff, start=1):
        should_retry = False
        with _DB_LOCK:
            con = _db()
            try:
                con.execute("BEGIN IMMEDIATE")
                result = fn(con)
                con.commit()
                return result
            except sqlite3.OperationalError as exc:
                try:
                    con.rollback()
                except Exception:
                    pass
                if _is_locked_error(exc):
                    if idx < len(backoff):
                        should_retry = True
                    else:
                        raise ValueError("db_locked")
                else:
                    raise
            except Exception:
                try:
                    con.rollback()
                except Exception:
                    pass
                raise
            finally:
                con.close()
        if should_retry:
            time.sleep(wait)
            continue
    raise ValueError("db_locked")


def _parse_money_to_cents(
    value: Any, *, field: str, allow_none: bool = False
) -> Optional[int]:
    if value is None or value == "":
        if allow_none:
            return None
        return 0
    if isinstance(value, bool):
        raise ValueError(f"{field}_invalid")
    try:
        dec = Decimal(str(value).strip())
        dec = dec.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        raise ValueError(f"{field}_invalid")
    return int(dec * 100)


def _cents_from_legacy(
    cents_value: Any, legacy_value: Any, *, nullable: bool = False
) -> Optional[int]:
    if cents_value is not None:
        try:
            cents_int = int(cents_value)
            if cents_int > 0:
                return cents_int
            if cents_int == 0 and not nullable:
                return 0
        except Exception:
            pass
    if legacy_value is None or str(legacy_value).strip() == "":
        return None if nullable else 0
    try:
        return _parse_money_to_cents(legacy_value, field="legacy", allow_none=nullable)
    except Exception:
        return None if nullable else 0


def _crm_require_customer(
    con: sqlite3.Connection, tenant_id: str, customer_id: str
) -> None:
    row = con.execute(
        "SELECT id FROM customers WHERE tenant_id=? AND id=?",
        (tenant_id, customer_id),
    ).fetchone()
    if not row:
        raise ValueError("not_found")


def _crm_require_contact(
    con: sqlite3.Connection, tenant_id: str, contact_id: str
) -> None:
    row = con.execute(
        "SELECT id FROM contacts WHERE tenant_id=? AND id=?",
        (tenant_id, contact_id),
    ).fetchone()
    if not row:
        raise ValueError("not_found")


def _crm_require_deal(
    con: sqlite3.Connection, tenant_id: str, deal_id: str
) -> Dict[str, Any]:
    row = con.execute(
        "SELECT * FROM deals WHERE tenant_id=? AND id=?",
        (tenant_id, deal_id),
    ).fetchone()
    if not row:
        raise ValueError("not_found")
    return dict(row)


def _crm_quote_items(
    con: sqlite3.Connection, tenant_id: str, quote_id: str
) -> List[Dict[str, Any]]:
    rows = con.execute(
        """
        SELECT id, quote_id, description, qty, unit_price_cents, line_total_cents, created_at, updated_at
        FROM quote_items
        WHERE tenant_id=? AND quote_id=?
        ORDER BY created_at, id
        """,
        (tenant_id, quote_id),
    ).fetchall()
    return [dict(r) for r in rows]


def customers_create(
    tenant_id: str,
    name: str,
    vat_id: Optional[str] = None,
    notes: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> str:
    tenant_id = _crm_tenant(tenant_id)
    customer_name = normalize_component(name)
    if not customer_name:
        raise ValueError("validation_error")

    customer_id = _crm_new_id()
    now = _now_iso()

    def _tx(con: sqlite3.Connection) -> str:
        con.execute(
            """
            INSERT INTO customers(id, tenant_id, name, vat_id, notes, created_at, updated_at)
            VALUES (?,?,?,?,?,?,?)
            """,
            (
                customer_id,
                tenant_id,
                customer_name,
                normalize_component(vat_id),
                normalize_component(notes),
                now,
                now,
            ),
        )
        event_append(
            event_type="crm_customer",
            entity_type="customer",
            entity_id=_crm_event_id(customer_id),
            payload={
                "schema_version": 1,
                "source": "core/customers_create",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "customer",
                "entity_id": customer_id,
                "action": "created",
                "data": {"name": customer_name},
            },
            con=con,
        )
        return customer_id

    return _run_write_txn(_tx)


def customers_get(tenant_id: str, customer_id: str) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT * FROM customers WHERE tenant_id=? AND id=?",
                (tenant_id, customer_id),
            ).fetchone()
            if not row:
                raise ValueError("not_found")
            return dict(row)
        finally:
            con.close()


def customers_list(
    tenant_id: str,
    limit: int = 100,
    offset: int = 0,
    query: Optional[str] = None,
) -> List[Dict[str, Any]]:
    tenant_id = _crm_tenant(tenant_id)
    lim = max(1, min(int(limit), 500))
    off = max(0, int(offset))
    q = normalize_component(query)

    with _DB_LOCK:
        con = _db()
        try:
            if q:
                rows = con.execute(
                    """
                    SELECT * FROM customers
                    WHERE tenant_id=? AND (LOWER(name) LIKE LOWER(?) OR LOWER(COALESCE(vat_id,'')) LIKE LOWER(?))
                    ORDER BY updated_at DESC, id DESC
                    LIMIT ? OFFSET ?
                    """,
                    (tenant_id, f"%{q}%", f"%{q}%", lim, off),
                ).fetchall()
            else:
                rows = con.execute(
                    """
                    SELECT * FROM customers
                    WHERE tenant_id=?
                    ORDER BY updated_at DESC, id DESC
                    LIMIT ? OFFSET ?
                    """,
                    (tenant_id, lim, off),
                ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()


def customers_update(
    tenant_id: str,
    customer_id: str,
    *,
    name: Optional[str] = None,
    vat_id: Optional[str] = None,
    notes: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)

    def _tx(con: sqlite3.Connection) -> Dict[str, Any]:
        row = con.execute(
            "SELECT * FROM customers WHERE tenant_id=? AND id=?",
            (tenant_id, customer_id),
        ).fetchone()
        if not row:
            raise ValueError("not_found")
        new_name = (
            normalize_component(name) if name is not None else str(row["name"] or "")
        )
        if not new_name:
            raise ValueError("validation_error")
        new_vat = (
            normalize_component(vat_id)
            if vat_id is not None
            else str(row["vat_id"] or "")
        )
        new_notes = (
            normalize_component(notes) if notes is not None else str(row["notes"] or "")
        )
        con.execute(
            """
            UPDATE customers
            SET name=?, vat_id=?, notes=?, updated_at=?
            WHERE tenant_id=? AND id=?
            """,
            (new_name, new_vat, new_notes, _now_iso(), tenant_id, customer_id),
        )
        event_append(
            event_type="crm_customer",
            entity_type="customer",
            entity_id=_crm_event_id(customer_id),
            payload={
                "schema_version": 1,
                "source": "core/customers_update",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "customer",
                "entity_id": customer_id,
                "action": "updated",
                "data": {"name": new_name},
            },
            con=con,
        )
        out = con.execute(
            "SELECT * FROM customers WHERE tenant_id=? AND id=?",
            (tenant_id, customer_id),
        ).fetchone()
        return dict(out) if out else {}

    return _run_write_txn(_tx)


def contacts_create(
    tenant_id: str,
    customer_id: str,
    name: str,
    email: Optional[str] = None,
    phone: Optional[str] = None,
    role: Optional[str] = None,
    notes: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> str:
    tenant_id = _crm_tenant(tenant_id)
    contact_name = normalize_component(name)
    if not contact_name:
        raise ValueError("validation_error")
    contact_id = _crm_new_id()
    now = _now_iso()

    def _tx(con: sqlite3.Connection) -> str:
        _crm_require_customer(con, tenant_id, customer_id)
        con.execute(
            """
            INSERT INTO contacts(id, tenant_id, customer_id, name, email, phone, role, notes, created_at, updated_at)
            VALUES (?,?,?,?,?,?,?,?,?,?)
            """,
            (
                contact_id,
                tenant_id,
                customer_id,
                contact_name,
                normalize_component(email),
                normalize_component(phone),
                normalize_component(role),
                normalize_component(notes),
                now,
                now,
            ),
        )
        event_append(
            event_type="crm_contact",
            entity_type="contact",
            entity_id=_crm_event_id(contact_id),
            payload={
                "schema_version": 1,
                "source": "core/contacts_create",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "contact",
                "entity_id": contact_id,
                "action": "created",
                "data": {"customer_id": customer_id},
            },
            con=con,
        )
        return contact_id

    return _run_write_txn(_tx)


def contacts_list_by_customer(tenant_id: str, customer_id: str) -> List[Dict[str, Any]]:
    tenant_id = _crm_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            _crm_require_customer(con, tenant_id, customer_id)
            rows = con.execute(
                """
                SELECT * FROM contacts
                WHERE tenant_id=? AND customer_id=?
                ORDER BY name, created_at
                """,
                (tenant_id, customer_id),
            ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()


def deals_create(
    tenant_id: str,
    customer_id: str,
    title: str,
    stage: str = "lead",
    value_cents: Optional[int] = None,
    currency: str = "EUR",
    notes: Optional[str] = None,
    project_id: Optional[int] = None,
    probability: Optional[int] = None,
    expected_close_date: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> str:
    tenant_id = _crm_tenant(tenant_id)
    title_norm = normalize_component(title)
    stage_norm = normalize_component(stage).lower() or "lead"
    if stage_norm not in {
        "lead",
        "qualified",
        "proposal",
        "negotiation",
        "won",
        "lost",
    }:
        raise ValueError("validation_error")
    if not title_norm:
        raise ValueError("validation_error")
    if probability is not None and (int(probability) < 0 or int(probability) > 100):
        raise ValueError("validation_error")
    if expected_close_date:
        try:
            datetime.fromisoformat(str(expected_close_date))
        except Exception:
            raise ValueError("validation_error")

    deal_id = _crm_new_id()
    now = _now_iso()
    cents = _cents_from_legacy(value_cents, None, nullable=True)

    def _tx(con: sqlite3.Connection) -> str:
        _crm_require_customer(con, tenant_id, customer_id)
        con.execute(
            """
            INSERT INTO deals(
              id, tenant_id, customer_id, title, stage, project_id,
              value_cents, currency, notes, probability, expected_close_date,
              created_at, updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                deal_id,
                tenant_id,
                customer_id,
                title_norm,
                stage_norm,
                int(project_id) if project_id is not None else None,
                cents,
                normalize_component(currency) or "EUR",
                normalize_component(notes),
                int(probability) if probability is not None else None,
                str(expected_close_date) if expected_close_date else None,
                now,
                now,
            ),
        )
        event_append(
            event_type="crm_deal",
            entity_type="deal",
            entity_id=_crm_event_id(deal_id),
            payload={
                "schema_version": 1,
                "source": "core/deals_create",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "deal",
                "entity_id": deal_id,
                "action": "created",
                "data": {"customer_id": customer_id, "stage": stage_norm},
            },
            con=con,
        )
        return deal_id

    return _run_write_txn(_tx)


def deals_update_stage(
    tenant_id: str,
    deal_id: str,
    stage: str,
    actor_user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)
    stage_norm = normalize_component(stage).lower()
    if stage_norm not in {
        "lead",
        "qualified",
        "proposal",
        "negotiation",
        "won",
        "lost",
    }:
        raise ValueError("validation_error")

    def _tx(con: sqlite3.Connection) -> Dict[str, Any]:
        deal = _crm_require_deal(con, tenant_id, deal_id)
        con.execute(
            "UPDATE deals SET stage=?, updated_at=? WHERE tenant_id=? AND id=?",
            (stage_norm, _now_iso(), tenant_id, deal_id),
        )
        event_append(
            event_type="crm_deal",
            entity_type="deal",
            entity_id=_crm_event_id(deal_id),
            payload={
                "schema_version": 1,
                "source": "core/deals_update_stage",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "deal",
                "entity_id": deal_id,
                "action": "updated",
                "data": {"before_stage": deal.get("stage"), "stage": stage_norm},
            },
            con=con,
        )
        out = con.execute(
            "SELECT * FROM deals WHERE tenant_id=? AND id=?",
            (tenant_id, deal_id),
        ).fetchone()
        return dict(out) if out else {}

    return _run_write_txn(_tx)


def deals_list(
    tenant_id: str,
    stage: Optional[str] = None,
    customer_id: Optional[str] = None,
) -> List[Dict[str, Any]]:
    tenant_id = _crm_tenant(tenant_id)
    clauses = ["tenant_id=?"]
    params: List[Any] = [tenant_id]
    st = normalize_component(stage).lower()
    if st:
        clauses.append("stage=?")
        params.append(st)
    cid = normalize_component(customer_id)
    if cid:
        clauses.append("customer_id=?")
        params.append(cid)
    where_sql = " AND ".join(clauses)

    with _DB_LOCK:
        con = _db()
        try:
            rows = con.execute(
                f"SELECT * FROM deals WHERE {where_sql} ORDER BY updated_at DESC, id DESC",
                tuple(params),
            ).fetchall()
            out = [dict(r) for r in rows]
            for item in out:
                item["value_cents"] = _cents_from_legacy(
                    item.get("value_cents"), None, nullable=True
                )
            return out
        finally:
            con.close()


def _next_quote_number(con: sqlite3.Connection, tenant_id: str) -> str:
    row = con.execute(
        "SELECT COUNT(*) AS cnt FROM quotes WHERE tenant_id=?",
        (tenant_id,),
    ).fetchone()
    nxt = int((row["cnt"] if row else 0) or 0) + 1
    return f"Q-{nxt:06d}"


def quotes_create(
    tenant_id: str,
    customer_id: str,
    deal_id: Optional[str] = None,
    currency: str = "EUR",
    notes: Optional[str] = None,
    quote_number: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> str:
    tenant_id = _crm_tenant(tenant_id)
    quote_id = _crm_new_id()
    now = _now_iso()

    def _tx(con: sqlite3.Connection) -> str:
        _crm_require_customer(con, tenant_id, customer_id)
        if deal_id:
            _crm_require_deal(con, tenant_id, deal_id)
        quote_no = normalize_component(quote_number) or _next_quote_number(
            con, tenant_id
        )
        try:
            con.execute(
                """
                INSERT INTO quotes(
                  id, tenant_id, customer_id, deal_id, status, currency,
                  quote_number, subtotal_cents, tax_cents, tax_amount_cents, total_cents, notes, created_at, updated_at
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    quote_id,
                    tenant_id,
                    customer_id,
                    deal_id,
                    "draft",
                    normalize_component(currency) or "EUR",
                    quote_no,
                    0,
                    0,
                    0,
                    0,
                    normalize_component(notes),
                    now,
                    now,
                ),
            )
        except sqlite3.IntegrityError as exc:
            if (
                "idx_quotes_tenant_quote_number" in str(exc)
                or "UNIQUE" in str(exc).upper()
            ):
                raise ValueError("duplicate")
            raise
        event_append(
            event_type="crm_quote",
            entity_type="quote",
            entity_id=_crm_event_id(quote_id),
            payload={
                "schema_version": 1,
                "source": "core/quotes_create",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "quote",
                "entity_id": quote_id,
                "action": "created",
                "data": {
                    "customer_id": customer_id,
                    "deal_id": deal_id,
                    "quote_number": quote_no,
                },
            },
            con=con,
        )
        return quote_id

    return _run_write_txn(_tx)


def quotes_add_item(
    tenant_id: str,
    quote_id: str,
    description: str,
    qty: float,
    unit_price_cents: int,
    actor_user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)
    description_norm = normalize_component(description)
    if not description_norm:
        raise ValueError("validation_error")
    try:
        qty_val = Decimal(str(qty))
    except Exception:
        raise ValueError("validation_error")
    if qty_val <= 0:
        raise ValueError("validation_error")
    unit_val = int(unit_price_cents)
    if unit_val < 0:
        raise ValueError("validation_error")
    line_total = int(
        (qty_val * Decimal(unit_val)).to_integral_value(rounding=ROUND_HALF_UP)
    )
    item_id = _crm_new_id()

    def _tx(con: sqlite3.Connection) -> Dict[str, Any]:
        qrow = con.execute(
            "SELECT id FROM quotes WHERE tenant_id=? AND id=?",
            (tenant_id, quote_id),
        ).fetchone()
        if not qrow:
            raise ValueError("not_found")
        now = _now_iso()
        con.execute(
            """
            INSERT INTO quote_items(id, tenant_id, quote_id, description, qty, unit_price_cents, line_total_cents, created_at, updated_at)
            VALUES (?,?,?,?,?,?,?,?,?)
            """,
            (
                item_id,
                tenant_id,
                quote_id,
                description_norm,
                float(qty_val),
                unit_val,
                line_total,
                now,
                now,
            ),
        )
        event_append(
            event_type="crm_quote_item",
            entity_type="quote_item",
            entity_id=_crm_event_id(item_id),
            payload={
                "schema_version": 1,
                "source": "core/quotes_add_item",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "quote_item",
                "entity_id": item_id,
                "action": "created",
                "data": {"quote_id": quote_id, "line_total_cents": line_total},
            },
            con=con,
        )

        subtotal_row = con.execute(
            "SELECT COALESCE(SUM(line_total_cents),0) AS subtotal FROM quote_items WHERE tenant_id=? AND quote_id=?",
            (tenant_id, quote_id),
        ).fetchone()
        subtotal = int((subtotal_row["subtotal"] if subtotal_row else 0) or 0)
        qrow = con.execute(
            "SELECT tax_amount_cents, tax_cents FROM quotes WHERE tenant_id=? AND id=?",
            (tenant_id, quote_id),
        ).fetchone()
        existing_tax = 0
        if qrow is not None:
            existing_tax = int(
                _cents_from_legacy(
                    qrow["tax_amount_cents"]
                    if "tax_amount_cents" in qrow.keys()
                    else None,
                    qrow["tax_cents"] if "tax_cents" in qrow.keys() else None,
                    nullable=False,
                )
                or 0
            )
        total = subtotal + max(0, existing_tax)
        con.execute(
            """
            UPDATE quotes
            SET subtotal_cents=?, tax_cents=?, tax_amount_cents=?, total_cents=?, updated_at=?
            WHERE tenant_id=? AND id=?
            """,
            (
                subtotal,
                existing_tax,
                existing_tax,
                total,
                _now_iso(),
                tenant_id,
                quote_id,
            ),
        )
        event_append(
            event_type="crm_quote",
            entity_type="quote",
            entity_id=_crm_event_id(quote_id),
            payload={
                "schema_version": 1,
                "source": "core/quotes_add_item",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "quote",
                "entity_id": quote_id,
                "action": "updated",
                "data": {
                    "subtotal_cents": subtotal,
                    "tax_amount_cents": existing_tax,
                    "total_cents": total,
                },
            },
            con=con,
        )
        q = con.execute(
            "SELECT * FROM quotes WHERE tenant_id=? AND id=?",
            (tenant_id, quote_id),
        ).fetchone()
        out = dict(q) if q else {}
        out["items"] = _crm_quote_items(con, tenant_id, quote_id)
        return out

    return _run_write_txn(_tx)


def quotes_recalculate_totals(
    tenant_id: str,
    quote_id: str,
    tax_rate: Optional[float] = None,
    actor_user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)

    def _tx(con: sqlite3.Connection) -> Dict[str, Any]:
        qrow = con.execute(
            "SELECT * FROM quotes WHERE tenant_id=? AND id=?",
            (tenant_id, quote_id),
        ).fetchone()
        if not qrow:
            raise ValueError("not_found")

        subtotal_row = con.execute(
            "SELECT COALESCE(SUM(line_total_cents),0) AS subtotal FROM quote_items WHERE tenant_id=? AND quote_id=?",
            (tenant_id, quote_id),
        ).fetchone()
        subtotal = int((subtotal_row["subtotal"] if subtotal_row else 0) or 0)
        if tax_rate is None:
            existing_tax = _cents_from_legacy(
                qrow["tax_amount_cents"] if "tax_amount_cents" in qrow.keys() else None,
                qrow["tax_cents"] if "tax_cents" in qrow.keys() else None,
                nullable=False,
            )
            tax_amount = max(0, int(existing_tax or 0))
        else:
            try:
                rate = Decimal(str(tax_rate))
            except Exception:
                raise ValueError("validation_error")
            if rate < 0:
                raise ValueError("validation_error")
            tax_amount = int(
                (Decimal(subtotal) * rate).to_integral_value(rounding=ROUND_HALF_UP)
            )
        total = subtotal + tax_amount

        con.execute(
            """
            UPDATE quotes
            SET subtotal_cents=?, tax_cents=?, tax_amount_cents=?, total_cents=?, updated_at=?
            WHERE tenant_id=? AND id=?
            """,
            (subtotal, tax_amount, tax_amount, total, _now_iso(), tenant_id, quote_id),
        )

        event_append(
            event_type="crm_quote",
            entity_type="quote",
            entity_id=_crm_event_id(quote_id),
            payload={
                "schema_version": 1,
                "source": "core/quotes_recalculate_totals",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "quote",
                "entity_id": quote_id,
                "action": "updated",
                "data": {
                    "subtotal_cents": subtotal,
                    "tax_amount_cents": tax_amount,
                    "total_cents": total,
                },
            },
            con=con,
        )

        q = con.execute(
            "SELECT * FROM quotes WHERE tenant_id=? AND id=?",
            (tenant_id, quote_id),
        ).fetchone()
        out = dict(q) if q else {}
        out["items"] = _crm_quote_items(con, tenant_id, quote_id)
        return out

    return _run_write_txn(_tx)


def quotes_get(tenant_id: str, quote_id: str) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT * FROM quotes WHERE tenant_id=? AND id=?",
                (tenant_id, quote_id),
            ).fetchone()
            if not row:
                raise ValueError("not_found")
            out = dict(row)
            out["subtotal_cents"] = _cents_from_legacy(
                out.get("subtotal_cents"), None, nullable=False
            )
            out["tax_amount_cents"] = _cents_from_legacy(
                out.get("tax_amount_cents"), out.get("tax_cents"), nullable=False
            )
            out["total_cents"] = _cents_from_legacy(
                out.get("total_cents"), None, nullable=False
            )
            out["items"] = _crm_quote_items(con, tenant_id, quote_id)
            return out
        finally:
            con.close()


def quotes_create_from_deal(
    tenant_id: str,
    deal_id: str,
    actor_user_id: Optional[int] = None,
) -> Dict[str, Any]:
    tenant_id = _crm_tenant(tenant_id)
    with _DB_LOCK:
        con = _db()
        try:
            deal = _crm_require_deal(con, tenant_id, deal_id)
        finally:
            con.close()

    quote_id = quotes_create(
        tenant_id=tenant_id,
        customer_id=str(deal.get("customer_id") or ""),
        deal_id=deal_id,
        currency=str(deal.get("currency") or "EUR"),
        notes="Automatisch aus Deal erstellt",
        actor_user_id=actor_user_id,
    )

    project_id = int(deal.get("project_id") or 0)
    if project_id > 0:
        summary = time_entries_summary_by_project(
            tenant_id=tenant_id, project_id=project_id
        )
        spent_hours = float(summary.get("spent_hours") or 0.0)
        if spent_hours > 0:
            default_hour_rate = 9000
            if deal.get("value_cents"):
                try:
                    deal_value = int(deal.get("value_cents") or 0)
                    estimated_hours = spent_hours if spent_hours > 0 else 1.0
                    default_hour_rate = max(
                        100, int(round(deal_value / estimated_hours))
                    )
                except Exception:
                    default_hour_rate = 9000
            quotes_add_item(
                tenant_id=tenant_id,
                quote_id=quote_id,
                description="Arbeitsstunden",
                qty=spent_hours,
                unit_price_cents=default_hour_rate,
                actor_user_id=actor_user_id,
            )
        else:

            def _tx_note(con: sqlite3.Connection) -> None:
                con.execute(
                    "UPDATE quotes SET notes=?, updated_at=? WHERE tenant_id=? AND id=?",
                    (
                        "Automatisch erstellt; keine Zeitdaten gefunden.",
                        _now_iso(),
                        tenant_id,
                        quote_id,
                    ),
                )
                event_append(
                    event_type="crm_quote",
                    entity_type="quote",
                    entity_id=_crm_event_id(quote_id),
                    payload={
                        "schema_version": 1,
                        "source": "core/quotes_create_from_deal",
                        "actor_user_id": actor_user_id,
                        "tenant_id": tenant_id,
                        "entity_type": "quote",
                        "entity_id": quote_id,
                        "action": "updated",
                        "data": {"note": "no_time_data"},
                    },
                    con=con,
                )
                return None

            _run_write_txn(_tx_note)
    else:

        def _tx_note2(con: sqlite3.Connection) -> None:
            con.execute(
                "UPDATE quotes SET notes=?, updated_at=? WHERE tenant_id=? AND id=?",
                (
                    "Automatisch erstellt; kein Projekt am Deal verknüpft.",
                    _now_iso(),
                    tenant_id,
                    quote_id,
                ),
            )
            event_append(
                event_type="crm_quote",
                entity_type="quote",
                entity_id=_crm_event_id(quote_id),
                payload={
                    "schema_version": 1,
                    "source": "core/quotes_create_from_deal",
                    "actor_user_id": actor_user_id,
                    "tenant_id": tenant_id,
                    "entity_type": "quote",
                    "entity_id": quote_id,
                    "action": "updated",
                    "data": {"note": "no_project_link"},
                },
                con=con,
            )
            return None

        _run_write_txn(_tx_note2)

    return quotes_get(tenant_id, quote_id)


def emails_import_eml(
    tenant_id: str,
    eml_bytes: bytes,
    customer_id: Optional[str] = None,
    contact_id: Optional[str] = None,
    source_notes: Optional[str] = None,
    actor_user_id: Optional[int] = None,
) -> str:
    tenant_id = _crm_tenant(tenant_id)
    if not isinstance(eml_bytes, (bytes, bytearray)) or not eml_bytes:
        raise ValueError("validation_error")

    email_id = _crm_new_id()
    now = _now_iso()
    msg_id = ""
    from_addr = ""
    to_addrs = ""
    subject = ""
    received_at = now
    body_text = ""
    notes = normalize_component(source_notes)
    attachment_meta: list[dict[str, Any]] = []

    try:
        msg = BytesParser(policy=policy.default).parsebytes(bytes(eml_bytes))
        msg_id = normalize_component(msg.get("Message-ID") or "")
        from_parsed = getaddresses([msg.get("From") or ""])
        from_addr = normalize_component(from_parsed[0][1] if from_parsed else "")
        to_parsed = getaddresses(msg.get_all("To", []))
        to_addrs = ", ".join(a for _, a in to_parsed if a)
        subject = normalize_component(msg.get("Subject") or "")

        dt_hdr = msg.get("Date")
        if dt_hdr:
            try:
                dt = parsedate_to_datetime(dt_hdr)
                if dt is not None:
                    received_at = dt.isoformat(timespec="seconds")
            except Exception:
                pass

        if msg.is_multipart():
            for part in msg.walk():
                filename = part.get_filename()
                if filename:
                    payload = part.get_payload(decode=True) or b""
                    attachment_meta.append(
                        {
                            "filename": normalize_component(filename),
                            "size": len(payload),
                        }
                    )
                if part.get_content_type() == "text/plain" and not body_text:
                    try:
                        body_text = part.get_content()
                    except Exception:
                        body_text = ""
        else:
            if msg.get_content_type() == "text/plain":
                try:
                    body_text = msg.get_content()
                except Exception:
                    body_text = ""

        if body_text and len(body_text) > 20000:
            body_text = body_text[:20000] + "\n[truncated]"

        if body_text and ("\x00" in body_text or any(ord(ch) < 9 for ch in body_text)):
            notes = f"{notes}; invalid_plain_text" if notes else "invalid_plain_text"
            body_text = ""

    except Exception as exc:
        err = f"parse_error: {exc}"
        notes = f"{notes}; {err}" if notes else err
        body_text = ""

    def _tx(con: sqlite3.Connection) -> str:
        if customer_id:
            _crm_require_customer(con, tenant_id, customer_id)
        if contact_id:
            _crm_require_contact(con, tenant_id, contact_id)
        try:
            con.execute(
                """
                INSERT INTO emails_cache(
                  id, tenant_id, customer_id, contact_id, message_id,
                  from_addr, to_addrs, subject, received_at, body_text,
                  raw_eml, notes, attachments_json, created_at
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    email_id,
                    tenant_id,
                    customer_id,
                    contact_id,
                    msg_id,
                    from_addr,
                    to_addrs,
                    subject,
                    received_at,
                    body_text,
                    sqlite3.Binary(bytes(eml_bytes[:65536])),
                    notes,
                    json.dumps(attachment_meta, ensure_ascii=False, sort_keys=True),
                    now,
                ),
            )
        except sqlite3.IntegrityError as exc:
            if (
                "idx_emails_cache_tenant_message" in str(exc)
                or "UNIQUE" in str(exc).upper()
            ):
                raise ValueError("duplicate")
            raise

        event_append(
            event_type="crm_email",
            entity_type="email",
            entity_id=_crm_event_id(email_id),
            payload={
                "schema_version": 1,
                "source": "core/emails_import_eml",
                "actor_user_id": actor_user_id,
                "tenant_id": tenant_id,
                "entity_type": "email",
                "entity_id": email_id,
                "action": "imported",
                "data": {
                    "message_id": msg_id,
                    "subject": subject,
                    "from_addr": from_addr,
                    "received_at": received_at,
                    "raw_size": len(eml_bytes),
                    "attachments": attachment_meta,
                },
            },
            con=con,
        )
        return email_id

    return _run_write_txn(_tx)


# ============================================================
# BOOTSTRAP DIRS (on import)
# ============================================================
def _bootstrap_dirs() -> None:
    EINGANG.mkdir(parents=True, exist_ok=True)
    BASE_PATH.mkdir(parents=True, exist_ok=True)
    PENDING_DIR.mkdir(parents=True, exist_ok=True)
    DONE_DIR.mkdir(parents=True, exist_ok=True)


_bootstrap_dirs()
# db_init() is intentionally not auto-called here; your Flask runner calls db_init() at startup.
