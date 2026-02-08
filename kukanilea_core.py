#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
KUKANILEA Core (v2.4 - A2 Multi-Tenant + More Formats) — DROP-IN replacement
============================================================================

Das ist dein `kukanilea_core.py` als kompletter, kopierfertiger Code.

Wichtig (Kompatibilität):
- Exportiert die gleichen Funktionen/Variablen wie vorher (Contract).
- Behält **TOPHANDWERK_*** ENV-Variablen für Abwärtskompatibilität.
- Zusätzlich werden **KUKANILEA_*** ENV-Variablen als Alias akzeptiert (haben Vorrang).

Hinweis:
- In `kukanilea_upload.py` musst du nur den Import anpassen:
    import kukanilea_core as core
  (statt `import tophandwerk_core as core`)
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import time
import unicodedata
import zipfile
from datetime import date, datetime
from difflib import SequenceMatcher
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# Optional libs
try:
    from pypdf import PdfReader  # type: ignore
except Exception:
    PdfReader = None

try:
    import fitz  # PyMuPDF  # type: ignore
except Exception:
    fitz = None

try:
    from PIL import Image  # type: ignore
except Exception:
    Image = None

try:
    import pytesseract  # type: ignore
except Exception:
    pytesseract = None

try:
    from docx import Document as DocxDocument  # type: ignore
except Exception:
    DocxDocument = None

try:
    import openpyxl  # type: ignore
except Exception:
    openpyxl = None

# Optional for Outlook .msg
try:
    import extract_msg  # type: ignore
except Exception:
    extract_msg = None


# ============================================================
# ENV helpers (KUKANILEA_* overrides TOPHANDWERK_*)
# ============================================================
def _env(key: str, default: str = "") -> str:
    """
    Read env with alias support:
      KUKANILEA_<key> overrides TOPHANDWERK_<key>.
    key is expected WITHOUT prefix, e.g. "TENANT_DEFAULT".
    """
    k1 = f"KUKANILEA_{key}"
    k2 = f"TOPHANDWERK_{key}"
    v = os.environ.get(k1)
    if v is not None:
        return str(v)
    v = os.environ.get(k2)
    if v is not None:
        return str(v)
    return default


def _env_bool(key: str, default: str = "0") -> bool:
    v = _env(key, default).strip()
    return v in ("1", "true", "TRUE", "yes", "YES", "on", "ON")


# ============================================================
# CONFIG / PATHS
# ============================================================
# (bewusst: Default-Pfade bleiben "Tophandwerk_*", damit bestehende Daten nicht verloren gehen)
EINGANG = Path.home() / _env("EINGANG_DIRNAME", "Tophandwerk_Eingang")
BASE_PATH = Path.home() / _env("BASE_DIRNAME", "Tophandwerk_Kundenablage")
PENDING_DIR = Path.home() / _env("PENDING_DIRNAME", "Tophandwerk_Pending")
DONE_DIR = Path.home() / _env("DONE_DIRNAME", "Tophandwerk_Done")
DB_PATH = Path.home() / _env("DB_FILENAME", "Tophandwerk_DB.sqlite3")

# Multi-tenant behavior
TENANT_DEFAULT = _env("TENANT_DEFAULT", "").strip()  # e.g. "FIRMA_X"
TENANT_REQUIRE = _env_bool("TENANT_REQUIRE", "0")

# A2: include more formats (some are store-only if extraction returns "")
SUPPORTED_EXT = {
    # Documents
    ".pdf",
    ".txt",
    ".md",
    ".rtf",
    ".docx",
    ".xlsx",
    ".csv",
    ".eml",
    ".html",
    ".htm",
    ".json",
    ".xml",
    # Images (OCR)
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    # Outlook (best-effort)
    ".msg",
    # Containers / non-text (accepted, but extraction may be empty; store-only)
    ".zip",
    ".7z",
    ".rar",
    ".dwg",
    ".dxf",
    ".ifc",
    ".p7m",
    ".p7s",
    ".psd",
    ".ai",
    ".mp4",
    ".mov",
    ".mp3",
}

# OCR / Extraction limits
OCR_MAX_PAGES = 2
MIN_TEXT_LEN_BEFORE_OCR = 200

# Duplicate-detection for object folders (name/addr/plzort variations)
DEFAULT_DUP_SIM_THRESHOLD = 0.93

# Assistant search result size
ASSISTANT_DEFAULT_LIMIT = 50

# Extraction size guards (prevent huge RAM / DB bloat)
MAX_EXTRACT_CHARS = 200_000
MAX_CSV_ROWS = 2000
MAX_CSV_COLS = 60
MAX_XLSX_ROWS = 2000
MAX_XLSX_COLS = 60
MAX_DOCX_PARAS = 4000


# ============================================================
# UTIL
# ============================================================
def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _sha256_bytes(b: bytes) -> str:
    h = hashlib.sha256()
    h.update(b)
    return h.hexdigest()


def _read_bytes(fp: Path) -> bytes:
    with open(fp, "rb") as f:
        return f.read()


def _token() -> str:
    raw = f"{time.time_ns()}:{os.getpid()}:{os.urandom(8).hex()}".encode("utf-8")
    return base64.urlsafe_b64encode(hashlib.sha256(raw).digest())[:22].decode("ascii")


def normalize_component(s: Any) -> str:
    """
    Strong normalization:
    - normalize unicode (NFKC)
    - collapse whitespace
    """
    if s is None:
        return ""
    s = str(s).strip()
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _norm_for_match(s: Any) -> str:
    """
    Aggressive matching normalization:
    - lower
    - replace umlauts
    - drop non-alnum
    """
    s = normalize_component(s).lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = re.sub(r"[^a-z0-9]+", "", s)
    return s


def _atomic_write_text(fp: Path, text: str, encoding: str = "utf-8") -> None:
    fp.parent.mkdir(parents=True, exist_ok=True)
    tmp = fp.with_suffix(fp.suffix + f".tmp_{os.getpid()}_{time.time_ns()}")
    tmp.write_text(text, encoding=encoding)
    tmp.replace(fp)


def _clip_text(s: str, limit: int = MAX_EXTRACT_CHARS) -> str:
    s = s or ""
    if len(s) <= limit:
        return s
    return s[:limit]


def _html_to_text(html: str) -> str:
    """
    Very pragmatic HTML -> text (no external deps).
    """
    if not html:
        return ""
    html = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>", "\n", html)
    html = re.sub(r"(?i)</p\s*>", "\n", html)
    html = re.sub(r"(?is)<[^>]+>", " ", html)
    html = (
        html.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    )
    html = re.sub(r"[ \t]+", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()


# --------------------------
# Tenant helpers (no DB migration)
# --------------------------
def _infer_tenant_from_path(fp: Path) -> str:
    """
    Convention:
      BASE_PATH/<TENANT>/<KUNDENORDNER>/...
    Returns tenant or "".
    """
    try:
        fp = Path(fp).resolve()
        base = Path(BASE_PATH).resolve()
        parts = fp.parts
        bparts = base.parts

        for i in range(len(parts) - len(bparts) + 1):
            if parts[i : i + len(bparts)] == bparts:
                if i + len(bparts) < len(parts):
                    tenant = normalize_component(parts[i + len(bparts)])
                    if tenant and not re.match(r"^\d{3,}_", tenant):
                        return tenant
                break
    except Exception:
        pass
    return ""


def _tenant_prefix_kdnr(tenant: str, kdnr: str) -> str:
    tenant = normalize_component(tenant)
    kdnr = normalize_component(kdnr)
    if not kdnr:
        return ""
    if ":" in kdnr:
        return kdnr
    if tenant:
        return f"{tenant}:{kdnr}"
    return kdnr


def _tenant_object_folder_tag(tenant: str, object_folder: str) -> str:
    tenant = normalize_component(tenant)
    object_folder = normalize_component(object_folder)
    if tenant and object_folder:
        return f"{tenant}/{object_folder}"
    return object_folder


def _effective_tenant(*candidates: Any) -> str:
    """
    Choose first non-empty tenant from candidates, else env default.
    """
    for c in candidates:
        t = normalize_component(c)
        if t:
            return t
    return normalize_component(TENANT_DEFAULT)


def _safe_fs(s: Any) -> str:
    s = normalize_component(s)
    if not s:
        return ""
    s = re.sub(r"[^\wäöüÄÖÜß\-\.\, ]+", "", s)
    s = s.replace(" ", "_")
    s = re.sub(r"_+", "_", s).strip("_")
    return s[:60]


_DATE_PATTERNS = [
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
]


def parse_excel_like_date(s: Any) -> str:
    """
    Accept common Excel-like date strings -> normalized YYYY-MM-DD or "" if invalid.
    """
    if not s:
        return ""
    s = str(s).strip()
    if not s:
        return ""
    s = re.sub(r"\s+", " ", s).replace("T", " ")

    for pat in _DATE_PATTERNS:
        try:
            dt = datetime.strptime(s, pat)
            d = dt.date()
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    m = re.search(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        try:
            y = int(m.group(1))
            mo = int(m.group(2))
            da = int(m.group(3))
            d = date(y, mo, da)
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
    if m:
        try:
            da = int(m.group(1))
            mo = int(m.group(2))
            y = int(m.group(3))
            if y < 100:
                y = 2000 + y if y <= 68 else 1900 + y
            d = date(y, mo, da)
            if d.year < 1900 or d.year > 2099:
                return ""
            return d.strftime("%Y-%m-%d")
        except Exception:
            pass

    return ""


# ============================================================
# PENDING / DONE store (JSON files)
# ============================================================
def _pending_path(token: str) -> Path:
    return PENDING_DIR / f"{token}.json"


def _done_path(token: str) -> Path:
    return DONE_DIR / f"{token}.json"


def read_pending(token: str) -> Optional[Dict[str, Any]]:
    fp = _pending_path(token)
    if not fp.exists():
        return None
    try:
        return json.loads(fp.read_text(encoding="utf-8"))
    except Exception:
        return None


def write_pending(token: str, payload: Dict[str, Any]) -> None:
    fp = _pending_path(token)
    _atomic_write_text(fp, json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def delete_pending(token: str) -> None:
    fp = _pending_path(token)
    try:
        fp.unlink()
    except Exception:
        pass


def list_pending() -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    PENDING_DIR.mkdir(parents=True, exist_ok=True)
    for fp in sorted(PENDING_DIR.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            j = json.loads(fp.read_text(encoding="utf-8"))
            j["_token"] = fp.stem
            out.append(j)
        except Exception:
            continue
    return out


def write_done(token: str, payload: Dict[str, Any]) -> None:
    fp = _done_path(token)
    _atomic_write_text(fp, json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def read_done(token: str) -> Optional[Dict[str, Any]]:
    fp = _done_path(token)
    if not fp.exists():
        return None
    try:
        return json.loads(fp.read_text(encoding="utf-8"))
    except Exception:
        return None


# ============================================================
# SQLITE DB
# ============================================================
_DB_LOCK = threading.Lock()
_FTS5_AVAILABLE: Optional[bool] = None


def _db() -> sqlite3.Connection:
    con = sqlite3.connect(str(DB_PATH))
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL;")
    con.execute("PRAGMA synchronous=NORMAL;")
    con.execute("PRAGMA foreign_keys=ON;")
    return con


def _table_exists(con: sqlite3.Connection, name: str) -> bool:
    row = con.execute(
        "SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name=?",
        (name,),
    ).fetchone()
    return bool(row)


def _has_fts5(con: sqlite3.Connection) -> bool:
    global _FTS5_AVAILABLE
    if _FTS5_AVAILABLE is not None:
        return _FTS5_AVAILABLE
    try:
        con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS __fts5_test USING fts5(x);")
        con.execute("DROP TABLE IF EXISTS __fts5_test;")
        _FTS5_AVAILABLE = True
    except Exception:
        _FTS5_AVAILABLE = False
    return bool(_FTS5_AVAILABLE)


def db_init() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS users(
                  username TEXT PRIMARY KEY,
                  pass_sha256 TEXT NOT NULL,
                  created_at TEXT NOT NULL
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS roles(
                  username TEXT NOT NULL,
                  role TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  PRIMARY KEY(username, role),
                  FOREIGN KEY(username) REFERENCES users(username) ON DELETE CASCADE
                );
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS audit(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ts TEXT NOT NULL,
                  user TEXT NOT NULL,
                  role TEXT NOT NULL,
                  action TEXT NOT NULL,
                  target TEXT,
                  meta_json TEXT
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS docs(
                  doc_id TEXT PRIMARY KEY,                  -- sha256(file bytes)
                  group_key TEXT NOT NULL,                  -- heuristic group for versioning
                  kdnr TEXT,
                  object_folder TEXT,
                  doctype TEXT,
                  doc_date TEXT,                            -- YYYY-MM-DD (OPTIONAL)
                  created_at TEXT NOT NULL
                );
                """
            )

            con.execute(
                """
                CREATE TABLE IF NOT EXISTS versions(
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  doc_id TEXT NOT NULL,
                  version_no INTEGER NOT NULL,
                  bytes_sha256 TEXT NOT NULL,
                  file_name TEXT NOT NULL,
                  file_path TEXT NOT NULL,
                  extracted_text TEXT,
                  used_ocr INTEGER NOT NULL DEFAULT 0,
                  note TEXT,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(doc_id) REFERENCES docs(doc_id) ON DELETE CASCADE
                );
                """
            )

            con.execute("CREATE INDEX IF NOT EXISTS idx_docs_group ON docs(group_key);")
            con.execute("CREATE INDEX IF NOT EXISTS idx_docs_kdnr ON docs(kdnr);")
            con.execute("CREATE INDEX IF NOT EXISTS idx_versions_doc ON versions(doc_id);")
            con.execute("CREATE INDEX IF NOT EXISTS idx_versions_path ON versions(file_path);")

            if _has_fts5(con):
                con.execute(
                    """
                    CREATE VIRTUAL TABLE IF NOT EXISTS docs_fts
                    USING fts5(
                        doc_id UNINDEXED,
                        kdnr UNINDEXED,
                        doctype UNINDEXED,
                        doc_date UNINDEXED,
                        file_name UNINDEXED,
                        file_path UNINDEXED,
                        content,
                        tokenize='unicode61'
                    );
                    """
                )

            con.commit()
        finally:
            con.close()


# ============================================================
# RBAC
# ============================================================
def _pw_hash(pw: str) -> str:
    return _sha256_bytes((pw or "").encode("utf-8"))


def rbac_create_user(username: str, password: str) -> str:
    username = normalize_component(username).lower()
    if not username:
        raise ValueError("username required")
    if not password:
        raise ValueError("password required")

    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                "INSERT OR REPLACE INTO users(username, pass_sha256, created_at) VALUES (?,?,?)",
                (username, _pw_hash(password), _now_iso()),
            )
            con.commit()
            return username
        finally:
            con.close()


def rbac_verify_user(username: str, password: str) -> bool:
    username = normalize_component(username).lower()
    if not username or not password:
        return False
    with _DB_LOCK:
        con = _db()
        try:
            row = con.execute(
                "SELECT pass_sha256 FROM users WHERE username=?", (username,)
            ).fetchone()
            if not row:
                return False
            return str(row["pass_sha256"]) == _pw_hash(password)
        finally:
            con.close()


def rbac_assign_role(username: str, role: str) -> None:
    username = normalize_component(username).lower()
    role = normalize_component(role).upper()
    if not username or not role:
        return
    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                "INSERT OR IGNORE INTO roles(username, role, created_at) VALUES (?,?,?)",
                (username, role, _now_iso()),
            )
            con.commit()
        finally:
            con.close()


def rbac_get_user_roles(username: str) -> List[str]:
    username = normalize_component(username).lower()
    if not username:
        return []
    with _DB_LOCK:
        con = _db()
        try:
            rows = con.execute(
                "SELECT role FROM roles WHERE username=? ORDER BY role", (username,)
            ).fetchall()
            return [str(r["role"]) for r in rows]
        finally:
            con.close()


# ============================================================
# AUDIT
# ============================================================
def audit_log(
    user: str, role: str, action: str, target: str = "", meta: Optional[dict] = None
) -> None:
    user = normalize_component(user).lower()
    role = normalize_component(role).upper()
    action = normalize_component(action)
    target = normalize_component(target)
    meta_json = json.dumps(meta or {}, ensure_ascii=False)

    with _DB_LOCK:
        con = _db()
        try:
            con.execute(
                "INSERT INTO audit(ts, user, role, action, target, meta_json) VALUES (?,?,?,?,?,?)",
                (_now_iso(), user, role, action, target, meta_json),
            )
            con.commit()
        finally:
            con.close()


# ============================================================
# FOLDER HELPERS
# ============================================================
def parse_folder_fields(folder_name: str) -> Dict[str, str]:
    folder_name = folder_name.strip()
    parts = folder_name.split("_")
    out = {"kdnr": "", "name": "", "addr": "", "plzort": ""}

    if not parts:
        return out

    if re.match(r"^\d{3,}$", parts[0]):
        out["kdnr"] = parts[0]
        rest = parts[1:]
    else:
        rest = parts

    plz_idx = None
    for i, t in enumerate(rest):
        if re.match(r"^\d{5}$", t):
            plz_idx = i
            break

    if plz_idx is not None:
        plz = rest[plz_idx]
        ort = "_".join(rest[plz_idx + 1 :]) if plz_idx + 1 < len(rest) else ""
        out["plzort"] = normalize_component(f"{plz} {ort.replace('_', ' ')}").strip()
        before = rest[:plz_idx]
    else:
        before = rest

    addr_start = None
    for i, t in enumerate(before):
        low = t.lower()
        if any(
            x in low
            for x in [
                "str",
                "straße",
                "strasse",
                "weg",
                "allee",
                "platz",
                "ring",
                "damm",
                "ufer",
                "gasse",
            ]
        ):
            addr_start = i
            break

    if addr_start is not None:
        out["name"] = normalize_component(" ".join(before[:addr_start]).replace("_", " "))
        out["addr"] = normalize_component(" ".join(before[addr_start:]).replace("_", " "))
    else:
        out["name"] = normalize_component(" ".join(before).replace("_", " "))

    return out


def find_existing_customer_folders(base_path: Path, kdnr: str) -> List[Path]:
    """
    Tenant-aware search without changing signature:
    - First: search direct child dirs: <base_path>/<kdnr>_*
    - If none: search one tenant level deep: <base_path>/<tenant>/<kdnr>_*
    """
    kdnr = normalize_component(kdnr)
    if not kdnr:
        return []
    base_path = Path(base_path)
    if not base_path.exists():
        return []

    out: List[Path] = []
    prefix = f"{kdnr}_"

    try:
        for p in base_path.iterdir():
            if p.is_dir() and p.name.startswith(prefix):
                out.append(p)
    except Exception:
        pass

    if out:
        return sorted(out)

    try:
        for tdir in base_path.iterdir():
            if not tdir.is_dir():
                continue
            if re.match(r"^\d{3,}_", tdir.name):
                continue
            for p in tdir.iterdir():
                if p.is_dir() and p.name.startswith(prefix):
                    out.append(p)
    except Exception:
        pass

    return sorted(out)


def best_match_object_folder(
    existing: List[Path], addr: str, plzort: str
) -> Tuple[Optional[Path], float]:
    addr_n = _norm_for_match(addr)
    plz_n = _norm_for_match(plzort)

    best: Optional[Path] = None
    best_score = 0.0

    for f in existing:
        fields = parse_folder_fields(f.name)
        a2 = _norm_for_match(fields.get("addr", ""))
        p2 = _norm_for_match(fields.get("plzort", ""))

        s1 = SequenceMatcher(None, addr_n, a2).ratio() if addr_n or a2 else 0.0
        s2 = SequenceMatcher(None, plz_n, p2).ratio() if plz_n or p2 else 0.0
        score = (0.6 * s2) + (0.4 * s1)

        if score > best_score:
            best_score = score
            best = f

    return best, best_score


def detect_object_duplicates_for_kdnr(
    kdnr: str, threshold: float = DEFAULT_DUP_SIM_THRESHOLD
) -> List[Dict[str, Any]]:
    kdnr = normalize_component(kdnr)
    folders = find_existing_customer_folders(BASE_PATH, kdnr)
    names = [(f, _norm_for_match(f.name)) for f in folders]
    out: List[Dict[str, Any]] = []

    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            f1, n1 = names[i]
            f2, n2 = names[j]
            if not n1 or not n2:
                continue
            sim = SequenceMatcher(None, n1, n2).ratio()
            if sim >= float(threshold):
                out.append(
                    {
                        "a": str(f1),
                        "b": str(f2),
                        "name_a": f1.name,
                        "name_b": f2.name,
                        "similarity": round(sim, 4),
                        "hint": "Mögliche Dublette (Admin: mergen/repair ohne Datenverlust).",
                    }
                )
    out.sort(key=lambda x: x["similarity"], reverse=True)
    return out


# ============================================================
# EXTRACTION / OCR
# ============================================================
def _extract_pdf_text(fp: Path) -> str:
    if PdfReader is None:
        return ""
    try:
        reader = PdfReader(str(fp))
        texts: List[str] = []
        for page in reader.pages[: max(1, OCR_MAX_PAGES)]:
            try:
                t = page.extract_text() or ""
                if t:
                    texts.append(t)
            except Exception:
                continue
        return "\n".join(texts).strip()
    except Exception:
        return ""


def _ocr_pdf(fp: Path) -> str:
    if fitz is None or pytesseract is None or Image is None:
        return ""
    try:
        doc = fitz.open(str(fp))
        texts: List[str] = []
        for i in range(min(len(doc), OCR_MAX_PAGES)):
            page = doc.load_page(i)
            pix = page.get_pixmap(dpi=250)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            txt = pytesseract.image_to_string(img, lang="deu+eng")
            if txt:
                texts.append(txt)
        return "\n".join(texts).strip()
    except Exception:
        return ""


def _ocr_image(fp: Path) -> str:
    if pytesseract is None or Image is None:
        return ""
    try:
        img = Image.open(str(fp))
        txt = pytesseract.image_to_string(img, lang="deu+eng")
        return (txt or "").strip()
    except Exception:
        return ""


def _extract_docx_text(fp: Path) -> str:
    if DocxDocument is not None:
        try:
            doc = DocxDocument(str(fp))
            paras = []
            for i, p in enumerate(doc.paragraphs):
                if i >= MAX_DOCX_PARAS:
                    break
                t = (p.text or "").strip()
                if t:
                    paras.append(t)
            return "\n".join(paras).strip()
        except Exception:
            pass

    try:
        with zipfile.ZipFile(str(fp), "r") as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        xml = xml.replace("</w:p>", "\n").replace("</w:tr>", "\n").replace("</w:tc>", " ")
        xml = re.sub(r"<[^>]+>", "", xml)
        xml = re.sub(r"[ \t]+", " ", xml)
        xml = re.sub(r"\n{3,}", "\n\n", xml)
        return xml.strip()
    except Exception:
        return ""


def _xlsx_shared_strings(z: zipfile.ZipFile) -> List[str]:
    try:
        sxml = z.read("xl/sharedStrings.xml").decode("utf-8", errors="ignore")
    except Exception:
        return []
    out = re.findall(r"<t[^>]*>(.*?)</t>", sxml, flags=re.IGNORECASE | re.DOTALL)
    return [re.sub(r"\s+", " ", x).strip() for x in out]


def _extract_xlsx_text(fp: Path) -> str:
    if openpyxl is not None:
        try:
            wb = openpyxl.load_workbook(str(fp), read_only=True, data_only=True)
            lines: List[str] = []
            for ws in wb.worksheets[:3]:
                lines.append(f"[Sheet] {ws.title}")
                rcount = 0
                for row in ws.iter_rows(values_only=True):
                    rcount += 1
                    if rcount > MAX_XLSX_ROWS:
                        break
                    vals = []
                    for v in row[:MAX_XLSX_COLS]:
                        if v is None:
                            continue
                        s = normalize_component(v)
                        if s:
                            vals.append(s)
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines).strip()
        except Exception:
            pass

    try:
        with zipfile.ZipFile(str(fp), "r") as z:
            sst = _xlsx_shared_strings(z)
            sheet_names = [
                n
                for n in z.namelist()
                if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
            ]
            sheet_names = sorted(sheet_names)[:3]
            out_lines: List[str] = []
            for sname in sheet_names:
                xml = z.read(sname).decode("utf-8", errors="ignore")
                out_lines.append(f"[SheetXML] {Path(sname).name}")
                vals = re.findall(r"<v>(.*?)</v>", xml, flags=re.IGNORECASE | re.DOTALL)
                cleaned: List[str] = []
                for v in vals[: MAX_XLSX_ROWS * 3]:
                    v = re.sub(r"\s+", " ", v).strip()
                    if not v:
                        continue
                    if v.isdigit():
                        idx = int(v)
                        if 0 <= idx < len(sst):
                            v = sst[idx]
                    cleaned.append(v)
                for i in range(0, min(len(cleaned), MAX_XLSX_ROWS * 10), MAX_XLSX_COLS):
                    chunk = cleaned[i : i + MAX_XLSX_COLS]
                    if chunk:
                        out_lines.append(" | ".join(chunk))
            return "\n".join(out_lines).strip()
    except Exception:
        return ""


def _extract_csv_text(fp: Path) -> str:
    raw = None
    for enc in ("utf-8-sig", "utf-8", "latin1"):
        try:
            raw = fp.read_text(encoding=enc, errors="strict")
            break
        except Exception:
            continue
    if raw is None:
        try:
            raw = fp.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""

    sio = io.StringIO(raw)
    try:
        sample = raw[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except Exception:
            dialect = csv.excel
        reader = csv.reader(sio, dialect)
        lines: List[str] = []
        for r_i, row in enumerate(reader):
            if r_i >= MAX_CSV_ROWS:
                break
            cols = [normalize_component(c) for c in row[:MAX_CSV_COLS]]
            cols = [c for c in cols if c]
            if cols:
                lines.append(" | ".join(cols))
        return "\n".join(lines).strip()
    except Exception:
        return _clip_text(raw.strip(), 50_000)


def _extract_eml_text(fp: Path) -> str:
    try:
        msg = BytesParser(policy=policy.default).parsebytes(_read_bytes(fp))
    except Exception:
        return ""

    hdr = []
    for k in ("Subject", "From", "To", "Cc", "Date"):
        try:
            v = msg.get(k)
            if v:
                hdr.append(f"{k}: {v}")
        except Exception:
            pass

    parts_text: List[str] = []

    def add_text(t: str) -> None:
        t = (t or "").strip()
        if t:
            parts_text.append(t)

    try:
        if msg.is_multipart():
            for part in msg.walk():
                ctype = (part.get_content_type() or "").lower()
                disp = (part.get_content_disposition() or "").lower()
                if disp == "attachment":
                    continue
                try:
                    payload = part.get_content()
                except Exception:
                    payload = None

                if ctype == "text/plain":
                    if isinstance(payload, str):
                        add_text(payload)
                elif ctype == "text/html":
                    if isinstance(payload, str):
                        add_text(_html_to_text(payload))
        else:
            ctype = (msg.get_content_type() or "").lower()
            try:
                payload = msg.get_content()
            except Exception:
                payload = None
            if isinstance(payload, str):
                add_text(_html_to_text(payload) if ctype == "text/html" else payload)
    except Exception:
        pass

    body = "\n\n".join(parts_text).strip()
    all_text = "\n".join(hdr + (["", body] if body else [])).strip()
    return all_text


def _extract_html_file(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
        return _html_to_text(raw)
    except Exception:
        return ""


def _extract_md_text(fp: Path) -> str:
    try:
        return (fp.read_text(encoding="utf-8", errors="ignore") or "").strip()
    except Exception:
        return ""


def _extract_rtf_text(fp: Path) -> str:
    raw = ""
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        try:
            raw = fp.read_text(encoding="latin1", errors="ignore")
        except Exception:
            return ""
    raw = re.sub(r"{\\\*[^}]*}", " ", raw)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def _extract_json_text(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""
    try:
        obj = json.loads(raw)
    except Exception:
        return _clip_text(raw.strip(), 50_000)

    out: List[str] = []

    def walk(x: Any, prefix: str = "") -> None:
        if isinstance(x, dict):
            for k, v in x.items():
                walk(v, f"{prefix}{k}.")
        elif isinstance(x, list):
            for i, v in enumerate(x[:2000]):
                walk(v, f"{prefix}{i}.")
        else:
            s = normalize_component(x)
            if s:
                out.append(f"{prefix}{s}")

    walk(obj)
    return "\n".join(out).strip()


def _extract_xml_text(fp: Path) -> str:
    try:
        raw = fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""
    raw = re.sub(r"(?is)<[^>]+>", " ", raw)
    raw = raw.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def _extract_msg_text(fp: Path) -> str:
    if extract_msg is None:
        return ""
    try:
        m = extract_msg.Message(str(fp))
        m.process()
        parts: List[str] = []
        if getattr(m, "subject", None):
            parts.append(f"Subject: {m.subject}")
        if getattr(m, "sender", None):
            parts.append(f"From: {m.sender}")
        if getattr(m, "to", None):
            parts.append(f"To: {m.to}")
        if getattr(m, "date", None):
            parts.append(f"Date: {m.date}")
        if getattr(m, "body", None):
            parts.append(str(m.body))
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _extract_text(fp: Path) -> Tuple[str, bool]:
    """
    Returns (text, used_ocr)
    """
    ext = fp.suffix.lower()

    if ext == ".txt":
        try:
            t = fp.read_text(encoding="utf-8", errors="ignore")
            return (t or "").strip(), False
        except Exception:
            return "", False

    if ext == ".md":
        t = _extract_md_text(fp)
        return _clip_text(t), False

    if ext == ".rtf":
        t = _extract_rtf_text(fp)
        return _clip_text(t), False

    if ext in (".html", ".htm"):
        t = _extract_html_file(fp)
        return _clip_text(t), False

    if ext == ".xml":
        t = _extract_xml_text(fp)
        return _clip_text(t), False

    if ext == ".json":
        t = _extract_json_text(fp)
        return _clip_text(t), False

    if ext == ".csv":
        t = _extract_csv_text(fp)
        return _clip_text(t), False

    if ext == ".docx":
        t = _extract_docx_text(fp)
        return _clip_text(t), False

    if ext == ".xlsx":
        t = _extract_xlsx_text(fp)
        return _clip_text(t), False

    if ext == ".eml":
        t = _extract_eml_text(fp)
        return _clip_text(t), False

    if ext == ".msg":
        t = _extract_msg_text(fp)
        return _clip_text(t), False

    if ext == ".pdf":
        t = _extract_pdf_text(fp)
        if len(t) >= MIN_TEXT_LEN_BEFORE_OCR:
            return _clip_text(t), False
        o = _ocr_pdf(fp)
        if o:
            return _clip_text(o), True
        return _clip_text(t), False

    if ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"):
        o = _ocr_image(fp)
        return _clip_text(o), True if o else False

    return "", False


# ============================================================
# HEURISTIC PARSING (SUGGESTIONS)
# ============================================================
_DOCTYPE_KEYWORDS = [
    (
        "H_RECHNUNG",
        [
            r"\bh[ _-]?rechnung\b",
            r"\bhändler[ _-]?rechnung\b",
            r"\bhaendler[ _-]?rechnung\b",
            r"\bdealer[ _-]?invoice\b",
        ],
    ),
    (
        "H_ANGEBOT",
        [
            r"\bh[ _-]?angebot\b",
            r"\bhändler[ _-]?angebot\b",
            r"\bhaendler[ _-]?angebot\b",
            r"\bdealer[ _-]?offer\b",
        ],
    ),
    ("RECHNUNG", [r"\brechnung\b", r"\binvoice\b"]),
    ("ANGEBOT", [r"\bangebot\b", r"\bquotation\b", r"\boffer\b"]),
    ("AUFTRAGSBESTAETIGUNG", [r"\bauftragsbest", r"\border confirmation\b"]),
    ("MAHNUNG", [r"\bmahnung\b", r"\breminder\b"]),
    ("NACHTRAG", [r"\bnachtrag\b"]),
    ("AW", [r"\baufmaß\b", r"\baufmass\b", r"\bmaß\b", r"\bmass\b"]),
    ("FOTO", [r"\bfoto\b", r"\bimage\b"]),
]


def _detect_doctype(text: str, filename: str) -> str:
    hay = f"{filename}\n{text}".lower()
    best = ("SONSTIGES", 0)
    for dt, pats in _DOCTYPE_KEYWORDS:
        score = 0
        for p in pats:
            if re.search(p, hay, flags=re.IGNORECASE):
                score += 1
        if score > best[1]:
            best = (dt, score)
    return best[0]


def _find_kdnr_candidates(text: str) -> List[Tuple[str, float]]:
    cands: List[str] = []
    for m in re.finditer(r"(kunden[\s\-]*nr\.?\s*[:#]?\s*)(\d{3,})", text, flags=re.IGNORECASE):
        cands.append(m.group(2))
    for m in re.finditer(r"(kdnr\.?\s*[:#]?\s*)(\d{3,})", text, flags=re.IGNORECASE):
        cands.append(m.group(2))

    if not cands:
        for m in re.finditer(r"\b(\d{4,6})\b", text):
            v = m.group(1)
            if v.startswith("20"):
                continue
            cands.append(v)

    freq: Dict[str, int] = {}
    for c in cands:
        freq[c] = freq.get(c, 0) + 1

    ranked = sorted(freq.items(), key=lambda x: (-x[1], x[0]))
    out: List[Tuple[str, float]] = []
    for num, f in ranked[:8]:
        score = round(min(1.0, 0.55 + 0.1 * f), 2)
        out.append((num, score))
    return out


def _find_dates(text: str) -> Tuple[str, List[Dict[str, Any]]]:
    cands: List[Dict[str, Any]] = []

    for m in re.finditer(r"\b(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})\b", text):
        raw = m.group(0)
        norm = parse_excel_like_date(raw)
        if norm:
            cands.append({"raw": raw, "date": norm, "reason": "DMY"})

    for m in re.finditer(r"\b(20\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})\b", text):
        raw = m.group(0)
        norm = parse_excel_like_date(raw)
        if norm:
            cands.append({"raw": raw, "date": norm, "reason": "YMD"})

    best = ""
    best_score = -1.0
    low = text.lower()

    for c in cands:
        raw = c["raw"]
        d = c["date"]
        idx = low.find(raw.lower())
        window = low[max(0, idx - 25) : idx + 25] if idx >= 0 else ""
        score = 0.1
        if "datum" in window:
            score += 0.9
        if "rechnung" in window or "angebot" in window:
            score += 0.2
        if score > best_score:
            best_score = score
            best = d

    seen = set()
    uniq: List[Dict[str, Any]] = []
    for c in cands:
        if c["date"] in seen:
            continue
        seen.add(c["date"])
        uniq.append(c)

    return best, uniq[:12]


def _find_name_addr_plzort(text: str) -> Tuple[List[str], List[str], List[str]]:
    lines = [normalize_component(x) for x in (text or "").splitlines()]
    lines = [line for line in lines if line]

    plzort: List[str] = []
    addr: List[str] = []
    name: List[str] = []

    for line in lines:
        m = re.search(r"\b(\d{5})\s+([A-Za-zÄÖÜäöüß\- ]{2,})\b", line)
        if m:
            candidate = normalize_component(f"{m.group(1)} {m.group(2)}")
            if candidate not in plzort:
                plzort.append(candidate)

    for line in lines:
        if re.search(
            r"\b(str\.?|straße|strasse|weg|allee|platz|ring|damm|ufer|gasse)\b",
            line,
            flags=re.IGNORECASE,
        ) and re.search(r"\b\d{1,4}[a-zA-Z]?\b", line):
            if line not in addr:
                addr.append(line)

    for line in lines[:15]:
        if len(line) < 3:
            continue
        if any(
            x in line.lower()
            for x in [
                "angebot",
                "rechnung",
                "datum:",
                "kunden-nr",
                "kunden nr",
                "projekt-nr",
                "bearbeiter",
            ]
        ):
            continue
        if re.search(r"(www\.|http|tel|fax|@)", line, flags=re.IGNORECASE):
            continue
        name.append(line)
        break

    return name[:8], addr[:8], plzort[:8]


# ============================================================
# INDEX / SEARCH
# ============================================================
def _fts_put(con: sqlite3.Connection, row: Dict[str, Any]) -> None:
    if not (_has_fts5(con) and _table_exists(con, "docs_fts")):
        return

    doc_id = str(row.get("doc_id", "") or "")
    if not doc_id:
        return

    con.execute("DELETE FROM docs_fts WHERE doc_id = ?", (doc_id,))
    con.execute(
        """
        INSERT INTO docs_fts(doc_id, kdnr, doctype, doc_date, file_name, file_path, content)
        VALUES (?,?,?,?,?,?,?)
        """,
        (
            doc_id,
            str(row.get("kdnr", "") or ""),
            str(row.get("doctype", "") or ""),
            str(row.get("doc_date", "") or ""),
            str(row.get("file_name", "") or ""),
            str(row.get("file_path", "") or ""),
            _clip_text(str(row.get("content", "") or ""), MAX_EXTRACT_CHARS),
        ),
    )


def _compute_group_key(kdnr: str, doctype: str, doc_date: str, file_name: str) -> str:
    k = normalize_component(kdnr)
    t = normalize_component(doctype).upper()
    d = parse_excel_like_date(doc_date) or ""
    stem = Path(file_name).stem
    stem = re.sub(r"_(v\d+|\d{6})$", "", stem, flags=re.IGNORECASE)
    stem_n = _norm_for_match(stem)
    raw = f"{k}|{t}|{d}|{stem_n}"
    return _sha256_bytes(raw.encode("utf-8"))


def index_upsert_document(
    *,
    doc_id: str,
    group_key: str,
    kdnr: str,
    object_folder: str,
    doctype: str,
    doc_date: str,
    file_name: str,
    file_path: str,
    extracted_text: str,
    used_ocr: bool,
    note: str = "",
) -> None:
    with _DB_LOCK:
        con = _db()
        try:
            exists = con.execute("SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)).fetchone()
            if not exists:
                con.execute(
                    "INSERT INTO docs(doc_id, group_key, kdnr, object_folder, doctype, doc_date, created_at) VALUES (?,?,?,?,?,?,?)",
                    (doc_id, group_key, kdnr, object_folder, doctype, doc_date or "", _now_iso()),
                )

            row = con.execute(
                "SELECT MAX(version_no) AS mx FROM versions WHERE doc_id=?", (doc_id,)
            ).fetchone()
            mx = int(row["mx"] or 0) if row else 0
            version_no = mx + 1

            con.execute(
                """
                INSERT INTO versions(doc_id, version_no, bytes_sha256, file_name, file_path, extracted_text, used_ocr, note, created_at)
                VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (
                    doc_id,
                    version_no,
                    doc_id,
                    file_name,
                    file_path,
                    _clip_text(extracted_text, MAX_EXTRACT_CHARS),
                    1 if used_ocr else 0,
                    note,
                    _now_iso(),
                ),
            )

            _fts_put(
                con,
                {
                    "doc_id": doc_id,
                    "kdnr": kdnr,
                    "doctype": doctype,
                    "doc_date": doc_date or "",
                    "file_name": file_name,
                    "file_path": file_path,
                    "content": extracted_text or "",
                },
            )
            con.commit()
        finally:
            con.close()


def assistant_search(
    query: str, kdnr: str = "", limit: int = ASSISTANT_DEFAULT_LIMIT, role: str = "ADMIN"
) -> List[Dict[str, Any]]:
    """
    Tenant note:
    - If you stored kdnr as "TENANT:1234", search by the same.
    - If you pass only "1234" and TENANT_DEFAULT is set, it will auto-prefix.
    """
    query = normalize_component(query)
    kdnr_in = normalize_component(kdnr)

    if kdnr_in and ":" not in kdnr_in:
        kdnr_in = _tenant_prefix_kdnr(TENANT_DEFAULT, kdnr_in)

    if not query:
        return []

    with _DB_LOCK:
        con = _db()
        try:
            use_fts = _has_fts5(con) and _table_exists(con, "docs_fts")
            rows: List[sqlite3.Row] = []

            if use_fts:
                tokens = [t for t in re.split(r"\s+", query) if t]
                q = " OR ".join(tokens) if tokens else query

                if kdnr_in:
                    rows = con.execute(
                        """
                        SELECT doc_id, kdnr, doctype, doc_date, file_name, file_path,
                               snippet(docs_fts, 6, '', '', ' … ', 12) AS snip
                        FROM docs_fts
                        WHERE docs_fts MATCH ? AND kdnr=?
                        LIMIT ?
                        """,
                        (q, kdnr_in, int(limit)),
                    ).fetchall()
                else:
                    rows = con.execute(
                        """
                        SELECT doc_id, kdnr, doctype, doc_date, file_name, file_path,
                               snippet(docs_fts, 6, '', '', ' … ', 12) AS snip
                        FROM docs_fts
                        WHERE docs_fts MATCH ?
                        LIMIT ?
                        """,
                        (q, int(limit)),
                    ).fetchall()
            else:
                like = f"%{query}%"
                if kdnr_in:
                    rows = con.execute(
                        """
                        SELECT d.doc_id, d.kdnr, d.doctype, d.doc_date,
                               v.file_name, v.file_path, substr(v.extracted_text,1,240) AS snip
                        FROM docs d
                        JOIN versions v ON v.doc_id=d.doc_id
                        WHERE d.kdnr=? AND (v.extracted_text LIKE ? OR v.file_name LIKE ? OR v.file_path LIKE ?)
                        GROUP BY d.doc_id
                        ORDER BY v.id DESC
                        LIMIT ?
                        """,
                        (kdnr_in, like, like, like, int(limit)),
                    ).fetchall()
                else:
                    rows = con.execute(
                        """
                        SELECT d.doc_id, d.kdnr, d.doctype, d.doc_date,
                               v.file_name, v.file_path, substr(v.extracted_text,1,240) AS snip
                        FROM docs d
                        JOIN versions v ON v.doc_id=d.doc_id
                        WHERE (v.extracted_text LIKE ? OR v.file_name LIKE ? OR v.file_path LIKE ?)
                        GROUP BY d.doc_id
                        ORDER BY v.id DESC
                        LIMIT ?
                        """,
                        (like, like, like, int(limit)),
                    ).fetchall()

            out: List[Dict[str, Any]] = []
            for r in rows:
                doc_id = str(r["doc_id"])
                vc = con.execute(
                    "SELECT COUNT(*) AS c FROM versions WHERE doc_id=?", (doc_id,)
                ).fetchone()
                version_count = int(vc["c"] or 0) if vc else 0
                out.append(
                    {
                        "doc_id": doc_id,
                        "kdnr": str(r["kdnr"] or ""),
                        "doctype": str(r["doctype"] or ""),
                        "doc_date": str(r["doc_date"] or ""),
                        "file_name": str(r["file_name"] or ""),
                        "file_path": str(r["file_path"] or ""),
                        "version_count": version_count,
                        "note": "",
                        "preview": str(r["snip"] or ""),
                    }
                )
            return out
        finally:
            con.close()


def index_run_full(base_path: Optional[Path] = None) -> Dict[str, Any]:
    base = Path(base_path) if base_path else BASE_PATH
    if not base.exists():
        return {
            "ok": True,
            "indexed": 0,
            "skipped": 0,
            "errors": 0,
            "skipped_by_reason": {
                "unsupported_ext": 0,
                "already_indexed": 0,
                "no_text": 0,
                "parse_failed": 0,
            },
        }

    indexed = 0
    skipped = 0
    errors = 0

    skipped_by_reason = {
        "unsupported_ext": 0,
        "already_indexed": 0,
        "no_text": 0,
        "parse_failed": 0,
    }

    try:
        for fp in base.rglob("*"):
            if not fp.is_file():
                continue

            ext = fp.suffix.lower()
            if ext not in SUPPORTED_EXT:
                skipped += 1
                skipped_by_reason["unsupported_ext"] += 1
                continue

            try:
                b = _read_bytes(fp)
                doc_id = _sha256_bytes(b)

                with _DB_LOCK:
                    con = _db()
                    try:
                        exists = con.execute(
                            "SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)
                        ).fetchone()
                    finally:
                        con.close()

                if exists:
                    skipped += 1
                    skipped_by_reason["already_indexed"] += 1
                    continue

                file_name = fp.name

                tenant = _effective_tenant(_infer_tenant_from_path(fp))

                kdnr_raw = ""
                object_folder = ""
                for part in reversed(fp.parts):
                    if re.match(r"^\d{3,}_", part):
                        kdnr_raw = part.split("_", 1)[0]
                        object_folder = part
                        break

                if TENANT_REQUIRE and not tenant:
                    skipped += 1
                    skipped_by_reason["parse_failed"] += 1
                    continue

                kdnr_idx = _tenant_prefix_kdnr(tenant, kdnr_raw) if kdnr_raw else ""
                object_folder_tag = (
                    _tenant_object_folder_tag(tenant, object_folder) if object_folder else ""
                )

                text, used_ocr = _extract_text(fp)
                if not text or len(text.strip()) < 3:
                    skipped += 1
                    skipped_by_reason["no_text"] += 1
                    continue

                doctype = _detect_doctype(text, file_name)
                best_date, _ = _find_dates(text)
                group_key = _compute_group_key(kdnr_idx, doctype, best_date, file_name)

                index_upsert_document(
                    doc_id=doc_id,
                    group_key=group_key,
                    kdnr=kdnr_idx,
                    object_folder=object_folder_tag,
                    doctype=doctype,
                    doc_date=best_date or "",
                    file_name=file_name,
                    file_path=str(fp),
                    extracted_text=text,
                    used_ocr=used_ocr,
                    note="indexed_by_full_scan",
                )
                indexed += 1

            except Exception:
                skipped += 1
                skipped_by_reason["parse_failed"] += 1
                continue

    except Exception:
        errors += 1

    return {
        "ok": True,
        "indexed": indexed,
        "skipped": skipped,
        "errors": errors,
        "skipped_by_reason": skipped_by_reason,
    }


# ============================================================
# BACKGROUND ANALYSIS -> PENDING
# ============================================================
def analyze_to_pending(src: Path) -> str:
    PENDING_DIR.mkdir(parents=True, exist_ok=True)

    src = Path(src)
    if not src.exists():
        raise FileNotFoundError(src)

    tenant = _effective_tenant(_infer_tenant_from_path(src))

    t = _token()
    payload: Dict[str, Any] = {
        "status": "ANALYZING",
        "progress": 1.0,
        "progress_phase": "Init…",
        "error": "",
        "path": str(src),
        "filename": src.name,
        "tenant_suggested": tenant,
        "used_ocr": False,
        "extracted_text": "",
        "preview": "",
        "doctype_suggested": "SONSTIGES",
        "doc_date_suggested": "",
        "doc_date_candidates": [],
        "kdnr_ranked": [],
        "name_suggestions": [],
        "addr_suggestions": [],
        "plzort_suggestions": [],
    }
    write_pending(t, payload)

    th = threading.Thread(target=_analyze_worker, args=(t,), daemon=True)
    th.start()
    return t


def start_background_analysis(src: Path) -> str:
    return analyze_to_pending(src)


def _set_progress(token: str, p: float, phase: str) -> None:
    d = read_pending(token) or {}
    d["progress"] = float(max(0.0, min(100.0, p)))
    d["progress_phase"] = phase
    write_pending(token, d)


def _analyze_worker(token: str) -> None:
    try:
        d = read_pending(token)
        if not d:
            return
        src = Path(d.get("path", ""))
        if not src.exists():
            d["status"] = "ERROR"
            d["error"] = "file_missing"
            d["progress"] = 0.0
            d["progress_phase"] = ""
            write_pending(token, d)
            return

        _set_progress(token, 5.0, "Datei lesen…")
        try:
            b = _read_bytes(src)
            doc_id = _sha256_bytes(b)
            d["doc_id"] = doc_id
        except Exception:
            d["doc_id"] = ""

        _set_progress(token, 18.0, "Text extrahieren…")
        text, used_ocr = _extract_text(src)
        d["used_ocr"] = bool(used_ocr)
        d["extracted_text"] = text
        d["preview"] = (text or "").strip()[:900].strip()

        _set_progress(token, 40.0, "Dokumenttyp/Datum erkennen…")
        d["doctype_suggested"] = _detect_doctype(text, src.name)
        best_date, date_cands = _find_dates(text)
        d["doc_date_suggested"] = best_date or ""
        d["doc_date_candidates"] = date_cands

        _set_progress(token, 60.0, "Kundendaten erkennen…")
        d["kdnr_ranked"] = _find_kdnr_candidates(text)
        names, addrs, plzs = _find_name_addr_plzort(text)
        d["name_suggestions"] = names
        d["addr_suggestions"] = addrs
        d["plzort_suggestions"] = plzs

        _set_progress(token, 85.0, "Fertigstellen…")
        d["status"] = "READY"
        d["progress"] = 100.0
        d["progress_phase"] = "Bereit"
        write_pending(token, d)

    except Exception as e:
        d = read_pending(token) or {}
        d["status"] = "ERROR"
        d["error"] = str(e)
        d["progress"] = 0.0
        d["progress_phase"] = ""
        write_pending(token, d)


# ============================================================
# ARCHIVE / PROCESS
# ============================================================
def _doctype_code(doctype: str) -> str:
    doctype = normalize_component(doctype).upper()

    if doctype in {"HAENDLERRECHNUNG", "HÄNDLERRECHNUNG", "H_RE"}:
        doctype = "H_RECHNUNG"
    if doctype in {"HAENDLERANGEBOT", "HÄNDLERANGEBOT", "H_ANG"}:
        doctype = "H_ANGEBOT"

    return {
        "H_RECHNUNG": "H_RE",
        "H_ANGEBOT": "H_ANG",
        "RECHNUNG": "RE",
        "ANGEBOT": "ANG",
        "AUFTRAGSBESTAETIGUNG": "AB",
        "AW": "AW",
        "MAHNUNG": "MAH",
        "NACHTRAG": "NTR",
        "FOTO": "FOTO",
        "SONSTIGES": "DOC",
    }.get(doctype, "DOC")


def _compose_object_folder(kdnr: str, name: str, addr: str, plzort: str) -> str:
    parts = [_safe_fs(kdnr), _safe_fs(name), _safe_fs(addr), _safe_fs(plzort)]
    parts = [p for p in parts if p]
    return "_".join(parts)[:180]


def _compose_filename(
    doctype: str, doc_date: str, kdnr: str, name: str, addr: str, plzort: str, ext: str
) -> str:
    code = _doctype_code(doctype)
    d = parse_excel_like_date(doc_date) or ""
    parts: List[str] = [code]
    if d:
        parts.append(d)
    for x in [kdnr, name, addr, plzort]:
        xs = _safe_fs(x)
        if xs:
            parts.append(xs)

    base = "_".join(parts)
    base = re.sub(r"_+", "_", base).strip("_")[:160]
    return f"{base}{ext}"


def _next_version_suffix(target_dir: Path, base_name: str, ext: str) -> str:
    stem = Path(base_name).stem
    n = 2
    while True:
        cand = f"{stem}_v{n}{ext}"
        if not (target_dir / cand).exists():
            return cand
        n += 1


def _db_has_doc(doc_id: str) -> bool:
    with _DB_LOCK:
        con = _db()
        try:
            r = con.execute("SELECT doc_id FROM docs WHERE doc_id=?", (doc_id,)).fetchone()
            return bool(r)
        finally:
            con.close()


def process_with_answers(src: Path, answers: Dict[str, Any]) -> Tuple[Path, Path, bool]:
    src = Path(src)
    if not src.exists():
        raise FileNotFoundError(src)

    tenant = _effective_tenant(
        answers.get("tenant"), answers.get("mandant"), _infer_tenant_from_path(src)
    )
    if TENANT_REQUIRE and not tenant:
        raise ValueError("tenant/mandant missing (TENANT_REQUIRE=1)")

    kdnr_raw = normalize_component(answers.get("kdnr", ""))
    use_existing = normalize_component(answers.get("use_existing", ""))
    name = normalize_component(answers.get("name", ""))
    addr = normalize_component(answers.get("addr", ""))
    plzort = normalize_component(answers.get("plzort", ""))
    doctype = normalize_component(answers.get("doctype", "SONSTIGES")).upper()
    doc_date = parse_excel_like_date(answers.get("document_date", "")) or ""

    if not kdnr_raw:
        raise ValueError("kdnr missing")

    kdnr_idx = _tenant_prefix_kdnr(tenant, kdnr_raw)

    BASE_PATH.mkdir(parents=True, exist_ok=True)
    tenant_dir = BASE_PATH / _safe_fs(tenant) if tenant else BASE_PATH
    tenant_dir.mkdir(parents=True, exist_ok=True)

    created_new_object = False
    if use_existing:
        folder = Path(use_existing)
        if not folder.exists() or not folder.is_dir():
            folder = tenant_dir / _compose_object_folder(kdnr_raw, name, addr, plzort)
            created_new_object = True
    else:
        folder_name = _compose_object_folder(kdnr_raw, name, addr, plzort)
        folder = tenant_dir / folder_name
        if not folder.exists():
            created_new_object = True

    folder.mkdir(parents=True, exist_ok=True)

    ext = src.suffix.lower()
    final_name = _compose_filename(doctype, doc_date, kdnr_raw, name, addr, plzort, ext)
    target = folder / final_name

    b = _read_bytes(src)
    doc_id = _sha256_bytes(b)

    object_folder_tag = _tenant_object_folder_tag(tenant, folder.name)

    if target.exists():
        try:
            if _sha256_bytes(_read_bytes(target)) == doc_id:
                try:
                    src.unlink()
                except Exception:
                    pass

                if not _db_has_doc(doc_id):
                    text, used_ocr = _extract_text(target)
                    group_key = _compute_group_key(kdnr_idx, doctype, doc_date, target.name)
                    index_upsert_document(
                        doc_id=doc_id,
                        group_key=group_key,
                        kdnr=kdnr_idx,
                        object_folder=object_folder_tag,
                        doctype=doctype,
                        doc_date=doc_date or "",
                        file_name=target.name,
                        file_path=str(target),
                        extracted_text=text,
                        used_ocr=used_ocr,
                        note="dedupe_same_bytes_existing_target",
                    )
                return folder, target, created_new_object
        except Exception:
            pass

        final_name = _next_version_suffix(folder, final_name, ext)
        target = folder / final_name

    try:
        src.replace(target)
    except Exception:
        target.write_bytes(b)
        try:
            src.unlink()
        except Exception:
            pass

    text, used_ocr = _extract_text(target)
    group_key = _compute_group_key(kdnr_idx, doctype, doc_date, target.name)

    note = ""
    with _DB_LOCK:
        con = _db()
        try:
            g = con.execute(
                "SELECT doc_id FROM docs WHERE group_key=? LIMIT 1", (group_key,)
            ).fetchone()
            if g and str(g["doc_id"]) != doc_id:
                note = "new_version_same_group_key"
        finally:
            con.close()

    index_upsert_document(
        doc_id=doc_id,
        group_key=group_key,
        kdnr=kdnr_idx,
        object_folder=object_folder_tag,
        doctype=doctype,
        doc_date=doc_date or "",
        file_name=target.name,
        file_path=str(target),
        extracted_text=text,
        used_ocr=used_ocr,
        note=note,
    )

    return folder, target, created_new_object


# ============================================================
# BOOTSTRAP DIRS (on import)
# ============================================================
def _bootstrap_dirs() -> None:
    EINGANG.mkdir(parents=True, exist_ok=True)
    BASE_PATH.mkdir(parents=True, exist_ok=True)
    PENDING_DIR.mkdir(parents=True, exist_ok=True)
    DONE_DIR.mkdir(parents=True, exist_ok=True)


_bootstrap_dirs()
# db_init() is intentionally not auto-called here; your Flask runner calls db_init() at startup.
